"""Build the IDF-B disclosure for length-invariance recalibration of roadside detectors.

Format: the inventor's accepted VIT IDF-B (masthead, eleven numbered sections with the
subsections of the accepted cyber-security IDF-B, captioned figures, claim tree), through the layout helpers of the CausalInterlock generator,
imported unchanged. Drafting rules from the guidance files in this folder are applied:
problem, definitions, split patent and non-patent tables, feature mapping, technical
effects, claimed functions against the example, numbered paragraphs, validation with
configuration, tests, effect summary and limits, a discussion, six protection aspects
before the claims, and no reference to patentability law in the text.

Every measured number is read from results/*.json at build time, and the operating
constants are read from the code that used them. The build refuses to run if a results
file is missing. After writing, it scans the document for placeholders, business-method
vocabulary, the banned phrases of writing-guidelines.md and section_3k_compliance.md, and
checks that more than 60 % of the paragraphs of Sections 6 and 7 name a physical element.
"""
import datetime
import json
import os
import re
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Inches, Pt, RGBColor

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
RESULTS = os.path.join(ROOT, "results")
sys.path.insert(0, os.path.join(ROOT, "src"))
import e3_steps                                                   # noqa: E402
import e8_offset_model as e8m                                     # noqa: E402
import pairs                                                      # noqa: E402

REF = r"C:\adi\patents\01_02_infosec\IDF-B_CausalInterlock"
sys.path.insert(0, REF)
import build_idfb as B                                            # noqa: E402

B.FIG = os.path.join(ROOT, "patent", "figures")
OUT = os.path.join(ROOT, "patent", "IDF-B_Length_Invariance_Recalibration.docx")

h1, h2, para, bullets, _run = B.h1, B.h2, B.para, B.bullets, B._run

FORBIDDEN = ["revenue", "cost saving", "billing", "pricing", "tariff", "customer",
             "marketing", "monetis", "monetiz", "business", "commercial benefit", r"\bROI\b"]
BANNED = ["it is worth noting", "in the realm of", "leverag", "harness", "utiliz",
          "plays a crucial", "pivotal role", "this approach ensures", "in order to achieve",
          "seamless", "robust", "state-of-the-art", "state of the art", "novel approach",
          "cutting-edge", "it should be noted", r"3\(k\)", "per se", "computer programme",
          "software module", "simulation", "virtual", "mathematical method", "algorithm for",
          "python", "patents act"]
OPENERS = re.compile(r"^(Furthermore|Moreover|Additionally)\b")
PLACEHOLDER = re.compile(r"\bTBD\b|\bTODO\b|\bXX+\b|\[name\]|\?\?")
PHYSICAL = re.compile(r"detector|loop|site|carriageway|link \(130\)|archive|outstation|"
                      r"card|vehicle|lane|slip road", re.I)

HDR_FILL, ALT_FILL, BORDER = "1565C0", "E3F2FD", "B0BEC5"


# ------------------------------------------------------------------ layout helpers
class Numbering:
    """Bracketed paragraph numbers, continued from Section 6 into Section 7."""
    def __init__(self):
        self.n = 0
        self.texts = []

    def __call__(self, doc, text):
        self.n += 1
        self.texts.append(text)
        return para(doc, text, lead="[%04d]" % self.n)


def figure(doc, name, title, expl, width=6.4):
    """Picture, then one caption in the style of the patent 2 IDF-B: 'Figure N.' bold,
    title and explanation following in the same paragraph."""
    num, rest = title.split(":", 1)
    B.FIG = B.FIG
    B.figure(doc, name, "%s. %s. %s" % (num, rest.strip(), expl), width=width)


def _borders(t):
    tblPr = t._tbl.tblPr
    bd = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement("w:" + side)
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), "4")
        e.set(qn("w:space"), "0")
        e.set(qn("w:color"), BORDER)
        bd.append(e)
    tblPr.append(bd)


def table(doc, headers, rows, cap, expl=None, widths=None, size=9.5, align_left_cols=()):
    """Plain grid table with bold header and caption below (patent 2 style)."""
    t = B.table(doc, headers, rows, cap=cap, widths=widths, size=size,
                align_left_cols=align_left_cols)
    # keep each row's last lines with the caption so the caption never stands alone
    for row in t.rows:
        trpr = row._tr.get_or_add_trPr()
        cs = OxmlElement("w:cantSplit")
        trpr.append(cs)
    # a table short enough for one page is kept whole; longer ones keep only their tail
    for row in (t.rows if len(t.rows) <= 18 and not cap.startswith(("Table 1:", "Table 2:"))
                else t.rows[-2:]):
        for c in row.cells:
            for pp in c.paragraphs:
                pp.paragraph_format.keep_with_next = True
    return t


def load(name):
    path = os.path.join(RESULTS, name)
    if not os.path.exists(path):
        raise SystemExit("results absent, refusing to build with invented figures: %s"
                         % name)
    with open(path, encoding="utf-8") as fh:
        return json.load(fh)


def pt(x, nd=1):
    """A share difference (fraction) as points."""
    return ("%." + str(nd) + "f") % (100.0 * x)


def span(xs):
    """'4' when a range collapses, '6-7' otherwise."""
    lo, hi = min(xs), max(xs)
    return "%d" % lo if lo == hi else "%d-%d" % (lo, hi)


def day_label(idx):
    d = datetime.date(2025, 3, 1) + datetime.timedelta(days=idx)
    return "%d %s 2025" % (d.day, d.strftime("%B"))


def claim(doc, n, text, independent=False):
    lead = "Claim %d%s." % (n, " (independent)" if independent else "")
    para(doc, text, lead=lead, indent=0.28, size=10.5)


def bias_spread(fit):
    b = sorted(abs(v["bias"][0]) for v in fit.values() if v["chain_size"] >= 3)
    n = len(b)
    return {"n": n, "q50": b[n // 2], "q90": b[int(0.9 * n)],
            "over1": sum(x > 0.01 for x in b) / n, "over5": sum(x > 0.05 for x in b) / n}


# ------------------------------------------------------------------ prior art
# Every row was written from the reference's raw text (patent/prior_art_raw/). Patents are
# cited in the text by number, publications by author, with the row tag in brackets.
PATENTS = [
    ["P1", "US 6,876,949 B2 (2005); divisional US 6,988,052 B2 (2006)",
     "Normalization of inductive vehicle detector outputs (Inductive Signature "
     "Technologies)",
     "Normalises several inductive detectors so that one vehicle reads alike at each, by "
     "probe vehicles or by ordinary vehicles re-identified at two detectors, the "
     "coefficients solved by linear algebra; and calibrates each detector's inductive "
     "length so that its local vehicle population distribution matches a standard table, "
     "which also exposes a wrong loop spacing. This is the closest reference on "
     "calibration from the vehicle stream, and it is conceded that a length error shifts "
     "the population distribution. It needs per-vehicle signatures, re-identification or "
     "a standard table; it does not use class counts, admit pairs, solve a chain, test "
     "whether one offset explains every class boundary, or correct recorded class counts."],
    ["P2", "GB 2385138 A (2003)",
     "Validation and calibration of vehicle loop detection systems (Golden River Traffic)",
     "Switched validation loops placed over the loop sensors simulate a vehicle. The "
     "background names loop geometry, turns, detector threshold behaviour and installation "
     "irregularities as causes of systematic over- or under-reading of length that "
     "inspection does not reveal. Test apparatus at each site; background to the "
     "technical problem."],
    ["P3", "GB 2377027 A (2002)",
     "Assessing the accuracy of roadside traffic measurement systems (Golden River "
     "Traffic)",
     "An instrumented probe vehicle records its own parameters, among them vehicle "
     "length, and the moment it passes a traffic monitoring station, whose measurements "
     "are then compared with them. Assessment by attendance with a known vehicle, one "
     "station at a time; no use of neighbouring stations and no correction of counts."],
    ["P4", "US 7,684,946 B2 (2010)",
     "Weigh-in-motion system with auto-calibration (Mettler-Toledo)",
     "Calibration factors of a weigh-in-motion scale are computed automatically from the "
     "weights of the same vehicles read by the scale and by a static scale. A second, "
     "reference instrument on the same vehicles; weight, not length; no class counts, no "
     "chain of sites, no test of the kind of fault."],
    ["P5", "EP 2790165 A1 (2014, withdrawn)",
     "Quality determination in data acquisition (SWARCO Traffic Systems)",
     "Models presence-pulse sequences with regard to sensor topology and environment, "
     "classifies and corrects errors, recommends detector sensitivity settings, with "
     "trained models including a truck detector. Pulse-level data and learned models; "
     "no comparison of length-class shares of detectors seeing the same vehicles and no "
     "bias per class boundary."],
    ["P6", "US 11,461,692 B2 (2022)",
     "Single-loop vehicle classification using inductive loop vehicle signatures (CLR "
     "Analytics)",
     "Classifies vehicles from single-loop signature waveforms against a signature "
     "library with a nearest-neighbour classifier; notes that length-based methods need "
     "re-calibration between stations. Classification at one site from signatures; no "
     "neighbour reference and no correction of recorded class counts."],
    ["P7", "US 6,999,886 B2 (2006)",
     "Vehicle speed estimation using inductive vehicle detection systems (Inductive "
     "Signature Technologies)",
     "A single-loop detector card synthesises a second output to mimic a two-loop speed "
     "trap so that a field controller can estimate speed and length. Measurement at one "
     "site; no calibration between sites."],
]

NPL = [
    ["N1", "A. Azin, X. Yang, Utah Department of Transportation, report UT-20.15, 2020",
     "Multi-Stage Algorithm for Detection-Error Identification and Data Screening",
     "Compares the average effective vehicle length (AEVL) distribution of each station "
     "with its adjacent stations, singles out the station that differs, uses its upstream "
     "and downstream neighbours to decide which is faulty and marks it for maintenance or "
     "recalibration; notes that on-ramps alter AEVL downstream; cites Lu, Yang and Chang "
     "(2014, TRR 2443) for the comparison of AEVL distributions. Closest reference on "
     "neighbour comparison: comparison, localisation and flagging are conceded. It does "
     "not size a signed error per detector at each class boundary, recover a length "
     "offset, test the kind of fault, correct any recorded count or date a change; AEVL "
     "mixes speed, occupancy and length."],
    ["N2", "D. J. Dailey, Z. Wall, Washington State DOT, report WA-RD 633.1, 2006 "
     "(describes Wall and Dailey, TRR 1855, 2003)",
     "A General Simulation Model for Use with Real Freeway Data to Perform Congestion "
     "Prediction, Phase 3",
     "Corrects archived loop volumes of poorly calibrated stations: candidate reference "
     "pairs have no on- or off-ramp between them, a reference station whose counts "
     "balance with its neighbour is found, adjacent stations are compared with it in turn, "
     "and a correction time series is applied. Conceded for the slip-free pair, the count "
     "balance, a reference station and correction of the archive. Total volumes only: no "
     "length classes, no chain solution, no length offset, no fault class."],
    ["N3", "N. L. Nihan, Y. Wang, P. Cheevarunothai, Washington State DOT, report "
     "WA-RD 647.1, 2006",
     "Improving Dual-Loop Truck (and Speed) Data",
     "Identifies dual-loop sensitivity problems from per-vehicle loop event data and "
     "retunes loop sensitivity against the short-vehicle length distribution of Wang and "
     "Nihan (2004); the ALEDA tool runs on a laptop at the controller cabinet. Needs "
     "per-vehicle event data, an external distribution and a visit to each station."],
    ["N4", "J. Lee, B. Coifman, Journal of Transportation Engineering, 138(7), 2012",
     "Quantifying Loop Detector Sensitivity and Correcting Detection Problems on Freeways",
     "Assesses each loop's sensitivity from its daily median on-time, assuming most "
     "vehicles are passenger vehicles of an expected length at free-flow speed; corrects "
     "speeds by a multiplicative factor. One detector at a time from per-vehicle on-times "
     "and an assumed length; no neighbour reference, no class counts, no fault test."],
    ["N5", "E. Minge et al., FHWA Pooled Fund TPF-5(192), MnDOT report 2012-33, 2012",
     "Loop- and Length-Based Vehicle Classification",
     "Recommends probe-vehicle runs as the primary calibration: vehicles of known speed "
     "and length are driven through the site and the detector is adjusted to match; "
     "documents causes of length error. Attendance at each site; no neighbours, no "
     "correction of recorded counts."],
    ["N6", "B. Coifman, S. Kim, Transportation Research Part C, 2009",
     "Speed estimation and length based vehicle classification from freeway single-loop "
     "detectors",
     "Speed and length classes from single-loop on-times, with a recalibration from "
     "successive vehicles' on-time ratios. Within one detector's own per-vehicle data; no "
     "comparison between detectors."],
    ["N7", "L. Wu, B. Coifman, Transportation Research Part C, 2014",
     "Improved vehicle classification from dual-loop detectors in congested traffic",
     "Dual-loop length measurement assumes negligible acceleration and degrades in "
     "congestion; the error is bounded from speed and acceleration. Basis of the "
     "free-flow condition used in admission; conceded."],
    ["N8", "X.-Y. Lu, P. Varaiya, R. Horowitz, J. Palen, California PATH (manuscript, "
     "file dated 2008)",
     "Faulty loop data analysis/correction and loop fault detection",
     "Reviews loop fault detection and data correction, including a test on the ratio of "
     "upstream and downstream flows and on-time threshold tests. Flows and on-times; no "
     "length-class comparison, no bias per class boundary, no reclassification."],
    ["N9", "V. Bychkovskiy, S. Megerian, D. Estrin, M. Potkonjak, Proc. IPSN, 2003",
     "A Collaborative Approach to In-Place Sensor Calibration",
     "Relative calibration of neighbouring sensors observing the same phenomenon, made "
     "consistent over groups; absolute calibration needs a calibrated subset. General "
     "principle of the chain solution, conceded; a continuous field, no class "
     "boundaries, no counts."],
    ["N10", "L. Balzano, R. Nowak, Proc. IPSN, 2007",
     "Blind Calibration of Sensor Networks",
     "Recovers unknown sensor gains from routine measurements of oversampled signals; the "
     "offset component in the signal subspace is generally not identifiable. Consistent "
     "with the need for a reference per chain; no traffic detectors."],
    ["N11", "FHWA, report FHWA-PL-18-015, 2018",
     "Weigh-in-Motion Pocket Guide, Part 3: WIM Calibration and Maintenance Guide",
     "Calibration of weigh-in-motion sites with test trucks of known weight and length; "
     "manufacturers' auto-calibration routines; drift monitored through the front-axle "
     "weight distribution of a common lorry class in the site's own traffic; length "
     "requirement of +-1.5 ft or 3 %. Conceded: checking a site against statistics of its "
     "own traffic. One site at a time, for weight, against an assumed fleet property."],
]


# ------------------------------------------------------------------ the document
def build():
    probe = load("probe.json")["roads"]
    e1 = load("e1_out_of_sample.json")
    e2 = load("e2_injection.json")
    e3 = load("e3_step_events.json")
    e4 = load("e4_ablation.json")
    e5 = load("e5_refit.json")
    fl = load("flag_rules.json")
    e7 = load("e7_timing.json")
    e8 = load("e8_offset_model.json")
    e9 = load("e9_consistency.json")
    e10 = load("e10_checks.json")
    e11 = load("e11_ground_truth.json")
    e12 = load("e12_tolerance.json")
    notes = load("idfb_fig_notes.json")
    fits = {r: load("fit_march_%s.json" % r) for r in ("M25", "M1")}
    spread = {r: bias_spread(fits[r]) for r in fits}
    for key in ("fig05", "fig08"):
        if key not in notes:
            raise SystemExit("figure notes missing %s; run make_figs.py" % key)
    if "by_chain_size" not in e2["roads"]["M25"]:
        raise SystemExit("e2_injection.json predates the chain-size breakdown; rerun it")
    if "per_detector" not in e8["roads"]["M25"]:
        raise SystemExit("e8_offset_model.json lacks per-detector values; rerun it")

    R = ("M25", "M1")
    MM = ["2025-03", "2025-04", "2025-05", "2025-06"]
    p25, p1 = probe["M25"], probe["M1"]
    E1 = e1["roads"]
    months = e1["eval_months"]
    mname = {"2025-03": "March", "2025-04": "April", "2025-05": "May", "2025-06": "June"}

    def e1v(road, m, pol="all/additive"):
        return E1[road][m][pol]["mean_abs_gap_touched"][0]

    def e1raw(road, m):
        return E1[road][m]["all/additive"]["mean_abs_gap_touched_raw"][0]

    def e1all(road, m, pol="all/additive"):
        return E1[road][m][pol]["mean_abs_gap_all"][0]

    def e1allraw(road, m):
        return E1[road][m]["none"]["mean_abs_gap_all"][0]

    def e5v(road, kind, m, key):
        return e5["roads"][road][kind][m][key]

    for r in R:
        for m in months:
            if not e1v(r, m) < e1v(r, m, "all/ratio"):
                raise SystemExit("text says constant-share beats proportional; data do not")
            if not (E1[r][m]["all/additive"]["mean_abs_gap_all"][0]
                    <= E1[r][m]["flagged/additive"]["mean_abs_gap_all"][0]):
                raise SystemExit("text says correcting all is no worse; data do not")
            c = e8["roads"][r]["correction"][m]
            if not (c["per_boundary"]["touched_corrected"][1]
                    < c["rank_one"]["touched_corrected"][1]):
                raise SystemExit("text says per-boundary beats one offset at 660; data do not")

    red_t = [1 - e1v(r, m) / e1raw(r, m) for r in R for m in months]
    red_a = [1 - e1all(r, m) / e1allraw(r, m) for r in R for m in months]
    p90rel = [c["p90_size_error_over_injection"] for r in R
              for c in e2["roads"][r]["by_chain_size"].values()]
    locsize = [c["localised"] for r in R for c in e2["roads"][r]["by_chain_size"].values()]
    locsz = [e2["roads"][r]["additive %+.2f" % s]["localised"] for r in R for s in e2["sizes"]]
    ev_all = [(r, e) for r in R for e in e3["roads"][r]["events"]]
    n_tested = sum(e3["roads"][r]["detectors_tested"] for r in R)
    n_ctrl = sum(e3["roads"][r]["control_trials"] for r in R)
    n_ctrl_ev = sum(e3["roads"][r]["control_events"] for r in R)
    cont = notes["fig08"]["continuity_5_days_either_side"]
    shown = notes["fig08"]["shown"]
    ev332 = [e for r, e in ev_all if (r, e["detector"]) == tuple(shown[0])][0]
    ch = notes["fig05"]
    t7 = e7["roads"]
    se25 = e4["variants"]["claimed"]["M25"]["median_se_520"]
    fl25 = fl["M25"]

    # offset model (E8)
    S8 = {r: e8["roads"][r]["structure"] for r in R}
    rank1 = [S8[r][m]["rank1_variance"] for r in R for m in MM]
    load660 = {r: [S8[r][m]["loadings"][1] for m in MM] for r in R}
    pred660 = {r: [S8[r][m]["cdf_density_ratio"][1] for m in MM] for r in R}
    off50 = [S8[r][m]["offset_cm_q50_q75_q90"][0] for r in R for m in MM]
    off90 = [S8[r][m]["offset_cm_q50_q75_q90"][2] for r in R for m in MM]
    SP = {r: e8["roads"][r]["speed"] for r in R}
    spx = [SP[r][m]["offset_cm_explained_by_speed_q50"] for r in R for m in MM]
    spf = [SP[r][m]["share_with_speed_bias_under_1pct"] for r in R for m in MM]
    PE = {r: e8["roads"][r]["fault_type_persistence"] for r in R}
    C8 = {r: e8["roads"][r]["correction"] for r in R}
    evsteps = {(r, e["detector"]): e for r in R for e in e8["roads"][r]["events"]}
    both_pairs = [e for e in evsteps.values()
                  if e["step_pair_downstream"] and e["step_pair_upstream"]]
    ratio_agree = sum(1 for e in both_pairs
                      if abs(e["step_pair_downstream"][1] / e["step_pair_downstream"][0]
                             - e["step_pair_upstream"][1] / e["step_pair_upstream"][0]) < 0.1)

    # cross-boundary test (E9)
    F9 = {r: e9["roads"][r]["field"] for r in R}
    I9 = {r: e9["roads"][r]["injection"] for r in R}
    c660 = [F9[r][m]["corr_660"] for r in R for m in MM]
    r2_660 = [F9[r][m]["variance_660_explained_by_prediction"] for r in R for m in MM]
    agree9 = [F9[r][m]["agreement_with_principal_direction_big"] for r in R for m in MM]
    det660 = [I9[r][k]["rate"] for r in R for k in ("boundary660 +100 cm", "boundary660 -100 cm")]
    det1160 = [I9[r][k]["rate"] for r in R for k in ("boundary1160 +150 cm",
                                                     "boundary1160 -150 cm")]
    keep40 = [I9[r][k]["rate"] for r in R for k in ("offset +40 cm", "offset -40 cm")]
    gauge9 = [tuple(int(x) for x in v.split("/")) for r in R
              for v in I9[r]["common_offset_class_changes"].values()]
    gauge_changed = sum(g[0] for g in gauge9)
    gauge_units = sum(g[1] for g in gauge9)      # units x two signs, both corridors
    P9 = {r: F9[r]["persistence"] for r in R}
    if not min(det660) > max(1 - x for x in keep40) or gauge_changed > 0.02 * gauge_units:
        raise SystemExit("text says the cross-boundary test separates faults and ignores "
                         "the chain constant; data do not")

    # review checks (E10)
    X = {r: e10["roads"][r] for r in R}
    nz = [X[r]["null"][m] for r in R for m in MM]
    std_real = [x["rank1_standardised"] for x in nz]
    std_null = [x["rank1_standardised_null_mean"] for x in nz]
    std_p99 = [x["rank1_standardised_null_p99"] for x in nz]
    raw_null = [x["rank1_null_mean"] for x in nz]
    r660 = [x["corr_520_660"] for x in nz]
    r1160 = [x["corr_520_1160"] for x in nz]
    if not min(std_real) > max(std_p99):
        raise SystemExit("text says the one-offset structure beats the shuffled baseline")
    drv = [X[r]["drivers"][m] for r in R for m in MM]
    off_fail = [x["median_abs_offset_cm_failing"] for x in drv]
    off_pass = [x["median_abs_offset_cm_passing"] for x in drv]
    iv = [v for r in R for m in MM for k, v in X[r]["interp"][m].items() if k != "baseline"]
    iv_ratio = [v["offset_ratio_to_baseline_median"] for v in iv]
    iv_agree = [v["class_agreement_with_baseline"] for v in iv]
    ev10 = [e for r in R for e in X[r]["events"]]
    ev_pass = [e for e in ev10 if e["passes"]]
    ev_fail = [e for e in ev10 if not e["passes"]]
    J = {r: X[r]["junction"] for r in R}
    jm = months

    def jr(r, k):
        return [J[r][m][k] for m in jm]

    if any(J["M25"][m]["improved"] != J["M25"][m]["pairs"] for m in jm):
        raise SystemExit("text says every M25 junction pair improves; data do not")
    if any(abs(e["step"][0]) <= 0.30 for e in ev_fail):
        raise SystemExit("text says failing events step by over 30 points; data do not")

    # ground truth (E11): manual classified counts of the Department for Transport
    G = e11["pooled"]
    G520, G660 = G["b520"], G["b660"]
    GR = {r: e11["roads"][r] for r in R}
    if not (abs(G520["fit_raw_deviation_on_bias"]["beta"] - 1)
            < 2 * G520["fit_raw_deviation_on_bias"]["beta_se"]
            and G520["mad_corrected_pts"] < G520["mad_raw_pts"]
            and G660["mad_corrected_pts"] < G660["mad_raw_pts"]):
        raise SystemExit("text says the biases match the manual counts; data do not")

    tol = [(v["units_over_tolerance"], v["units"]) for r in R
           for v in e12["roads"][r].values()]
    tol_share = [a / b for a, b in tol]

    e8_CDF_TOP = e8m.CDF_TOP

    # admission conditions (E4): fit made with the balance test or free-flow condition dropped
    V4 = e4["variants"]

    def e4v(var, r, m):
        return V4[var][r]["scores"][m]["corrected_mean_abs_gap_520"]

    nogate_m1 = [e4v("no volume gate", "M1", m) for m in months]
    claimed_m1 = [e4v("claimed", "M1", m) for m in months]
    nogate_m25_diff = max(abs(e4v("no volume gate", "M25", m) - e4v("claimed", "M25", m))
                          for m in months)
    noff_worse = all(e4v("no free-flow condition", r, m) > e4v("claimed", r, m)
                     for r in R for m in months)
    if not (all(a > b for a, b in zip(nogate_m1, claimed_m1)) and noff_worse
            and nogate_m25_diff < 0.0002):
        raise SystemExit("text says the admission conditions help; data do not")
    gate_pairs_m25 = max(V4["no volume gate"]["M25"]["scores"][m]["touched_pairs"]
                         - V4["claimed"]["M25"]["scores"][m]["touched_pairs"] for m in months)

    # ranges used in several sections
    rng_raw_t = (pt(min(e1raw(r, m) for r in R for m in months)),
                 pt(max(e1raw(r, m) for r in R for m in months)))
    rng_cor_t = (pt(min(e1v(r, m) for r in R for m in months), 2),
                 pt(max(e1v(r, m) for r in R for m in months), 2))
    g520 = G520["fit_raw_deviation_on_bias"]
    us_rec = max(t7[r]["microseconds_per_record"] for r in R)
    fit_s = (t7["M25"]["fit_seconds"], t7["M1"]["fit_seconds"])
    n_inj = sum(c["injections"] for r in R for c in e2["roads"][r]["by_chain_size"].values())
    n_inj9 = sum(I9[r]["offset +40 cm"]["injections"] for r in R)
    TN = {"pat": 1, "npl": 2, "map": 3, "impl": 4, "ref": 5, "cfg": 6, "tests": 7,
          "corr": 8, "refit": 9, "inj": 10, "offset": 11, "xb": 12, "steps": 13,
          "metrics": 14, "effects": 15, "plan": 16}

    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = B.BODY_FONT
    st.font.size = Pt(11)
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Inches(0.9)
    sec.top_margin = sec.bottom_margin = Inches(0.85)
    B.running_heads(sec)
    B.letterhead(doc)
    num = Numbering()

    # ============================================================== 1
    h1(doc, "1. Title of the Invention")
    para(doc,
         "A System and Method for Determining the Length-Measurement Offset of Each Roadside "
         "Inductive-Loop Detector in a Chain and Recalibrating Its Vehicle-Length "
         "Classification, Using the Invariance of Vehicle Length Between Adjacent Detectors.")

    # ============================================================== 2
    h1(doc, "2. Field / Area of Invention")
    para(doc,
         "The present invention relates to roadside vehicle-detection hardware, in "
         "particular inductive-loop detector sites that measure the length of each passing "
         "vehicle and transmit counts of vehicles by length class to a traffic-data "
         "archive. More particularly, it pertains to a system and method that determine, "
         "from those class counts alone, the length-measurement offset in centimetres of "
         "each detector in a chain of such sites and the class bias it causes at every "
         "class boundary, test whether that offset is the detector's fault, and correct the "
         "recorded counts, using the invariance of vehicle length between adjacent "
         "detectors that observe the same vehicles.")
    para(doc,
         "A dual-loop detector site computes the speed of a vehicle from the times at which "
         "it reaches the two loops of a lane and its length from the time it occupies a "
         "loop, less the loop length. The result depends on the stored loop dimensions and "
         "spacing, on the path of vehicles over the loops, on the number of turns of each "
         "loop, on reinforcement near the loops and on the threshold behaviour of the "
         "detector card; GB 2385138 A [P2] names each of these as a cause of systematic "
         "over- or under-reading of length that inspection does not reveal. A site that "
         "reads long moves vehicles from each length class into the next. Its total count "
         "stays right, its fleet mix goes wrong, and nothing in its own output shows it.")
    para(doc,
         "The error is large and common. On carriageway A of the M25 in March 2025, %d "
         "pairs of consecutive detectors have no slip road between them and a median ratio "
         "of total counts of %.4f, so they count the same vehicles; their shares of "
         "vehicles below 520 cm differ by a median of %s points and by %s at the 90th "
         "percentile. Each pair's difference in March correlates with the same pair's "
         "difference in June at r = %.3f, and hardly at all with the ratio of the two "
         "detectors' speeds (r = %.2f). On carriageway A of the M1 the figures are %d "
         "pairs, %s and %s points. The technical problem addressed here is to determine, "
         "from the length-class counts that the detectors already report, by how much each "
         "detector unit of a chain mis-measures vehicle length, whether that is its fault "
         "or a fault that a retune of its sensitivity would not cure, and to correct the "
         "class counts it has recorded and reports, without test vehicles, apparatus at "
         "the site or per-vehicle data."
         % (p25["pairs_march"], p25["volume_ratio_median"], pt(p25["abs_gap520_median"]),
            pt(p25["abs_gap520_p90"]), p25["persistence"]["2025-06"]["corr_gap520_vs_march"],
            p25["corr_gap520_speed_ratio"], p1["pairs_march"], pt(p1["abs_gap520_median"]),
            pt(p1["abs_gap520_p90"])))

    h2(doc, "2.1 Technical Problem Addressed by the Invention")
    bullets(doc, [
        ("Hidden length offset.", "A detector that over- or under-measures every vehicle "
         "length keeps a correct total count and a wrong class mix, and site inspection "
         "does not reveal it [P2]."),
        ("Calibration by attendance.", "Known practice brings probe vehicles, a validation "
         "apparatus, an instrumented vehicle or a laptop to each site, one site at a time "
         "[N5, P2, P3, N3]."),
        ("Archives hold only class counts.", "Methods that work from a detector's own "
         "records need its per-vehicle on-times or signatures and an assumed vehicle length "
         "or a standard length table [N4, P1]; a classifier reports, per interval, only "
         "the number of vehicles in each length class."),
        ("A flag, not a number.", "Comparison with neighbouring stations singles out the "
         "odd station for attendance [N1], but gives no size of the error at any class "
         "boundary and leaves the recorded counts wrong."),
        ("Wrong remedy.", "A class boundary set wrongly in the detector card, or a fault "
         "confined to one lane or one class, is not cured by retuning the detector's "
         "sensitivity, and no known method tells it from a length offset using class "
         "counts."),
        ("Undated changes.", "A retune, a repair or a fault changes a unit's length "
         "measurement on one day; the archive does not show which of its records lie "
         "before the change and which after it."),
    ])

    h2(doc, "2.2 Definitions")
    bullets(doc, [
        ("Detector unit (110, 111):", "a roadside site with a pair of inductive loops in "
         "every lane and a detector card that computes each vehicle's length and counts it "
         "into length classes, reporting class counts, total count and mean speed per "
         "interval."),
        ("Class boundary e:", "a length separating two classes; 520, 660 and 1160 cm in "
         "the embodiment. The first boundary is the most populated one, 520 cm."),
        ("Share below a boundary, S_i(e):", "the fraction of the vehicles counted by "
         "detector i over the used intervals that it reports shorter than e."),
        ("Admitted pair:", "two consecutive detectors with no slip road (120) between "
         "them whose total counts agree within a balance tolerance over free-flowing "
         "intervals; its two detectors measure the same vehicles."),
        ("Chain:", "a sequence of admitted pairs in which consecutive pairs share a "
         "detector."),
        ("Share bias, b_i(e):", "the error in S_i(e), solved along the chain, relative to "
         "the chain's reference: its median detector, or a detector whose length "
         "measurement has been verified."),
        ("Length distribution, F:", "the cumulative distribution of the lengths of the "
         "vehicles of the stream, interpolated through the class shares; f is its slope."),
        ("Length offset, d:", "the distance by which a detector over-measures every "
         "vehicle length relative to the chain's reference, so that b(e) = F(e - d) - "
         "F(e)."),
        ("Cross-boundary test:", "the comparison of a unit's share biases at the other "
         "boundaries with those its offset predicts, to first order b(520) f(e) / f(520)."),
        ("Fault class:", "a length-offset fault, where the unit passes the test and its "
         "offset is a sensitivity correction, or a fault of another kind, where it fails."),
        ("Total-preserving reclassification:", "moving b_i(e) n vehicles across each "
         "boundary in each interval, n being the interval total, which leaves n unchanged."),
        ("Dated change:", "equal and opposite steps on one day in the two pairs of one "
         "detector, attributing a change of its length measurement to that detector and "
         "that day."),
    ])

    # ============================================================== 3
    h1(doc, "3. Prior Patents and Publications from Literature")
    para(doc,
         "Seven patent documents and eleven publications were found closest to the "
         "invention, and each was read in full text before any statement about it was "
         "written. Patents are cited below by number and publications by author, with the "
         "row of Table 1 (P) or Table 2 (N) in brackets.")
    h2(doc, "A. Prior Patent Literature")
    table(doc, ["S.No.", "Patent No. / Publication", "Title", "Comments"],
          [list(r) for r in PATENTS],
          cap="Table 1: Prior Patent Literature",
          widths=[0.45, 1.15, 1.45, 3.55], size=8.5, align_left_cols=(1, 2, 3))
    h2(doc, "B. Non-Patent Literature")
    table(doc, ["S.No.", "Source / Publisher / Citation", "Title", "Comments"],
          [list(r) for r in NPL],
          cap="Table 2: Non-Patent Literature",
          widths=[0.45, 1.45, 1.35, 3.35], size=8.5, align_left_cols=(1, 2, 3))
    para(doc,
         "Patent searches were run on Google Patents and FreePatentsOnline with the terms "
         "vehicle length, effective vehicle length, length class, length bin, population "
         "distribution, reclassification, self-calibration and auto-calibration, combined "
         "with adjacent, neighbouring, upstream and downstream detector and with "
         "calibration, normalisation and correction, in English and Chinese; the "
         "literature search covered loop sensitivity, archived-data correction, "
         "weigh-in-motion calibration and blind sensor calibration.")

    # ============================================================== 4
    h1(doc, "4. Summary and Background of the Invention")
    para(doc,
         "Motorway networks classify traffic by vehicle length at roadside inductive-loop "
         "sites, and the class counts they archive feed pavement, freight and "
         "traffic-management analyses. Each site reads length with an offset peculiar to "
         "its installation, so that consecutive sites counting the same vehicles report "
         "shares of vehicles below 520 cm differing by a median of %s points on one "
         "corridor and %s on another. The invention uses the vehicles passing consecutive "
         "detectors as a moving reference: it solves for each detector's signed class bias "
         "at every boundary along a chain of detector pairs, converts those biases into a "
         "length offset in centimetres through the traffic stream's own length "
         "distribution, tests that this one offset predicts the biases at every other "
         "boundary, and reclassifies the recorded counts with every interval total kept. "
         "A flag for inspection thereby becomes a corrected record and, for each detector "
         "unit, a length offset and a fault class that say what is wrong with the unit "
         "and by how much."
         % (pt(p25["abs_gap520_median"]), pt(p1["abs_gap520_median"])))

    h2(doc, "4.1 Practical Background and Limitations of Existing Solutions")
    para(doc,
         "Finding a length error has meant going to the site. The practice recommended for "
         "length-based classification is to drive probe vehicles of known length through "
         "the detector and adjust it until it agrees (Minge et al. [N5]); alternatives "
         "bring switched validation loops (GB 2385138 A [P2]), an instrumented probe "
         "vehicle (GB 2377027 A [P3]) or a laptop that records per-vehicle event data at "
         "the controller cabinet and compares it with an external length distribution "
         "(Nihan et al. [N3]). Methods that work from the detector's own records need its "
         "per-vehicle on-times or signatures and an assumed vehicle length or a standard "
         "population table (Lee and Coifman [N4], US 6,876,949 B2 [P1]). Archives of "
         "classified counts hold none of these.")
    para(doc,
         "The art that does compare neighbouring detectors stops at a flag or at total "
         "volumes. Azin and Yang [N1] single out the station whose length statistic differs "
         "from its neighbours and mark it for recalibration; Dailey and Wall [N2] correct "
         "archived volumes station by station from a reference station; Bychkovskiy et al. "
         "[N9] and Balzano and Nowak [N10] calibrate generic sensors against one another; "
         "US 6,876,949 B2 [P1] normalises detectors with re-identified vehicles. None "
         "determines how far each detector's classification is wrong at each boundary, none "
         "recovers a length offset from class counts, none tests whether the error is a "
         "length offset at all, and none corrects the class record.")
    figure(doc, "fig01.png",
           "Figure 1: Roadside detector sites on a carriageway and the length-invariance "
           "principle",
           "Six detector sites (110), each a loop pair in every lane with a roadside "
           "detector card, report class counts over the link (130). Consecutive sites "
           "with no slip road between them form admitted pairs and two chains; a site "
           "reading every length 45 cm long (d5) moves vehicles across each boundary in "
           "proportion to the length density there, and disagrees with both neighbours by "
           "the same amount. The length distributions are schematic.")
    figure(doc, "fig02.png",
           "Figure 2: Architecture of the roadside hardware and the data-processing system",
           "Class counts cross from the roadside hardware (100) to the data-processing "
           "system (900), pass through pair admission (300, 310), chain solution (320) and "
           "the offset estimator and fault classifier (350, 360), and are reclassified (400); "
           "outputs (810-830) are the corrected archive, supply to traffic management and "
           "the maintenance record, and the sensitivity correction (840) returns to the "
           "detector unit.")

    h2(doc, "4.2 Distinctions Over the Closest Reported Systems")
    para(doc,
         "Azin and Yang [N1] compare a length statistic of each station with its "
         "neighbours and flag the odd station for recalibration. The comparison and the "
         "flag are conceded. The present invention solves, along a chain of pairs that see "
         "the same vehicles, a signed bias for every detector at every class boundary, "
         "turns it into a length offset in centimetres, tests whether the unit's fault is a "
         "length offset, and corrects the recorded class counts, so that the output is a "
         "corrected record and a sensitivity correction rather than a request for "
         "attendance.", lead="Compared with Azin and Yang [N1].")
    para(doc,
         "US 6,876,949 B2 [P1] calibrates each detector's inductive length so that its "
         "population distribution of per-vehicle lengths matches a standard table, or "
         "normalises detectors with re-identified vehicles. It needs per-vehicle "
         "signatures and a standard or a re-identification, which a class-count archive "
         "does not hold and which a local fleet mix would defeat. The present invention "
         "takes the length distribution from the class shares of the stream itself, fixes "
         "the offset relative to the chain's own reference, and uses the distribution a "
         "second time, to predict the bias at every other boundary and so tell a length "
         "offset from another fault, which a matched distribution cannot do.",
         lead="Relative to US 6,876,949 B2 [P1].")
    para(doc,
         "Dailey and Wall [N2] pair stations with no ramp between them, confirm them by "
         "count balance and correct archived total volumes from a reference station; Lu et "
         "al. [N8] test upstream and downstream flows. Total counts are what these methods "
         "correct. In the present invention the totals only admit a pair and are kept "
         "unchanged by the correction; what is compared, solved and corrected is the share "
         "of vehicles below each length-class boundary.",
         lead="Against Dailey and Wall [N2] and Lu et al. [N8].")
    para(doc,
         "Nihan et al. [N3], Lee and Coifman [N4], Minge et al. [N5], GB 2385138 A [P2] and "
         "GB 2377027 A [P3] calibrate one site at a time against vehicles, apparatus or an "
         "assumed length brought to it. The weigh-in-motion guide [N11] watches one site's "
         "weight calibration through a statistic of its own traffic, and US 7,684,946 B2 "
         "[P4] compares a scale with a second instrument on the same vehicles. The present "
         "invention needs no second instrument, no attendance and no assumed fleet "
         "property: each detector's neighbours, seeing the same vehicles, are its "
         "reference.",
         lead="In contrast to site-by-site calibration [N3, N4, N5, P2, P3, N11, P4].")

    h2(doc, "4.3 Inventive Combination / Non-Obvious Technical Arrangement")
    para(doc,
         "No single element carries the inventive step, and the following are conceded. "
         "Comparing a length statistic between adjacent detectors and flagging the odd one "
         "is Azin and Yang [N1]. Pairing stations with no ramp between them, confirming "
         "them by count balance, working outward from a reference station and correcting "
         "the archive is Dailey and Wall [N2], for total volumes. Calibrating sensors "
         "against their neighbours is Bychkovskiy et al. [N9] and Balzano and Nowak [N10]. "
         "Watching a site's calibration through a statistic of its own traffic is "
         "weigh-in-motion practice [N11]. Calibrating a detector's length measurement so "
         "that its vehicle population distribution matches a standard one, or normalising "
         "detectors with re-identified vehicles, is US 6,876,949 B2 [P1], and retuning a "
         "loop from the length distribution of per-vehicle data is Nihan et al. [N3]. The "
         "features claimed in combination address the one technical problem stated in "
         "Section 2, and the invention resides in their co-operation:")
    bullets(doc, [
        ("(a)", "the invariant: the share of vehicles below each class boundary of the "
                "detectors' own classification, compared between two detectors over "
                "matched intervals, on the class counts already archived and nothing else;"),
        ("(b)", "admission of a pair only where no slip road lies between the detectors "
                "and their totals agree, the physical condition for the two detectors to "
                "measure the same vehicles;"),
        ("(c)", "determination along a chain of admitted pairs of a signed bias for every "
                "detector at every boundary, with the chain's constant fixed by its median "
                "detector or by a verified site;"),
        ("(d)", "conversion of each detector's bias into a length offset in centimetres, "
                "relative to the chain's reference, through the cumulative length "
                "distribution of the traffic stream interpolated from the class shares "
                "themselves, so that no standard distribution, per-vehicle record or "
                "re-identification is needed;"),
        ("(e)", "a cross-boundary test: the offset fixed at one boundary predicts, with no "
                "further fitted parameter, the bias at every other boundary; a unit whose "
                "biases follow the prediction has a length offset, which is written to it "
                "as a sensitivity correction, and a unit whose biases do not has a fault of "
                "another kind, which a retune of its sensitivity would not cure;"),
        ("(f)", "correction by reclassifying, in every interval, a number of vehicles "
                "equal to the bias times the interval total, which keeps every total; and"),
        ("(g)", "attribution of a change in length measurement to one detector and one day "
                "from equal and opposite steps in its two pairs, with refitting each "
                "period and from the day of the change."),
    ])
    para(doc,
         "The elements are interdependent, and the interdependence is physical. Elements "
         "(d) and (e) are obtainable only because (a) compares class shares at the "
         "boundaries where a length error acts, (b) guarantees that both detectors sample "
         "one length distribution, and (c) yields each detector's bias at several "
         "boundaries at once, so that the chain supplies both the reference and the length "
         "distribution. The test of (e) rests on a structure found in the data and taught "
         "by no reference: across units, the biases at 520 cm and 660 cm move together "
         "(r = %.2f to %.2f in every corridor-month); with each boundary scaled to unit "
         "spread, one offset per unit explains %.0f %% to %.0f %% of the bias variance, "
         "against at most %.0f %% when the link between a unit's boundaries is broken by "
         "shuffling; and the bias at 660 cm predicted from each unit's offset correlates "
         "with the solved bias at r = %.2f to %.2f. The test does not depend on the chain's "
         "unknown constant: moving every detector of a chain by a common offset of +30 or "
         "-30 cm changes the class in %d of %d unit tests. Combining Azin and Yang [N1], "
         "Dailey and Wall [N2] and US 6,876,949 B2 [P1] does not reach (d) and (e): the "
         "last matches a whole distribution of per-vehicle lengths to a standard table, "
         "which a class-count archive does not hold, and has no means of telling a length "
         "offset from another fault."
         % (min(r660), max(r660), 100 * min(std_real), 100 * max(std_real),
            100 * max(std_p99), min(c660), max(c660), gauge_changed, gauge_units))
    para(doc,
         "Each element earns its place in measurement (Section 8.10). Without the balance "
         "test of (b) the corrected disagreement on the M1 is larger in every month, and "
         "without its free-flow condition it is larger in every month on both corridors; "
         "the per-boundary correction of (f) leaves less disagreement than a correction "
         "proportional to the share and, at 660 cm, less than one computed from the offset "
         "alone; (e) classes an injected 660 cm boundary fault as not an offset in %.0f %% "
         "to %.0f %% of cases while keeping an injected 40 cm length offset in %.0f %% to "
         "%.0f %%; and against independent manual classified counts each detector's fitted "
         "bias at 520 cm accounts for %.2f +- %.2f of its real departure per point."
         % (100 * min(det660), 100 * max(det660), 100 * min(keep40), 100 * max(keep40),
            g520["beta"], g520["beta_se"]))

    h2(doc, "4.4 Mapping of Key Features Against Closest Prior Art")
    Y, P, N = "Yes", "Partial", "No"
    table(doc, ["Feature", "N1", "N2", "N3", "N4", "P1", "N11"], [
        ["F1 Works on the class counts already archived", P, N, N, N, N, N],
        ["F2 Same-vehicle pair: no slip road between, totals agree", P, Y, N, N, P, N],
        ["F3 Neighbouring detectors as the reference", Y, Y, N, N, Y, N],
        ["F4 Signed bias per detector and per class boundary, solved along a chain",
         N, P, N, N, P, N],
        ["F5 Length offset in cm from the stream's own length distribution", N, N, P, P,
         P, N],
        ["F6 Cross-boundary test of whether the fault is a length offset", N, N, N, N, N,
         N],
        ["F7 Total-preserving reclassification of recorded class counts", N, N, N, N, N,
         N],
        ["F8 Change dated to one detector and one day", N, N, P, P, N, P],
        ["F9 No attendance, test vehicle or per-vehicle data", Y, Y, N, N, N, P],
    ], cap="Table 3: Feature Mapping (Yes = disclosed, Partial = related feature, "
           "No = not identified)",
        widths=[3.0, 0.55, 0.55, 0.55, 0.55, 0.55, 0.55], size=8.8, align_left_cols=(0,))
    para(doc,
         "Table 3 shows that F1, F2, F3 and F9 exist in related forms and that F6 and F7, "
         "the test of the kind of fault and the correction of the class record, were not "
         "found in any mapped reference. F4 is Partial for Dailey and Wall [N2], who "
         "correct each station's total volume from a reference station, and for "
         "US 6,876,949 B2 [P1], which solves normalisation coefficients over detectors "
         "from re-identified vehicles; neither works per class boundary. F5 is Partial "
         "where a length error is found against an external distribution [N3], an assumed "
         "vehicle length [N4] or a standard table [P1], never from the stream's own class "
         "shares. F8 is Partial where a site's measurement is tracked over time, before "
         "and after a known tune-up [N3], by its daily median [N4] or by its drift [N11], "
         "never by equal and opposite steps in two pairs.")

    h2(doc, "4.5 Technical Effects Achieved")
    bullets(doc, [
        ("Adjacent detectors agree.", "In months the fit did not see, the mean difference "
         "below 520 cm between adjacent detectors falls from %s-%s to %s-%s points over "
         "the pairs touching a corrected detector, and by %.0f %% to %.0f %% over all "
         "admitted pairs." % (rng_raw_t + rng_cor_t + (100 * min(red_a), 100 * max(red_a)))),
        ("Each unit's length error in centimetres.", "Median |d| %.0f to %.0f cm, 90th "
         "percentile %.0f to %.0f cm; %.0f %% to %.0f %% of fitted units read beyond the "
         "published length requirement of +-1.5 ft [N11] relative to their chain's "
         "reference." % (min(off50), max(off50), min(off90), max(off90),
                         100 * min(tol_share), 100 * max(tol_share))),
        ("The kind of fault told apart.", "An injected 660 cm boundary fault is classed as "
         "not an offset in %.0f %% to %.0f %% of cases; an injected 40 cm offset keeps its "
         "class in %.0f %% to %.0f %%." % (100 * min(det660), 100 * max(det660),
                                          100 * min(keep40), 100 * max(keep40))),
        ("The faulty unit localised.", "An injected fault shows its largest change at the "
         "injected site in %.0f %% to %.0f %% of injections by chain length."
         % (100 * min(locsize), 100 * max(locsize))),
        ("Confirmed by independent counts.", "At %d manual classified counts the fitted "
         "bias at 520 cm predicts the recorded departure with slope %.2f +- %.2f, and the "
         "spread of departures falls from %.1f to %.1f points after correction."
         % (G["matched"], g520["beta"], g520["beta_se"], G520["mad_raw_pts"],
            G520["mad_corrected_pts"])),
        ("Changes dated.", "%d of %d sites show equal and opposite steps on one day; the "
         "same rule finds %d in %d mismatched pairings." % (len(ev_all), n_tested,
                                                            n_ctrl_ev, n_ctrl)),
        ("Live operation.", "Each fifteen-minute record is corrected in %.1f microseconds "
         "and a corridor-month is fitted in %.1f to %.1f s." % (us_rec, min(fit_s),
                                                               max(fit_s))),
    ])

    h2(doc, "4.6 Industrial / Practical Applicability")
    para(doc,
         "The system runs on a processor attached to an existing count archive and adds "
         "no equipment at the roadside. It applies wherever consecutive length-classifying "
         "detectors observe the same traffic stream: motorway and expressway detector "
         "chains, classified-count stations on national highways, and radar, video, "
         "piezoelectric or axle-based classifiers that report class counts. Its corrected "
         "counts serve the uses of classified counts, the heavy-vehicle and fleet-mix "
         "inputs of pavement design, freight counts and traffic-management records, and "
         "its maintenance record tells maintenance staff which detector unit to retune, "
         "by how many centimetres, and which unit has a fault a retune would not cure.")

    h2(doc, "4.7 Patentability Summary")
    para(doc,
         "None of the seven patent documents and eleven publications reviewed describes "
         "a system in which the length-class shares of detectors that see the same "
         "vehicles are solved along a chain into a signed bias per detector and per class "
         "boundary, converted into a length offset through the stream's own length "
         "distribution, tested against the biases that offset predicts at the other "
         "boundaries, and used to reclassify the recorded counts with every total kept.",
         lead="Novelty.")
    para(doc,
         "The objective technical problem is single: to find, from class counts alone, how "
         "far each detector unit mis-measures length, whether that is its fault, and to "
         "correct its record. The art offers two responses: calibrate one site at a time "
         "against a known reference brought to it [N3, N4, N5, P2, P3], or flag the "
         "station that differs from its neighbours [N1]. Neither yields an offset per unit "
         "from class counts, and neither can tell a length offset from another fault. The "
         "present invention does both, and the test that does the second is predicted "
         "from the counts' own length distribution with no fitted parameter, rather than "
         "chosen as the direction of a residual; its effect is measured in Table 12.",
         lead="Inventive step.")
    para(doc,
         "The system is implemented and runs end to end on four months of field records "
         "of two motorway corridors (Section 8).", lead="Industrial applicability.")
    para(doc,
         "Protection is sought for a method and system that receive the output of "
         "roadside measuring hardware, rely on a physical invariant of the vehicles that "
         "hardware measures, determine the length offset of each detector unit in "
         "centimetres, test against the physics of a length offset whether that is the "
         "unit's fault, and change both the record the detectors produce and the "
         "sensitivity setting of the units.", lead="Technical character.")

    # ============================================================== 5
    h1(doc, "5. Objective(s) of the Invention")
    para(doc,
         "The primary objective is a roadside detector system in which the length-class "
         "counts that a chain of detectors already reports are used to find, for every "
         "detector unit, how far it mis-measures vehicle length and whether that is its "
         "fault, and to correct its recorded and incoming class counts, without attending "
         "the site.")
    bullets(doc, [
        ("(a)", "To admit pairs of consecutive detectors that measure the same vehicles: "
                "no slip road (120) between them and total counts that agree."),
        ("(b)", "To solve along a chain of admitted pairs a signed share bias for every "
                "detector at every class boundary, relative to the chain's reference."),
        ("(c)", "To convert each bias into a length offset in centimetres through the "
                "stream's own length distribution, interpolated from the class shares."),
        ("(d)", "To test, per unit, whether the offset predicts its biases at the other "
                "boundaries, and so classify its fault."),
        ("(e)", "To reclassify each detector's recorded and incoming counts with every "
                "interval total kept, writing the result to the count archive (200) and "
                "supplying it to traffic management (820)."),
        ("(f)", "To write each unit's offset as a sensitivity correction (840) of its "
                "detector card, date any change of its measurement, and refit from it."),
        ("(g)", "To validate the system out of sample, against injected faults, on pairs "
                "never fitted and against independent manual classified counts."),
    ])

    # ============================================================== 6
    h1(doc, "6. Working Principle of the Invention (in Brief)")
    num(doc,
        "A vehicle's length does not change between two roadside detector sites a few "
        "hundred metres apart. Where no slip road joins or leaves the carriageway between "
        "them, both sites observe the same vehicles, and the share of those vehicles "
        "shorter than any given length is the same at both. A persistent difference "
        "between the shares the two sites report can only arise in their measurement. "
        "Each fitting period the system (900) runs the sequence below.")
    p_ = para(doc,
              "class counts of detector sites (110) → pair admission (300) → share "
              "differences (310) → chain solution (320) → length offset (350) → "
              "cross-boundary test (360) → reclassification, totals kept (400) → corrected "
              "archive (810), traffic management (820), maintenance record (830), "
              "sensitivity correction (840)", size=10.5, italic=True)
    p_.alignment = WD_ALIGN_PARAGRAPH.CENTER
    num(doc,
        "Each site measures length with a pair of inductive loops in every lane and a "
        "detector card that times each vehicle's passage. An installation-specific error "
        "adds an offset d to every length the card computes. At a class boundary e, a site "
        "that reads long counts below e only the vehicles shorter than e - d, so its share "
        "below e falls by the share of vehicles whose length lies between e - d and e; the "
        "vehicle count itself is untouched.")
    num(doc,
        "Two consecutive sites are paired only if no slip road lies between them and, "
        "over the intervals both reported in free-flowing traffic, their total counts "
        "agree. Agreement of totals confirms that the pair sees the same vehicles; total "
        "counts are used for this admission and for nothing else.")
    num(doc,
        "For an admitted pair (i, i+1) and a boundary e, S_i(e) is the share of vehicles "
        "that detector i reports below e and b_i(e) the error in that share. The true "
        "shares being equal, the measured difference is the difference of the two biases: "
        "g_i(e) = S_i(e) - S_(i+1)(e) = b_i(e) - b_(i+1)(e).")
    num(doc,
        "Admitted pairs that share a detector form a chain. A chain of k+1 sites gives k "
        "such relations per boundary in k+1 unknowns, so the biases are fixed up to one "
        "constant, set by taking the median site of the chain as correct, or a site whose "
        "length measurement has been verified. Each interior site belongs to two pairs, "
        "and its own error appears in both with opposite signs, which localises it.")
    num(doc,
        "Because the vehicles are the same along the chain, one cumulative length "
        "distribution F describes them all, and it is known at the class boundaries from "
        "the shares the chain reports. A detector with offset d has bias F(e - d) - F(e) "
        "at every boundary e: one number per detector, scaled at each boundary by the "
        "length density there. Solving this relation at the most populated boundary gives "
        "d in centimetres, the unit's length offset relative to the chain's reference: "
        "differences between units are physical, and the common level becomes absolute "
        "when the reference is a verified site.")
    num(doc,
        "At the other boundaries the same d predicts the bias, with no further parameter: "
        "to first order it is the bias at the first boundary times the ratio of the length "
        "densities at the two boundaries, which needs the distribution only at the class "
        "boundaries and not its shape between them, which class counts do not fix. "
        "Where the detector's biases depart from that prediction by more than a limit, its "
        "fault is not a length offset, for instance a class boundary set wrongly in the "
        "detector card or a fault confined to one lane or one class; its sensitivity is "
        "not the thing to retune, and the unit is classified accordingly. A common offset "
        "of the whole chain moves every unit along the same predicted pattern, so the "
        "test does not depend on the chain's reference.")
    num(doc,
        "A loop-spacing error scales the speed the detector computes and, through it, the "
        "length; a detection-threshold (sensitivity) error changes the time a vehicle "
        "occupies the loop and so the length alone. The speed bias of each detector, "
        "solved along the chain from the ratio of mean speeds in each pair in the same way "
        "as the share biases, separates the two causes of an offset.")
    num(doc,
        "The recorded counts are corrected by reclassifying, in each interval, a number of "
        "vehicles across boundary e equal to b_i(e) times the interval total; the "
        "corrected share becomes S_i(e) - b_i(e) and the total count and mean speed are "
        "not changed.")
    num(doc,
        "A site whose length measurement changes on a given day, through a retune, a "
        "repair or a fault, produces on that day a step of one sign in the pair where it "
        "is downstream and an equal step of the other sign in the pair where it is "
        "upstream, at every boundary; a change in traffic moves both sites of a pair "
        "together and leaves their difference unchanged. Biases and offsets are refitted "
        "each period, and a site with a dated change is refitted from the day of the "
        "change, so the correction follows the physical state of each detector unit.")
    num(doc,
        "The outputs are the corrected class counts, written back to the count archive "
        "and supplied to traffic management in place of the received counts; a "
        "maintenance record holding, for each detector unit, its length offset, fault "
        "class, the part of the offset due to loop spacing, its bias at every boundary "
        "and any dated change; and, for each unit with a length-offset fault, a "
        "sensitivity correction sent to its detector card.")

    h2(doc, "6.1 What Is Being Invented vs One Example Implementation")
    table(doc, ["Functional element", "Example implementation"], [
        ["Length-classifying detector reporting class counts",
         "Dual inductive-loop sites of the English motorway network; four classes at 520, "
         "660 and 1160 cm; fifteen-minute records with total count and mean speed"],
        ["Order of detectors and positions of slip roads",
         "Read from the site identifiers: marker number = chainage, marker letter = "
         "mainline or slip road"],
        ["Condition that a pair sees the same vehicles",
         "No slip-road site between, at most 3 km apart; totals within %.0f %% over "
         "intervals above %.0f mph with at least %d vehicles; at least %d intervals a month"
         % (100 * pairs.BALANCE_TOL, pairs.MIN_SPEED_MPH, pairs.MIN_VOLUME,
            pairs.MIN_INTERVALS)],
        ["Reference of each chain", "Median of the chain's biases; or a site verified by "
                                    "probe vehicles or a manual classified count"],
        ["Length distribution of the stream",
         "Monotone piecewise-cubic curve through the median class shares, anchored at 0 "
         "below 250 cm and 1 at %d cm; linear or other anchors change offsets by x%.2f to "
         "x%.2f" % (e8_CDF_TOP, min(iv_ratio), max(iv_ratio))],
        ["Length offset", "Bisection of F(520 - d) - F(520) = b(520)"],
        ["Cross-boundary test", "First-order prediction b(520) f(e) / f(520); fails above "
                                "%.1f of max(|b(520)|, %s points); or the principal "
                                "direction of the units' bias vectors"
         % (e9["limit"], pt(e9["floor"], 0))],
        ["Correction of the record", "round((s - b) n) below each boundary, bounded to "
                                     "[0, n], non-decreasing; totals and mean speed kept"],
        ["Dated change", "Largest standardised step, at least %d days each side; two steps "
                         "within %d day, each over %s points, opposite, sum under %.0f %% "
                         "of the larger" % (e3_steps.MIN_SIDE, e3_steps.STEP_DAY_TOL,
                                            pt(e3_steps.MIN_STEP, 0),
                                            100 * e3_steps.CANCEL_TOL)],
        ["Fitting period", "Calendar month, applied to the next; a week or rolling window"],
        ["Data-processing system (900)", "Desktop processor attached to the archive"],
    ], cap="Table 4: Claimed Functional Combination vs Example Implementation",
        widths=[2.1, 4.4], size=9, align_left_cols=(0, 1))
    num(doc,
        "Table 4 separates the functional combination for which protection is sought from "
        "the particular values and parts of this embodiment, any of which can be replaced "
        "by the equivalents listed in Section 7.11.")

    h2(doc, "6.2 Technical Contribution")
    bullets(doc, [
        ("The vehicles are the reference.", "Adjacent detectors that see the same "
         "vehicles, whose lengths do not change, replace probe vehicles, test apparatus "
         "and standard length tables."),
        ("Class counts suffice.", "The share below each boundary, solved along a chain, "
         "gives each detector's error per boundary without any per-vehicle record."),
        ("The stream supplies its own length distribution.", "The class shares fix the "
         "distribution at the boundaries, which turns a share bias into a length offset in "
         "centimetres."),
        ("The physics of a length offset is tested per unit.", "The offset predicts the "
         "bias at every other boundary with no fitted parameter, which separates a "
         "retunable offset from a fault of another kind."),
        ("The correction keeps the totals.", "The class record is corrected without "
         "changing any total count or mean speed, so admission and every total-based use "
         "of the counts are undisturbed."),
        ("Changes are dated to one unit and one day.", "Equal and opposite steps in a "
         "detector's two pairs separate a change of the detector from a change of the "
         "traffic."),
    ])

    h2(doc, "6.3 Brief Description of the Drawings")
    num(doc,
        "Figure 1 shows detector sites on a carriageway and the length-invariance "
        "principle, and Figure 2 the architecture of the roadside hardware and the "
        "data-processing system with reference numerals 100 to 900. Figure 3 gives the "
        "sequence of operations from the loop pairs to the maintenance record. Figures 4 "
        "to 10 present the validation: disagreement between adjacent detectors (4), "
        "correction of later months (5), injected faults (6), the length offset and the "
        "cross-boundary test (7), dated changes (8), admission conditions and the "
        "attendance mark (9), and comparison with manual classified counts (10). Figure "
        "11 is the claim tree.")

    # ============================================================== 7
    h1(doc, "7. Description of the Invention in Detail")
    h2(doc, "7.1 System Scope and Embodiment")
    num(doc,
        "The system (900) and the roadside hardware it serves are shown in Figure 2, the "
        "sequence of operations in Figure 3, and the reference numerals in Table 5. The "
        "operating values below are those of the embodiment and of every measurement in "
        "Section 8.")
    table(doc, ["Ref", "Element", "Ref", "Element"], [
        ["100", "roadside detector hardware", "340", "refit scheduler"],
        ["110", "detector site, loop pair in each lane", "350", "offset estimator"],
        ["111", "detector card", "360", "fault classifier"],
        ["112", "neighbouring detector sites", "400", "reclassification unit"],
        ["120", "slip road", "405", "interval total kept"],
        ["130", "communication link", "810", "corrected dataset"],
        ["200", "count archive", "820", "traffic-management process"],
        ["300", "pair-admission unit", "830", "maintenance record"],
        ["310", "share-difference unit", "840", "sensitivity correction of detector unit"],
        ["320", "chain solver", "900", "data-processing system"],
        ["330", "step-event detector", "", ""],
    ], cap="Table 5: Reference Numerals Used in the Description and the Figures",
        widths=[0.55, 2.45, 0.55, 2.45], size=9.5, align_left_cols=(1, 3))

    h2(doc, "7.2 Detector Sites, Communication Link and Count Archive")
    num(doc,
        "Detector sites (110). Each site has a pair of inductive loops in every lane and a "
        "detector card (111) in a roadside cabinet that computes each vehicle's length and "
        "counts it into one of four classes separated by boundaries at 520, 660 and "
        "1160 cm. Every fifteen minutes the card reports the four class counts, the total "
        "count and the mean speed.")
    num(doc,
        "Communication link (130) and count archive (200). The records travel over the "
        "link to the archive, which also holds the site register: the order of the sites "
        "along the carriageway and the positions of the slip roads (120). In the "
        "embodiment both are read from the site identifiers, whose marker number is the "
        "chainage and whose marker letter distinguishes mainline from slip-road sites.")

    h2(doc, "7.3 Pair Admission and Share Differences")
    num(doc,
        "Pair-admission unit (300). Consecutive mainline sites no more than 3 km apart with "
        "no slip-road site between them are candidates. An interval is used if both sites "
        "reported it, both report a mean speed above %.0f mph, and the upstream site "
        "counted at least %d vehicles; the speed condition keeps to free flow, where "
        "dual-loop length measurement is valid (Wu and Coifman [N7]). A candidate is "
        "admitted for a fitting period, one calendar month in the embodiment, if it has at "
        "least %d such intervals and its two totals agree within %.0f %%. Figure 4(c) "
        "shows why: the largest difference among the M25 candidates, %s points, belongs "
        "to a pair whose totals differ by %.0f %%, where the two sites do not see the same "
        "vehicles."
        % (pairs.MIN_SPEED_MPH, pairs.MIN_VOLUME, pairs.MIN_INTERVALS,
           100 * pairs.BALANCE_TOL, pt(p25["abs_gap520_max"]),
           100 * (1 - p25["abs_gap520_max_pair_volume_ratio"])))
    num(doc,
        "Share-difference unit (310). For each admitted pair and each boundary, the share "
        "below the boundary at each site is the sum of its counts below the boundary over "
        "the used intervals divided by the sum of all its counts, and g is the upstream "
        "share less the downstream share, over the whole period and day by day over days "
        "with at least eight used intervals. Figure 4(a) gives the spread of g at 520 cm "
        "across the candidate pairs of March, and Figure 4(b) its persistence: r = %.3f "
        "between March and June on the M25 and %.3f on the M1."
        % (p25["persistence"]["2025-06"]["corr_gap520_vs_march"],
           p1["persistence"]["2025-06"]["corr_gap520_vs_march"]))

    h2(doc, "7.4 Chain Solution and Its Reference")
    num(doc,
        "Chain solver (320). Consecutive admitted pairs sharing a site form a chain; a pair "
        "failing admission breaks it. For each chain of at least two pairs and each "
        "boundary, relative biases are accumulated from the first site, b_(i+1) = b_i - "
        "g_i, and the chain median is subtracted. The same solution is formed for each day "
        "on which every pair of the chain has data, and each site's standard error is the "
        "standard deviation of its daily values divided by the square root of their "
        "number: %s points at 520 cm (median, M25) against a median bias magnitude of %s "
        "points (Figure 4(d))."
        % (pt(se25, 2), pt(spread["M25"]["q50"])))
    num(doc,
        "The median reference has a measured limit. With a known fault injected into one "
        "real detector's record, the injected site shows the largest change on its chain "
        "in %.0f %% to %.0f %% of injections by chain length (Figure 6, Table 10); the "
        "median error in the fitted size is close to zero, but the 90th percentile is "
        "%.0f %% to %.0f %% of the fault, because a fault near the middle of a short chain "
        "moves the median itself. Agreement between neighbours does not depend on the "
        "chain's constant; absolute size on short chains does, and a verified site removes "
        "the limit."
        % (100 * min(locsize), 100 * max(locsize), 100 * min(p90rel), 100 * max(p90rel)))
    figure(doc, "fig03.png",
           "Figure 3: Sequence of operations from the loop pairs to the maintenance record",
           "Acquisition (S1, S2), admission of pairs (S3 and two decisions), chain solution "
           "(S4, S5), correction (S6, S7) and output (S8, S9): each unit's biases become one "
           "length offset and a fault class; the refit engine (330, 340) dates changes and "
           "restarts S5, and S4 to S6 repeat each fitting period.", width=4.9)

    h2(doc, "7.5 Length Offset from the Stream's Own Length Distribution")
    num(doc,
        "Offset estimator (350). The cumulative length distribution of the stream is "
        "interpolated through the median class shares of the fitted detectors, with a "
        "monotone piecewise-cubic (PCHIP) curve anchored at zero below 250 cm and at one "
        "at %d cm, above the longest permitted vehicle. Each unit's offset d solves "
        "F(520 - d) - F(520) = b(520) by bisection; d is relative to the chain's "
        "reference. Over all fitted units with a bias at 520 cm above 2 points, the "
        "median |d| is %.0f to %.0f cm and the 90th percentile %.0f to %.0f cm across the "
        "eight corridor-months (Table 11, Figure 7(c)). Against the published requirement "
        "for vehicle-length measurement of +-1.5 ft, %.0f cm [N11], %.0f %% to %.0f %% of "
        "all fitted units read beyond it relative to their chain's reference in March and "
        "June. The centimetre value depends on the interpolation: moving the lower anchor "
        "to 200 or 300 cm, the upper anchor to 2000 or 3500 cm, or interpolating linearly "
        "changes the median offset by a factor of %.2f to %.2f, and the fault class of "
        "%.0f %% to %.0f %% of units stays the same."
        % (e8_CDF_TOP, min(off50), max(off50), min(off90), max(off90), e12["tolerance_cm"],
           100 * min(tol_share), 100 * max(tol_share), min(iv_ratio), max(iv_ratio),
           100 * min(iv_agree), 100 * max(iv_agree)))
    num(doc,
        "The single-offset structure is measured, not assumed. Taking all fitted units of "
        "a corridor-month as rows and the three boundaries as columns, one component "
        "explains %.1f %% to %.1f %% of the variance of the bias matrix; since the spread "
        "at 520 cm dominates, this share is still %.1f %% to %.1f %% when the 660 cm and "
        "1160 cm columns are shuffled across units, and is not itself the evidence. With "
        "each boundary scaled to unit spread, one component explains %.0f %% to %.0f %% "
        "against %.0f %% to %.0f %% shuffled (99th percentile at most %.0f %%). The link "
        "is strong at 660 cm, r = %.2f to %.2f with the bias at 520 cm across units, and "
        "weak at 1160 cm, r = %.2f to %.2f, where few vehicles lie near the boundary. Its "
        "ratio between the 660 cm and 520 cm boundaries is %.2f-%.2f "
        "on the M25 and %.2f-%.2f on the M1, against %.2f-%.2f and %.2f-%.2f predicted "
        "from the length density of the interpolated distribution (Figure 7(a), (b)); the "
        "prediction is higher than the fit by %.2f to %.2f in every corridor-month, a "
        "consistent gap that the coarse interpolation between class boundaries would "
        "produce. The prediction at 1160 cm is looser, since the interpolation has no "
        "class boundary between 1160 cm and its upper anchor."
        % (100 * min(rank1), 100 * max(rank1), 100 * min(raw_null), 100 * max(raw_null),
           100 * min(std_real), 100 * max(std_real), 100 * min(std_null),
           100 * max(std_null), 100 * max(std_p99), min(r660), max(r660), min(r1160),
           max(r1160), min(load660["M25"]), max(load660["M25"]),
           min(load660["M1"]), max(load660["M1"]), min(pred660["M25"]), max(pred660["M25"]),
           min(pred660["M1"]), max(pred660["M1"]),
           min(pr - fi for r in R for pr, fi in zip(pred660[r], load660[r])),
           max(pr - fi for r in R for pr, fi in zip(pred660[r], load660[r]))))

    h2(doc, "7.6 Cross-Boundary Test and Fault Class")
    num(doc,
        "Fault classifier (360). For each unit the offset d predicts its biases at 660 cm "
        "and 1160 cm, to first order, as b(520) f(e) / f(520), f being the density of the "
        "interpolated distribution at the boundary. Where the larger departure of the "
        "solved biases from the prediction exceeds %.1f of the larger of the unit's bias "
        "at 520 cm and %s points, the unit is classified as having a fault other than a "
        "length offset, and its offset is not written as a sensitivity correction. Nothing "
        "in the test is fitted to the biases: across the eight corridor-months the "
        "predicted bias at 660 cm correlates with the solved bias at r = %.2f to %.2f and "
        "explains %.0f %% to %.0f %% of its variance (Figure 7(d), Table 12). In March the "
        "test fails %d of %d M25 units and %d of %d M1 units. The class is a property of "
        "the unit: of the units failing in March and fitted again in June, %d of %d fail "
        "again on the M25 and %d of %d on the M1. Failing units are not those with large "
        "offsets, which an error of the interpolated distribution would pick out: their "
        "median offset is %.0f to %.0f cm against %.0f to %.0f cm for passing units."
        % (e9["limit"], pt(e9["floor"], 0), min(c660), max(c660), 100 * min(r2_660),
           100 * max(r2_660), F9["M25"]["2025-03"]["fail_all"],
           F9["M25"]["2025-03"]["detectors"], F9["M1"]["2025-03"]["fail_all"],
           F9["M1"]["2025-03"]["detectors"], P9["M25"]["fail_both"],
           P9["M25"]["fail_march"], P9["M1"]["fail_both"], P9["M1"]["fail_march"],
           min(off_fail), max(off_fail), min(off_pass), max(off_pass)))
    num(doc,
        "The test was checked by injecting faults into the records of real detector units "
        "of April chains of four or more and re-solving the chain. A fault confined to the "
        "660 cm boundary, as a 660 cm threshold set 100 cm wrong would cause, turns a "
        "passing unit into a failing one in %.0f %% to %.0f %% of injections; one confined "
        "to the 1160 cm boundary, where few vehicles lie near the boundary, in %.0f %% to "
        "%.0f %%. A length offset of 40 cm leaves a passing unit passing in %.0f %% to "
        "%.0f %%. Moving every unit of a chain by a common 30 cm offset changes the class "
        "in %d of %d unit tests, both signs and corridors (Table 12)."
        % (100 * min(det660), 100 * max(det660), 100 * min(det1160), 100 * max(det1160),
           100 * min(keep40), 100 * max(keep40), gauge_changed, gauge_units))

    h2(doc, "7.7 Loop Spacing Against Detection Threshold")
    num(doc,
        "Speed attribution. Each unit's speed bias is solved along the chain from the log "
        "ratio of the mean speeds of each pair, with the median gauge. A spacing error "
        "large enough to shift speed by a factor k shifts length at 520 cm by about "
        "(k - 1) x 720 cm with 2 m loops. The median offset such speed biases could "
        "explain is %.0f to %.0f cm, against median offsets of %.0f to %.0f cm, and %.0f %% "
        "to %.0f %% of units with a bias above 2 points show a speed bias under 1 %% "
        "(Figure 7(c)); most offsets are therefore detection-threshold errors, corrected "
        "by the detector's sensitivity or length constant rather than its loop spacing. "
        "The speed bias also contains any real change of mean speed between the two "
        "sites, so the spacing part it gives is an upper bound."
        % (min(spx), max(spx), min(off50), max(off50), 100 * min(spf), 100 * max(spf)))

    h2(doc, "7.8 Reclassification, Corrected Dataset and Supply")
    num(doc,
        "Reclassification unit (400). For every interval of a site with a fitted bias, the "
        "count below each boundary is set to round((s - b) n), where s is the interval's "
        "share below the boundary, b the site's bias there and n the interval total, "
        "bounded to [0, n] and kept non-decreasing across boundaries; the class counts are "
        "the successive differences. The total count (405) and the mean speed are not "
        "altered. The operation takes %.1f microseconds per fifteen-minute record on one "
        "desktop processor (Table 14), against a reporting interval of 900 s." % us_rec)
    num(doc,
        "The correction restores agreement between sites. On the longest May chain of the "
        "M25, %d detector sites carrying the same traffic, the share below 520 cm ranges "
        "over %.1f points as recorded and %.1f points after correction with April biases "
        "(Figure 5(a)). Biases fitted on March remove %.0f %% to %.0f %% of the "
        "disagreement over the pairs touching a corrected detector in April, May and June, "
        "and %.0f %% to %.0f %% over all admitted pairs, including pairs whose detectors "
        "lie on chains too short to be corrected (Table 8); the other boundaries follow "
        "(Figure 5(d))."
        % (len(ch["detectors"]), ch["raw_range_pct"], ch["corrected_range_pct"],
           100 * min(red_t), 100 * max(red_t), 100 * min(red_a), 100 * max(red_a)))
    num(doc,
        "Two alternatives were measured against the per-boundary correction. The "
        "proportional form leaves more disagreement in every month on both corridors "
        "(in June %s against %s points on the M1, Figure 5(c)). The single-offset form "
        "matches it at 520 cm but not at the smaller boundaries: in April on the M25 it "
        "leaves %s points at 660 cm against %s (Table 11). The offset is therefore the "
        "unit's physical correction and diagnosis, and the per-boundary biases correct "
        "the record."
        % (pt(e1v("M1", "2025-06"), 2), pt(e1v("M1", "2025-06", "all/ratio"), 2),
           pt(C8["M25"]["2025-04"]["rank_one"]["touched_corrected"][1], 2),
           pt(C8["M25"]["2025-04"]["per_boundary"]["touched_corrected"][1], 2)))
    num(doc,
        "Corrected dataset (810) and supply (820). The corrected counts are written to the "
        "archive in place of the received class counts. In live operation the same "
        "reclassification is applied to each record from a detector card as it arrives, "
        "with the biases of the latest fit, and the corrected record is supplied to the "
        "traffic-management process; the results of Section 8 are this computation "
        "applied one to three months after the fit.")

    h2(doc, "7.9 Maintenance Record, Dated Changes and Refit")
    num(doc,
        "Maintenance record (830). Each unit's offset, fault class, spacing part, biases, "
        "standard errors and fitting period are written to the record. A unit may be "
        "marked for attendance where its bias at 520 cm exceeds three standard errors and "
        "a tolerance set by the operator; at 1 point the mark falls on %d of %d fitted M25 "
        "units in March and at 5 points on %d (Figure 9(c)). A local rule, marking a unit "
        "only when both its pairs disagree beyond the tolerance, disturbs other marks less "
        "often (%.2f against %.2f per injected fault on the M25 at 2 points) but newly "
        "marks only %s previously unmarked injected sites against %s (Figure 9(d))."
        % (fl25["march_counts"]["0.01"]["median"], fl25["march_counts"]["0.01"]["detectors"],
           fl25["march_counts"]["0.05"]["median"],
           fl25["injection"]["local T=0.02"]["other_flips_per_injection"],
           fl25["injection"]["median T=0.02"]["other_flips_per_injection"],
           fl25["injection"]["local T=0.02"]["newly_flagged"].replace("/", " of "),
           fl25["injection"]["median T=0.02"]["newly_flagged"].replace("/", " of ")))
    num(doc,
        "Step-event detector (330). For each site belonging to two candidate pairs, the "
        "daily difference series of each pair at 520 cm is searched for its largest step, "
        "the split maximising the difference of the means on either side scaled by the "
        "pooled standard deviation, with at least %d days on each side. An event is "
        "attributed to the site when the two steps fall within %d day of each other, each "
        "exceeds %s points, they are of opposite sign and their sum is under %.0f %% of "
        "the larger. Of %d sites tested, %d show the signature and the same rule applied "
        "to %d mismatched pairings finds %d (Figure 8, Table 13). At M1 site %s the pairs "
        "step by %+.1f and %+.1f points on %s. The steps appear at 660 cm as well, and in "
        "%d of the %d events with both pairs measured, the two pairs agree on the ratio of "
        "the 660 cm step to the 520 cm step within 0.1, as one change of one unit should. "
        "Applied to the step of each event, the cross-boundary test finds %d of %d "
        "consistent with a change of length offset, of %.0f to %.0f cm; the steps of the "
        "other %d, over 30 points at 520 cm, are larger than any length offset within "
        "300 cm could produce, a fault of another kind."
        % (e3_steps.MIN_SIDE, e3_steps.STEP_DAY_TOL, pt(e3_steps.MIN_STEP, 0),
           100 * e3_steps.CANCEL_TOL, n_tested, len(ev_all), n_ctrl, n_ctrl_ev,
           ev332["detector"], 100 * ev332["step_upstream_pair"],
           100 * ev332["step_downstream_pair"], day_label(ev332["day_index"]),
           ratio_agree, len(both_pairs), len(ev_pass), len(ev10),
           min(abs(e["offset_change_cm"]) for e in ev_pass),
           max(abs(e["offset_change_cm"]) for e in ev_pass), len(ev_fail)))
    num(doc,
        "Refit scheduler (340). All biases and offsets are refitted each period and "
        "applied to the next, and a site with a dated event is refitted from the day of "
        "the event. In June, a fit on May leaves %s points of disagreement on the M1 "
        "against %s for a fit on March, and %s against %s on the M25 (Table 9, "
        "Figure 5(b))."
        % (pt(e5v("M1", "rolling", "2025-06", "corrected_mean_abs_gap_520"), 2),
           pt(e5v("M1", "fixed", "2025-06", "corrected_mean_abs_gap_520"), 2),
           pt(e5v("M25", "rolling", "2025-06", "corrected_mean_abs_gap_520"), 2),
           pt(e5v("M25", "fixed", "2025-06", "corrected_mean_abs_gap_520"), 2)))

    h2(doc, "7.10 Sensitivity Correction of the Detector Unit and Platform")
    num(doc,
        "Sensitivity correction (840). For a unit classed as having a length-offset fault, "
        "the offset is the length its detector card (111) should subtract from each "
        "vehicle length it computes, or, where the card holds a length constant or a "
        "sensitivity setting, the change of that setting which shortens every length by "
        "the offset; it is sent to the card over the link (130) or applied at the next "
        "attendance. Until then the reclassification unit applies the biases to every "
        "record of the unit. The retune shows in the counts as equal and opposite steps "
        "in the unit's two pairs, which the step-event detector (330) dates, and the "
        "refit takes the unit's new state from that day, so the correction is verified "
        "from the same counts that prompted it. A unit that fails the cross-boundary test "
        "receives no sensitivity correction; its record names the kind of fault.")
    num(doc,
        "Platform. The system (900) is a server or desktop processor attached to the "
        "archive: a month's fit takes %.1f s for April's %d admitted M25 pairs and %.1f s "
        "for the M1's %d."
        % (t7["M25"]["fit_seconds"], t7["M25"]["admitted_pairs"], t7["M1"]["fit_seconds"],
           t7["M1"]["admitted_pairs"]))

    h2(doc, "7.11 Alternative Embodiments and Equivalent Implementations")
    num(doc,
        "Reference site. Where one site of a chain has had its length measurement verified "
        "with probe vehicles, or against a manual classified count on the same link as in "
        "Section 8.12, its bias may be set to that found in place of the median condition, "
        "which gives the chain, and every unit's offset, an absolute calibration.")
    num(doc,
        "Other corrections, tests and length distributions. Vehicles may be moved in "
        "proportion to the share below the boundary, or by the single offset alone; both "
        "were measured and are inferior on the data of Section 8. The cross-boundary "
        "pattern may instead be taken as the principal direction of the units' bias "
        "vectors, which is fitted to the data rather than predicted; on units with a bias "
        "above 2 points the two tests agree on %.0f %% to %.0f %% of units. The length "
        "distribution may be interpolated by other monotone curves or taken from a "
        "classifier with more boundaries." % (100 * min(agree9), 100 * max(agree9)))
    num(doc,
        "Other detector hardware, class schemes, fitting periods and networks: radar, "
        "video, piezoelectric and axle-based classifiers; any number of length classes; "
        "per-lane records; a week or a rolling window; any road on which consecutive "
        "detector sites observe the same traffic stream, including classified-count "
        "stations on national highways.")

    # ============================================================== 8
    h1(doc, "8. Experimental Validation Results")
    h2(doc, "8.1 Experimental / Validation Setup")
    para(doc,
         "The system was evaluated on field records from motorway inductive-loop detector "
         "sites: fifteen-minute class counts, totals and mean speeds from the WebTRIS "
         "service of National Highways, England, for carriageway A of the M25 (reference) "
         "and of the M1 (replication), March to June 2025. Validation ran at four levels, "
         "and the scope of each is stated so that the evidence is read for what it is.")
    bullets(doc, [
        ("(i) Out of sample.", "Biases are fitted on one month and scored on later months "
         "the fit did not see. The score is agreement between adjacent sites that measure "
         "the same vehicles, on two populations: the pairs touching a corrected detector, "
         "and all admitted pairs. Since totals are kept, the correction cannot change which "
         "pairs are admitted."),
        ("(ii) Known faults.", "Faults of known size and kind are injected into the "
         "records of real detector units and the chain is solved again: share faults for "
         "localisation, and length offsets, class-boundary faults and common shifts of a "
         "whole chain for the cross-boundary test."),
        ("(iii) Pairs never fitted.", "Consecutive sites with a slip road between them "
         "are never admitted; their two sites lie on different chains, each referred to "
         "its own median, and they share the through traffic."),
        ("(iv) Independent ground truth.", "No record of individual vehicle lengths exists "
         "for these sites. The independent reference is the manual classified counts of "
         "the Department for Transport, made by trained enumerators by vehicle type over "
         "twelve hours at count points on the same roads and published as open data; "
         "carriageway B was fitted near each count point for this comparison."),
    ])

    h2(doc, "8.2 Configuration / Operating Parameter Values")
    table(doc, ["Parameter", "Value", "Basis"], [
        ["Class boundaries", "520, 660, 1160 cm", "Detector classification"],
        ["Reporting interval", "15 min (900 s)", "Detector card"],
        ["Candidate pair", "No slip-road site between; at most 3 km", "Site register"],
        ["Free-flow condition", "Mean speed above %.0f mph at both sites"
         % pairs.MIN_SPEED_MPH, "Wu and Coifman [N7]"],
        ["Interval volume, intervals per month", "At least %d vehicles; at least %d"
         % (pairs.MIN_VOLUME, pairs.MIN_INTERVALS), "Standard error of the shares"],
        ["Balance tolerance of totals", "%.0f %%" % (100 * pairs.BALANCE_TOL),
         "Same vehicles at both sites"],
        ["Fitting period", "One calendar month, applied to the next", "Section 7.9"],
        ["Chain reference", "Median of the chain's biases", "Section 7.4"],
        ["Length distribution anchors", "0 below 250 cm; 1 at %d cm" % e8_CDF_TOP,
         "Longest permitted vehicle"],
        ["Cross-boundary limit and floor", "%.1f of max(|b(520)|, %s points)"
         % (e9["limit"], pt(e9["floor"], 0)), "Section 7.6"],
        ["Step detection", "%d days each side; steps within %d day, over %s points, sum "
         "under %.0f %%" % (e3_steps.MIN_SIDE, e3_steps.STEP_DAY_TOL,
                            pt(e3_steps.MIN_STEP, 0), 100 * e3_steps.CANCEL_TOL),
         "Section 7.9"],
        ["Manual-count match", "Nearest fitted detector within %d m; totals within %.0f %%"
         % (e11["max_dist_m"], 100 * e11["tol"]), "Same vehicles as the count"],
    ], cap="Table 6: Configuration Parameters", widths=[1.9, 2.8, 1.8], size=9,
        align_left_cols=(0, 1, 2))
    para(doc, "Table 6 holds for every test unless a test states otherwise.")

    h2(doc, "8.3 Validation Tests and Outcome Measures")
    table(doc, ["ID", "Test", "Size", "Condition"], [
        ["V1", "Disagreement between adjacent detectors",
         "%d + %d candidate pairs" % (p25["pairs_march"], p1["pairs_march"]),
         "March 2025, persistence to June"],
        ["V2", "Correction of months the fit did not see",
         "%d-%d admitted pairs a month" % (min(E1[r][m]["admitted_pairs"] for r in R
                                               for m in months),
                                           max(E1[r][m]["admitted_pairs"] for r in R
                                               for m in months)),
         "Fit on March, scored April to June"],
        ["V3", "Refit each period", "Both corridors, May and June",
         "Fit on the month before against fit on March"],
        ["V4", "Localisation of injected faults", "%d injections" % n_inj,
         "Share faults of %s to %s points, April chains of four or more"
         % (pt(min(e2["sizes"]), 0), pt(max(e2["sizes"]), 0))],
        ["V5", "One offset per unit and cross-boundary test",
         "8 corridor-months, %d-%d units each"
         % (min(F9[r][m]["detectors"] for r in R for m in MM),
            max(F9[r][m]["detectors"] for r in R for m in MM)),
         "Prediction with no fitted parameter; shuffled control"],
        ["V6", "Injected faults of known kind", "%d units, four kinds of fault, both signs" % n_inj9,
         "40 cm offsets, 660 and 1160 cm boundary faults, 30 cm chain shifts"],
        ["V7", "Dated changes", "%d sites; %d control pairings" % (n_tested, n_ctrl),
         "Daily series, March to June"],
        ["V8", "Admission conditions", "3 variants, both corridors",
         "Balance test dropped or relaxed; free-flow condition dropped"],
        ["V9", "Pairs across a junction, never fitted",
         "%s + %s pairs a month" % (span(jr("M25", "pairs")), span(jr("M1", "pairs"))),
         "March biases; shuffled-bias control"],
        ["V10", "Independent manual classified counts", "%d matched counts" % G["matched"],
         "Both carriageways, March to June 2025"],
        ["V11", "Computation", "One corridor-month", "One desktop processor"],
    ], cap="Table 7: Validation Tests", widths=[0.45, 2.2, 1.6, 2.3], size=8.8,
        align_left_cols=(1, 2, 3))
    para(doc,
         "The outcome measures are the mean absolute difference below 520 cm between "
         "adjacent sites, in points; the share of injections in which the injected site "
         "shows the largest change; the share of injected faults given the right class; "
         "and, against manual counts, the slope of each detector's recorded departure on "
         "its fitted bias, 1 when the fitted bias is the whole error.")

    h2(doc, "8.4 Disagreement Between Adjacent Detectors (V1)")
    para(doc,
         "Consecutive sites that count the same vehicles disagree in their share below "
         "520 cm by a median of %s points on the M25 and %s on the M1, and the "
         "disagreement of each pair persists from March to June (Figure 4(b)): it belongs "
         "to the sites, not the traffic. Pairs whose totals do not agree carry the largest "
         "differences (Figure 4(c)), which is why admission requires the balance of totals."
         % (pt(p25["abs_gap520_median"]), pt(p1["abs_gap520_median"])))
    figure(doc, "fig04.png",
           "Figure 4: Disagreement in length classification between adjacent detector "
           "sites that count the same vehicles",
           "(a) Differences in share below 520 cm across the candidate pairs of March; (b) "
           "the same pairs in March and June, showing the difference persists as a property "
           "of the sites; (c) difference against the ratio of totals, with the admission "
           "band; (d) fitted biases of all units on chains of three or more.")

    h2(doc, "8.5 Correction of Months the Fit Did Not See (V2)")
    rows = []
    for road in R:
        for m in months:
            rows.append([road, mname[m], E1[road][m]["all/additive"]["touched_pairs"],
                         pt(e1raw(road, m), 2), pt(e1v(road, m), 2),
                         pt(e1v(road, m, "all/ratio"), 2),
                         E1[road][m]["admitted_pairs"],
                         "%s -> %s" % (pt(e1allraw(road, m), 2), pt(e1all(road, m), 2))])
    table(doc, ["Corridor", "Month", "Touched pairs", "Recorded", "Claimed correction",
                "Proportional (7.11)", "All pairs", "All pairs: recorded -> corrected"],
          rows,
          cap="Table 8: Mean Difference Below 520 cm Between Adjacent Sites, Points, Biases "
              "Fitted on March",
          widths=[0.7, 0.6, 0.65, 0.75, 0.85, 0.85, 0.6, 1.25], size=8.6)
    para(doc,
         "Table 8 and Figure 5 show the claimed correction removing %.0f %% to %.0f %% of "
         "the disagreement on the pairs it touches and %.0f %% to %.0f %% over all admitted "
         "pairs, in every month on both corridors, and leaving less than the proportional "
         "alternative in every month."
         % (100 * min(red_t), 100 * max(red_t), 100 * min(red_a), 100 * max(red_a)))
    figure(doc, "fig05.png",
           "Figure 5: Correction of months the fit did not see",
           "(a) One chain of the M25 in May, as recorded and corrected with April biases; "
           "(b) mean difference over pairs touching a corrected detector, as recorded, with a "
           "March fit and with a fit on the month before; (c) the claimed per-boundary "
           "correction against the proportional alternative; (d) all three boundaries in "
           "April.")

    h2(doc, "8.6 Refit Each Period (V3)")
    rows = []
    for road in R:
        for m in months:
            rows.append([road, mname[m],
                         pt(e5v(road, "fixed", m, "raw_mean_abs_gap_520"), 2),
                         pt(e5v(road, "fixed", m, "corrected_mean_abs_gap_520"), 2),
                         e5v(road, "fixed", m, "worse_by_1pt"),
                         pt(e5v(road, "rolling", m, "raw_mean_abs_gap_520"), 2),
                         pt(e5v(road, "rolling", m, "corrected_mean_abs_gap_520"), 2),
                         e5v(road, "rolling", m, "worse_by_1pt")])
    table(doc, ["Corridor", "Month", "Recorded", "Fit on March", "Worse", "Recorded",
                "Fit on month before", "Worse"], rows,
          cap="Table 9: Fixed Fit Against Refit on the Month Before, Pairs Touching a "
              "Corrected Detector, Points",
          widths=[0.75, 0.6, 0.8, 0.8, 0.55, 0.8, 1.05, 0.55], size=9)
    para(doc,
         "Table 9 shows refitting roughly halving what a March fit leaves in June. The "
         "recorded columns differ slightly because a refit covers the sites fitted in the "
         "month before; 'Worse' counts pairs made worse by more than 1 point.")

    h2(doc, "8.7 Localisation of Injected Faults (V4)")
    figure(doc, "fig06.png",
           "Figure 6: Known faults injected into the records of real detector sites",
           "(a) Share of injections in which the injected site shows the largest change on "
           "its chain; (b) error in the fitted size; (c) and (d) the same by chain length, "
           "where the median reference limits absolute size on short chains.")
    rows = []
    for road in R:
        for k, c in e2["roads"][road]["by_chain_size"].items():
            rows.append([road, k, c["injections"], "%.0f %%" % (100 * c["localised"]),
                         "%.0f %%" % (100 * c["median_size_error_over_injection"]),
                         "%.0f %%" % (100 * c["p90_size_error_over_injection"])])
    table(doc, ["Corridor", "Sites on chain", "Injections", "Largest change at injected site",
                "Size error / injection, median", "90th percentile"], rows,
          cap="Table 10: Injected Faults by Chain Length, Both Shapes and All Sizes Pooled",
          widths=[0.8, 0.9, 0.8, 1.3, 1.3, 1.0], size=9)
    para(doc,
         "Faults were injected into each site of every April chain of four or more, one at "
         "a time; by fault size the injected site shows the largest change in %.0f %% to "
         "%.0f %% of injections (Table 14), and by chain length in %.0f %% to %.0f %% "
         "(Table 10). Localisation holds on every chain length; absolute size on short "
         "chains carries the error of the median reference."
         % (100 * min(locsz), 100 * max(locsz), 100 * min(locsize), 100 * max(locsize)))

    h2(doc, "8.8 One Length Offset per Unit and the Cross-Boundary Test (V5, V6)")
    figure(doc, "fig07.png",
           "Figure 7: One length offset per detector unit, recovered from class counts",
           "(a) Bias at 660 cm against bias at 520 cm for every fitted unit, with the fitted "
           "ratio and the ratio predicted from the stream's length density; (b) the two "
           "ratios in every corridor-month; (c) offsets in cm and the part a speed error "
           "could explain; (d) the cross-boundary test: bias at 660 cm predicted from each "
           "unit's offset alone against the bias solved along the chain, units failing the "
           "test marked.")
    rows = []
    for road in R:
        for m in MM:
            s_ = S8[road][m]
            n_ = X[road]["null"][m]
            rows.append([road, mname[m], s_["detectors"],
                         "%.0f %% (%.0f %%)" % (100 * n_["rank1_standardised"],
                                              100 * n_["rank1_standardised_null_mean"]),
                         "%.2f / %.2f" % (n_["corr_520_660"], n_["corr_520_1160"]),
                         "%.3f" % s_["loadings"][1], "%.3f" % s_["cdf_density_ratio"][1],
                         "%.0f / %.0f" % (s_["offset_cm_q50_q75_q90"][0],
                                          s_["offset_cm_q50_q75_q90"][2]),
                         "%.1f" % SP[road][m]["offset_cm_explained_by_speed_q50"]])
    table(doc, ["Corridor", "Month", "Units", "One offset, scaled (shuffled)",
                "r with 520 cm: 660 / 1160", "Ratio 660/520 fitted",
                "Ratio predicted", "|Offset| median / p90, cm",
                "Offset from speed, median cm"], rows,
          cap="Table 11: The Single Length-Offset Model in Every Corridor-Month",
          widths=[0.6, 0.55, 0.45, 0.7, 0.7, 0.7, 0.85, 0.85, 0.85], size=8.3)
    para(doc,
         "In Table 11 one offset per unit explains far more of the scaled bias variance "
         "than the shuffled control in every corridor-month. Correcting from the offset "
         "alone, instead of per boundary, leaves %s against %s points at 660 cm and %s "
         "against %s at 1160 cm on the pairs it touches (M25, April), which is why the "
         "per-boundary biases correct the record."
         % (pt(C8["M25"]["2025-04"]["rank_one"]["touched_corrected"][1], 2),
            pt(C8["M25"]["2025-04"]["per_boundary"]["touched_corrected"][1], 2),
            pt(C8["M25"]["2025-04"]["rank_one"]["touched_corrected"][2], 2),
            pt(C8["M25"]["2025-04"]["per_boundary"]["touched_corrected"][2], 2)))
    rows = []
    for road in R:
        for m in MM:
            f_ = F9[road][m]
            rows.append([road, mname[m], f_["detectors"], "%.2f" % f_["corr_660"],
                         "%.0f %%" % (100 * f_["variance_660_explained_by_prediction"]),
                         "%.2f" % f_["slope_660_observed_on_predicted"],
                         "%d / %d" % (f_["fail_all"], f_["detectors"]),
                         "%.0f %%" % (100 * f_["agreement_with_principal_direction_big"])])
    for road in R:
        i_ = I9[road]
        rows.append([road, "April, injected", i_["offset +40 cm"]["injections"],
                     "660 cm fault: %.0f / %.0f %%" % (100 * i_["boundary660 +100 cm"]["rate"],
                                                     100 * i_["boundary660 -100 cm"]["rate"]),
                     "1160 cm fault: %.0f / %.0f %%" % (
                         100 * i_["boundary1160 +150 cm"]["rate"],
                         100 * i_["boundary1160 -150 cm"]["rate"]),
                     "offset 40 cm kept: %.0f / %.0f %%" % (
                         100 * i_["offset +40 cm"]["rate"], 100 * i_["offset -40 cm"]["rate"]),
                     "chain moved 30 cm: %s, %s changed" % tuple(
                         i_["common_offset_class_changes"].values()), "-"])
    table(doc, ["Corridor", "Month", "Units", "Corr. 660 cm, predicted vs solved",
                "Variance explained", "Slope solved / predicted", "Units failing",
                "Agreement with principal direction (> 2 pt)"], rows,
          cap="Table 12: The Cross-Boundary Test: One Offset Predicts the Other Boundaries",
          widths=[0.6, 0.85, 0.5, 1.0, 0.9, 0.95, 0.75, 0.95], size=8.2,
          align_left_cols=(3, 4, 5, 6))
    para(doc,
         "Upper rows of Table 12: field data, every fitted unit; the prediction uses no "
         "parameter fitted to the biases. Lower rows: faults injected into real units (+ / "
         "- sign of the fault), the share of passing units turned failing, or kept passing "
         "for a true offset, and the class changes when the whole chain is moved by a "
         "common offset. The test separates a 660 cm boundary fault from a length offset "
         "and ignores the chain's constant; at 1160 cm it detects about half of the "
         "injected faults.")

    h2(doc, "8.9 Dated Changes (V7)")
    figure(doc, "fig08.png",
           "Figure 8: Dated changes in length measurement attributed to one detector unit",
           "(a), (b) Daily differences of the two pairs of M1 sites %s and %s, stepping on "
           "the same day by equal and opposite amounts; (c) the two steps of every event "
           "against the equal-and-opposite line; (d) the signature in tested sites and in "
           "mismatched control pairings." % (shown[0][1], shown[1][1]))
    rows = []
    for road, ev in ev_all:
        c = [x for x in cont if x["road"] == road and x["detector"] == ev["detector"]][0]
        if c["both_pairs_all_days"]:
            note = "both pairs report every day"
        elif c["missing_upstream_pair"] == c["missing_downstream_pair"]:
            note = "site gave no usable data on %d days" % len(c["missing_upstream_pair"])
        else:
            note = "one pair missing before the step"
        s_ = evsteps[(road, ev["detector"])]
        rat = []
        for v in (s_["step_pair_downstream"], s_["step_pair_upstream"]):
            rat.append("%.2f" % (v[1] / v[0]) if v else "-")
        rows.append([road, ev["detector"], day_label(ev["day_index"]),
                     "%+.1f" % (100 * ev["step_upstream_pair"]),
                     "%+.1f" % (100 * ev["step_downstream_pair"]), " / ".join(rat), note])
    table(doc, ["Corridor", "Site", "Day of step", "Pair where downstream, pts",
                "Pair where upstream, pts", "Step ratio 660/520, both pairs",
                "Data within 5 days"], rows,
          cap="Table 13: Changes in Length Measurement Attributed to One Site and One Day",
          widths=[0.7, 0.55, 1.05, 0.9, 0.9, 0.9, 1.25], size=8.5, align_left_cols=(6,))
    para(doc,
         "In Table 13 every event steps by nearly equal and opposite amounts in the two "
         "pairs. Where both pairs lack the same days, the site itself gave no usable "
         "intervals around the step, which is consistent with work at the site; "
         "maintenance logs would confirm it (Table 16).")

    h2(doc, "8.10 Admission Conditions and the Attendance Mark (V8)")
    para(doc,
         "The conditions of admission were tested by fitting without them (Figure 9(a), "
         "(b)). Without the balance test, the M1 fit admits more pairs and leaves %s, %s "
         "and %s points on the pairs it touches in April, May and June, against %s, %s and "
         "%s with it; on the M25, where the test removes at most %d touched pairs, the "
         "difference is under %s points. Without the free-flow condition the corrected "
         "disagreement is larger in every month on both corridors. The attendance mark of "
         "Section 7.9 is shown in Figure 9(c), (d)."
         % (pt(nogate_m1[0], 2), pt(nogate_m1[1], 2), pt(nogate_m1[2], 2),
            pt(claimed_m1[0], 2), pt(claimed_m1[1], 2), pt(claimed_m1[2], 2),
            gate_pairs_m25, pt(0.0002, 2)))
    figure(doc, "fig09.png",
           "Figure 9: Admission conditions and the attendance mark",
           "(a), (b) Corrected disagreement when the fit is made with the balance test or "
           "the free-flow condition relaxed; (c) share of fitted units marked for attendance "
           "at each tolerance under the median and local rules; (d) injected faults newly "
           "marked against marks of other units disturbed.")

    h2(doc, "8.11 Pairs Across a Junction, Never Fitted (V9)")
    para(doc,
         "The agreement above is on pairs the fit used. A test the fit never saw is given "
         "by consecutive sites with a slip road between them: they are never admitted, their "
         "two sites lie on different chains, each referred to its own median, and they see "
         "the same through traffic with the slip-road traffic added or removed. With March "
         "biases, the mean difference below 520 cm across such pairs falls on the M25 from "
         "%.1f-%.1f to %.1f-%.1f points in April to June, every one of the %s pairs "
         "improving, while the biases of randomly chosen units raise it to %.1f-%.1f. On the "
         "M1 it falls only from %.1f-%.1f to %.1f-%.1f, and %s of %s pairs improve: "
         "there the median references of separate chains, or the class mix of the "
         "slip-road traffic, differ by several points. The biases are therefore real "
         "beyond the pairs fitted, and a common absolute level across chains is not "
         "established without a verified reference site."
         % (min(jr("M25", "raw")), max(jr("M25", "raw")), min(jr("M25", "corrected")),
            max(jr("M25", "corrected")), span(jr("M25", "pairs")),
            min(jr("M25", "control_shuffled")), max(jr("M25", "control_shuffled")),
            min(jr("M1", "raw")), max(jr("M1", "raw")), min(jr("M1", "corrected")),
            max(jr("M1", "corrected")), span(jr("M1", "improved")), span(jr("M1", "pairs"))))

    h2(doc, "8.12 Comparison with Independent Manual Classified Counts (V10)")
    para(doc,
         "At %d manual counts in March to June 2025 on the two motorways, both "
         "carriageways, the nearest detector on a chain of three or more within %d m was "
         "taken where its total over the counted free-flow hours agreed with the manual "
         "total within %.0f %%, so that it counted the same vehicles. Vehicle types are not "
         "lengths, but cars, motorcycles and vans lie below 660 cm and heavier vehicles "
         "above it, and the share below 520 cm is that of cars and motorcycles plus the "
         "shorter vans, the van share being taken as a covariate. The bias of each "
         "detector, fitted from detector pairs alone, predicts how far its recorded share "
         "departs from the manual count: at 520 cm the departure rises by %.2f +- %.2f "
         "points per point of fitted bias (r = %.2f, rank correlation %.2f), and at 660 cm "
         "by %.2f +- %.2f (r = %.2f), against 1 if the fitted bias is the whole error "
         "(Figure 10(a), (b)). After correction the spread of the departures across the "
         "counts, at detectors on separate chains each referred to its own median, falls "
         "from %.1f to %.1f points at 520 cm and from %.1f to %.1f at 660 cm (median "
         "absolute deviation, Figure 10(c), (d)). On the M25 alone the slope at 520 cm is "
         "%.2f +- %.2f with r = %.2f; on the M1 it is %.2f +- %.2f with r = %.2f, and at "
         "660 cm the M1 shows no relation, as across its junctions. At 1160 cm, where the "
         "biases are a fraction of a point, the manual counts do not resolve them. The "
         "common level of the departure, %.1f points at 660 cm, holds the vans longer than "
         "660 cm and any error common to all chains, which the counts cannot separate."
         % (G["matched"], e11["max_dist_m"], 100 * e11["tol"],
            g520["beta"], g520["beta_se"], g520["r"], G520["spearman"],
            G660["fit_raw_deviation_on_bias"]["beta"], G660["fit_raw_deviation_on_bias"]["beta_se"],
            G660["fit_raw_deviation_on_bias"]["r"],
            G520["mad_raw_pts"], G520["mad_corrected_pts"], G660["mad_raw_pts"],
            G660["mad_corrected_pts"],
            GR["M25"]["b520"]["fit_raw_deviation_on_bias"]["beta"],
            GR["M25"]["b520"]["fit_raw_deviation_on_bias"]["beta_se"],
            GR["M25"]["b520"]["fit_raw_deviation_on_bias"]["r"],
            GR["M1"]["b520"]["fit_raw_deviation_on_bias"]["beta"],
            GR["M1"]["b520"]["fit_raw_deviation_on_bias"]["beta_se"],
            GR["M1"]["b520"]["fit_raw_deviation_on_bias"]["r"],
            G660["fit_raw_deviation_on_bias"]["alpha_pts"]))
    figure(doc, "fig10.png",
           "Figure 10: Fitted biases against independent manual classified counts",
           "(a), (b) Departure of each detector's recorded share from the manual count at "
           "the same place and hours, against the bias fitted from detector pairs alone, "
           "at 520 cm and 660 cm, with the slope-1 line; (c), (d) the departures about "
           "their median, as recorded and after correction.")

    h2(doc, "8.13 Overall Performance Metrics")
    rows = [
        ["Candidate pairs in March (no slip road between, enough data)", p25["pairs_march"],
         p1["pairs_march"]],
        ["Difference below 520 cm between paired sites, median (March), points",
         pt(p25["abs_gap520_median"]), pt(p1["abs_gap520_median"])],
        ["Admitted pairs, April / May / June",
         " / ".join(str(E1["M25"][m]["admitted_pairs"]) for m in months),
         " / ".join(str(E1["M1"][m]["admitted_pairs"]) for m in months)],
        ["All admitted pairs, mean difference below 520 cm, recorded -> corrected (Apr-Jun), "
         "points",
         "%s-%s -> %s-%s" % (pt(min(e1allraw("M25", m) for m in months), 2),
                            pt(max(e1allraw("M25", m) for m in months), 2),
                            pt(min(e1all("M25", m) for m in months), 2),
                            pt(max(e1all("M25", m) for m in months), 2)),
         "%s-%s -> %s-%s" % (pt(min(e1allraw("M1", m) for m in months), 2),
                            pt(max(e1allraw("M1", m) for m in months), 2),
                            pt(min(e1all("M1", m) for m in months), 2),
                            pt(max(e1all("M1", m) for m in months), 2))],
        ["Pairs touching a corrected detector, the same, points",
         "%s-%s -> %s-%s" % (pt(min(e1raw("M25", m) for m in months), 2),
                            pt(max(e1raw("M25", m) for m in months), 2),
                            pt(min(e1v("M25", m) for m in months), 2),
                            pt(max(e1v("M25", m) for m in months), 2)),
         "%s-%s -> %s-%s" % (pt(min(e1raw("M1", m) for m in months), 2),
                            pt(max(e1raw("M1", m) for m in months), 2),
                            pt(min(e1v("M1", m) for m in months), 2),
                            pt(max(e1v("M1", m) for m in months), 2))],
        ["Unit biases at 520 and 660 cm across units, correlation, Mar-Jun",
         "%.2f-%.2f" % (min(X["M25"]["null"][m]["corr_520_660"] for m in MM),
                        max(X["M25"]["null"][m]["corr_520_660"] for m in MM)),
         "%.2f-%.2f" % (min(X["M1"]["null"][m]["corr_520_660"] for m in MM),
                        max(X["M1"]["null"][m]["corr_520_660"] for m in MM))],
        ["Manual counts: departure per point of fitted bias at 520 cm (1 = whole error)",
         "%.2f +- %.2f (n %d)" % (GR["M25"]["b520"]["fit_raw_deviation_on_bias"]["beta"],
                                 GR["M25"]["b520"]["fit_raw_deviation_on_bias"]["beta_se"],
                                 GR["M25"]["b520"]["n"]),
         "%.2f +- %.2f (n %d)" % (GR["M1"]["b520"]["fit_raw_deviation_on_bias"]["beta"],
                                 GR["M1"]["b520"]["fit_raw_deviation_on_bias"]["beta_se"],
                                 GR["M1"]["b520"]["n"])],
        ["Manual counts: spread of departure below 520 cm, recorded -> corrected, points",
         "%.1f -> %.1f" % (GR["M25"]["b520"]["mad_raw_pts"], GR["M25"]["b520"]["mad_corrected_pts"]),
         "%.1f -> %.1f" % (GR["M1"]["b520"]["mad_raw_pts"], GR["M1"]["b520"]["mad_corrected_pts"])],
        ["Pairs across a junction, never fitted: recorded -> corrected (Apr-Jun), points",
         "%.1f-%.1f -> %.1f-%.1f" % (min(jr("M25", "raw")), max(jr("M25", "raw")),
                                    min(jr("M25", "corrected")), max(jr("M25", "corrected"))),
         "%.1f-%.1f -> %.1f-%.1f" % (min(jr("M1", "raw")), max(jr("M1", "raw")),
                                    min(jr("M1", "corrected")), max(jr("M1", "corrected")))],
        ["Length offset of units with bias > 2 points, median / 90th pct (March), cm",
         "%.0f / %.0f" % (S8["M25"]["2025-03"]["offset_cm_q50_q75_q90"][0],
                          S8["M25"]["2025-03"]["offset_cm_q50_q75_q90"][2]),
         "%.0f / %.0f" % (S8["M1"]["2025-03"]["offset_cm_q50_q75_q90"][0],
                          S8["M1"]["2025-03"]["offset_cm_q50_q75_q90"][2])],
        ["Injected faults: injected site shows largest change, range over fault sizes",
         "%.0f-%.0f %%" % (100 * min(e2["roads"]["M25"]["additive %+.2f" % s]["localised"]
                                    for s in e2["sizes"]),
                          100 * max(e2["roads"]["M25"]["additive %+.2f" % s]["localised"]
                                    for s in e2["sizes"])),
         "%.0f-%.0f %%" % (100 * min(e2["roads"]["M1"]["additive %+.2f" % s]["localised"]
                                    for s in e2["sizes"]),
                          100 * max(e2["roads"]["M1"]["additive %+.2f" % s]["localised"]
                                    for s in e2["sizes"]))],
        ["Cross-boundary test: bias at 660 cm predicted from each unit's offset, "
         "correlation with solved bias, Mar-Jun",
         "%.2f-%.2f" % (min(F9["M25"][m]["corr_660"] for m in MM),
                        max(F9["M25"][m]["corr_660"] for m in MM)),
         "%.2f-%.2f" % (min(F9["M1"][m]["corr_660"] for m in MM),
                        max(F9["M1"][m]["corr_660"] for m in MM))],
        ["Injected 660 cm boundary fault classed as not an offset (+-100 cm)",
         "%.0f-%.0f %%" % (100 * min(I9["M25"][k]["rate"] for k in
                                     ("boundary660 +100 cm", "boundary660 -100 cm")),
                           100 * max(I9["M25"][k]["rate"] for k in
                                     ("boundary660 +100 cm", "boundary660 -100 cm"))),
         "%.0f-%.0f %%" % (100 * min(I9["M1"][k]["rate"] for k in
                                     ("boundary660 +100 cm", "boundary660 -100 cm")),
                           100 * max(I9["M1"][k]["rate"] for k in
                                     ("boundary660 +100 cm", "boundary660 -100 cm")))],
        ["Dated changes / sites tested; the same in control pairings",
         "%d / %d; %d / %d" % (e3["roads"]["M25"]["n_events"],
                               e3["roads"]["M25"]["detectors_tested"],
                               e3["roads"]["M25"]["control_events"],
                               e3["roads"]["M25"]["control_trials"]),
         "%d / %d; %d / %d" % (e3["roads"]["M1"]["n_events"],
                               e3["roads"]["M1"]["detectors_tested"],
                               e3["roads"]["M1"]["control_events"],
                               e3["roads"]["M1"]["control_trials"])],
        ["Fit of one corridor-month (April admitted pairs), s",
         "%.1f (%d)" % (t7["M25"]["fit_seconds"], t7["M25"]["admitted_pairs"]),
         "%.1f (%d)" % (t7["M1"]["fit_seconds"], t7["M1"]["admitted_pairs"])],
        ["Reclassification time per 15-min record, microseconds",
         "%.2f" % t7["M25"]["microseconds_per_record"],
         "%.2f" % t7["M1"]["microseconds_per_record"]],
    ]
    table(doc, ["Performance metric", "M25", "M1"], rows,
          cap="Table 14: System Performance Metrics on Recorded Motorway Detector Data",
          widths=[3.9, 1.2, 1.2], size=8.8, align_left_cols=(0,))
    para(doc,
         "Every value in Table 14 is computed from the recorded class counts of the two "
         "corridors or, for the manual-count rows, from the Department for Transport's "
         "counts. 'All admitted pairs' includes pairs whose detectors lie on chains too "
         "short to be corrected; the reclassification time is the latency the correction "
         "adds to a record in live use.")

    h2(doc, "8.14 Integrated Technical-Effect Summary")
    table(doc, ["Effect", "How achieved", "Evidence"], [
        ["Adjacent detectors agree in months the fit did not see",
         "Per-boundary biases solved along chains; total-preserving reclassification",
         "%s-%s -> %s-%s points on touched pairs; %.0f-%.0f %% less over all pairs "
         "(Table 8)" % (rng_raw_t + rng_cor_t + (100 * min(red_a), 100 * max(red_a)))],
        ["Faulty unit localised", "Each interior detector in two pairs",
         "Largest change at injected site %.0f-%.0f %% (Table 10)"
         % (100 * min(locsize), 100 * max(locsize))],
        ["Length offset of each unit in cm", "Stream's own length distribution from the "
                                             "class shares",
         "Median %.0f-%.0f cm, p90 %.0f-%.0f cm; interpolation x%.2f-%.2f (Table 11)"
         % (min(off50), max(off50), min(off90), max(off90), min(iv_ratio), max(iv_ratio))],
        ["Kind of fault told apart", "Cross-boundary test, no fitted parameter",
         "660 cm fault %.0f-%.0f %% flagged, 40 cm offset %.0f-%.0f %% kept, %d of %d "
         "changed by chain shift (Table 12)"
         % (100 * min(det660), 100 * max(det660), 100 * min(keep40), 100 * max(keep40),
            gauge_changed, gauge_units)],
        ["Biases carry the real error", "Fitted from detector pairs alone",
         "Slope %.2f +- %.2f against manual counts; spread %.1f -> %.1f points (Figure 10)"
         % (g520["beta"], g520["beta_se"], G520["mad_raw_pts"], G520["mad_corrected_pts"])],
        ["Correction holds beyond the pairs fitted", "Biases of separate chains",
         "M25 junction pairs %.1f-%.1f -> %.1f-%.1f points, shuffled %.1f-%.1f (8.11)"
         % (min(jr("M25", "raw")), max(jr("M25", "raw")), min(jr("M25", "corrected")),
            max(jr("M25", "corrected")), min(jr("M25", "control_shuffled")),
            max(jr("M25", "control_shuffled")))],
        ["Change dated to one unit and one day", "Equal and opposite steps in two pairs",
         "%d of %d sites; %d of %d control pairings (Table 13)"
         % (len(ev_all), n_tested, n_ctrl_ev, n_ctrl)],
        ["Correction follows the unit", "Refit each period",
         "June: %s against %s points (M1), %s against %s (M25) (Table 9)"
         % (pt(e5v("M1", "rolling", "2025-06", "corrected_mean_abs_gap_520"), 2),
            pt(e5v("M1", "fixed", "2025-06", "corrected_mean_abs_gap_520"), 2),
            pt(e5v("M25", "rolling", "2025-06", "corrected_mean_abs_gap_520"), 2),
            pt(e5v("M25", "fixed", "2025-06", "corrected_mean_abs_gap_520"), 2))],
        ["Correction in live operation", "Per-record reclassification",
         "%.1f microseconds per record against a 900 s interval (Table 14)" % us_rec],
    ], cap="Table 15: Technical Effects, Mechanism and Evidence",
        widths=[1.9, 2.1, 2.5], size=8.8, align_left_cols=(0, 1, 2))
    para(doc, "Table 15 ties each effect to the mechanism behind it and to the test that "
              "measured it.")

    h2(doc, "8.15 Key Validation Insights")
    bullets(doc, [
        "Adjacent detectors that see the same vehicles disagree persistently (r = %.3f "
        "between March and June on the M25): the error belongs to the site, not the "
        "traffic." % p25["persistence"]["2025-06"]["corr_gap520_vs_march"],
        "One offset per unit explains most of the bias pattern across boundaries, %.0f %% "
        "to %.0f %% against at most %.0f %% shuffled, strongly at 660 cm and weakly at "
        "1160 cm." % (100 * min(std_real), 100 * max(std_real), 100 * max(std_p99)),
        "The cross-boundary test separates a 660 cm boundary fault from a length offset "
        "and does not depend on the chain's reference; at 1160 cm it detects about half.",
        "The per-boundary correction beats both the proportional and the single-offset "
        "corrections: the offset diagnoses the unit and the biases correct the record.",
        "The fitted biases carry the real error: against manual counts the slope is %.2f "
        "at 520 cm, strongly on the M25 and weakly on the M1." % g520["beta"],
        "Biases drift, and refitting each period roughly halves what an old fit leaves.",
        "The median reference limits absolute size on short chains, not localisation; a "
        "verified site removes the limit.",
    ])

    h2(doc, "8.16 Limitations and Planned Validation")
    para(doc,
         "Agreement between corrected neighbours is established; whether a chain's median "
         "site is correct is not, so each offset is relative to its chain's reference, and "
         "a verified reference site makes it absolute. The cross-boundary test does not "
         "depend on that reference. The offsets in centimetres rest on the interpolated "
         "length distribution, whose prediction at 660 cm runs higher than the solved "
         "biases, and are physically consistent across boundaries, months and corridors; "
         "the biases they rest on match independent manual counts in size on the M25 and "
         "less closely on the M1, but no individual vehicle length has been measured. The "
         "test detects a fault at 660 cm far more reliably than one at 1160 cm, where few "
         "vehicles lie near the boundary and the link to the 520 cm bias is weak. Across "
         "junctions the correction is confirmed on the M25 and only partly on the M1. Only "
         "sites on chains of three or more are corrected, which is why the reduction over "
         "all admitted pairs is smaller than over the pairs touched. Absolute size on short "
         "chains carries the error of Table 10. Biases change, so the fit must be "
         "refreshed. No sensitivity correction has yet been sent to a detector card. The "
         "data cover two corridors, four months and one national detector network. The "
         "measurements of Table 16 close these limits.")
    table(doc, ["#", "Planned measurement", "What it would establish",
                "Minimum useful result", "Claims"], [
        ["1", "Class shares and offsets at a site co-located with a weigh-in-motion or "
              "camera classification station, or after a probe-vehicle calibration",
         "Absolute accuracy of the corrected record and of the offsets in cm",
         "Recorded departure from the reference rises with the fitted bias at a slope "
         "within two standard errors of 1, as at the manual counts (%.2f +- %.2f)"
         % (g520["beta"], g520["beta_se"]), "1, 6"],
        ["2", "Maintenance logs for the sites of Table 13 and the units classed as not an "
              "offset", "That dated events and fault classes match work done at the site",
         "A logged intervention on or next to the day of each dated step", "1, 3"],
        ["3", "Sensitivity correction sent to the detector cards of units with a "
              "length-offset fault",
         "That the retune removes the unit's bias",
         "Equal and opposite steps on the day of the retune; the unit's bias at 520 cm "
         "below three standard errors afterwards", "5"],
        ["4", "Live operation, correction per reporting interval at the archive",
         "Latency and continuity in service",
         "Every record corrected within its 900 s interval; %.1f microseconds measured "
         "offline" % us_rec, "4, 7"],
        ["5", "A second national network, for example classified counts on an Indian "
              "national highway", "Transfer to other detector hardware and fleet mixes",
         "Out-of-sample reduction of adjacent-detector disagreement over all admitted "
         "pairs, as on the M25 and M1 (%.0f-%.0f %%)" % (100 * min(red_a), 100 * max(red_a)),
         "1, 7"],
    ], cap="Table 16: Planned Measurements and Acceptance Criteria",
        widths=[0.3, 1.75, 1.45, 2.3, 0.6], size=8.8, align_left_cols=(1, 2, 3))

    # ============================================================== 9
    h1(doc, "9. Discussion")
    para(doc,
         "A length-classifying detector that reads long keeps its total count and loses "
         "its fleet mix, and nothing in its own output shows it. The art finds such a "
         "detector by attending it with vehicles or apparatus of known length [N5, P2, "
         "P3], by reading its per-vehicle records against an assumed length or a standard "
         "table [N3, N4, P1], or by noticing that its length statistic differs from its "
         "neighbours' and sending someone to look [N1]. An archive of class counts, which "
         "is what most networks keep, supports only the last, and that gives a flag, not a "
         "number.")
    para(doc,
         "The present invention makes the vehicles themselves the reference. Where no slip "
         "road intervenes and totals agree, two consecutive detectors measure the same "
         "vehicles, whose lengths do not change, so their shares below each class boundary "
         "must agree. Solving along a chain gives each detector's signed bias at every "
         "boundary. The stream's own length distribution, known at the boundaries from the "
         "same counts, turns that bias into a length offset in centimetres, and the same "
         "offset predicts, with no fitted parameter, the bias at every other boundary. A "
         "unit that follows the prediction has a length offset, which is sent to it as a "
         "sensitivity correction; a unit that does not has a fault a retune would not "
         "cure. The class counts are corrected with every total kept.")
    para(doc,
         "This differs from each family in the art. Neighbour comparison [N1] flags; "
         "archived-volume correction [N2] corrects totals from a reference station; "
         "stream-based length calibration [P1] matches per-vehicle distributions to a "
         "standard table; weigh-in-motion practice watches one site's own statistic [N11] "
         "or compares it with a second instrument [P4]. None sizes the error at each class "
         "boundary from class counts, recovers a length offset without a standard, tests "
         "whether the error is a length offset, or corrects the class record.")
    para(doc,
         "The tests of Section 8 check that each part earns its place. The balance test and "
         "the free-flow condition each lower the corrected disagreement (8.10); the chain "
         "localises an injected fault in %.0f %% to %.0f %% of cases (8.7); the per-boundary "
         "correction beats the proportional and single-offset forms (8.5, 8.8); the "
         "cross-boundary test separates a 660 cm boundary fault from a length offset "
         "without depending on the chain's reference (8.8); and biases fitted from pairs "
         "alone predict the departure of each detector from independent manual counts "
         "(8.12). The limits are stated in Section 8.16, with the measurements that "
         "would close them."
         % (100 * min(locsize), 100 * max(locsize)))

    # ============================================================== 10
    h1(doc, "10. What Aspect(s) of the Invention Need(s) Protection?")
    bullets(doc, [
        ("1.", "Admission of consecutive roadside detector-site pairs without an "
               "intervening slip road, confirmed by agreement of their total counts."),
        ("2.", "Signed length-class bias per detector site and per class boundary, solved "
               "along a chain of admitted pairs."),
        ("3.", "Length offset in centimetres of each detector unit relative to the "
               "chain's reference, recovered from class counts through the stream's own "
               "length distribution, and a cross-boundary test that the one offset predicts "
               "every other boundary, which classifies the fault."),
        ("4.", "Total-preserving reclassification of each site's archived and incoming "
               "length-class counts by its fitted boundary biases."),
        ("5.", "Sensitivity correction sent to the detector card of each unit with a "
               "length-offset fault, and dating of changes by equal and opposite steps."),
        ("6.", "Data-processing system coupling inductive-loop detector sites, count archive "
               "and maintenance record of offsets and fault classes."),
    ])
    figure(doc, "fig11.png", "Figure 11: Claim tree",
           "One independent method claim with five dependents, and one independent system "
           "claim with one dependent.", width=6.3)
    h2(doc, "A. Method claims")
    CLAIMS_A(doc)
    h2(doc, "B. System claims")
    CLAIMS_B(doc)
    para(doc,
         "Evidence for each claim: Claim 1, Tables 8, 10, 11, 12 and 14, Figures 4 to 7 "
         "and 10; Claim 2, Table 9 and Figure 5(b); Claim 3, Table 13 and Figure 8; Claim "
         "4, Table 8 (the correction of later months is this computation) and the latency "
         "of Table 14; Claim 5, Tables 11 and 13 (the offsets, and the dated steps by which "
         "a change of a unit shows and is verified); Claim 6, Table 11 and Figure 7(c); "
         "Claims 7 and 8, Section 7.", size=10.5)

    # ============================================================== 11
    h1(doc, "11. Technology Readiness Level of the Invention")
    para(doc,
         "The method is implemented in full and validated on archived field records from "
         "two motorway corridors at full corridor scale, each claimed element measured out "
         "of sample, against injected faults, against a control or for physical "
         "consistency, and against independent manual classified counts (Section 8). It "
         "has not run on a live feed, sent a correction to a detector card or been "
         "compared with measured vehicle lengths. Readiness is assessed at TRL 4, "
         "technology validated in the laboratory; measurements 1, 3 and 4 of Table 16 "
         "would support TRL 5.")
    trl_table(doc)
    p_ = doc.add_paragraph()
    p_.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_.paragraph_format.space_before = Pt(18)
    _run(p_, "---------- END OF THE DOCUMENT ----------", size=10.5, bold=True)
    return doc, num


def trl_table(doc):
    trl = doc.add_table(rows=4, cols=9)
    trl.style = "Table Grid"
    trl.alignment = WD_TABLE_ALIGNMENT.CENTER
    for lo, hi, lab in ((0, 2, "Research"), (3, 5, "Development"), (6, 8, "Deployment")):
        m = trl.rows[0].cells[lo].merge(trl.rows[0].cells[hi])
        m.text = ""
        pp = m.paragraphs[0]
        pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _run(pp, lab, size=9.5, bold=True)
        pp.paragraph_format.keep_with_next = True
    names = ["Basic principles observed", "Technology concept formulated",
             "Experimental proof of concept", "Technology validated in lab",
             "Validated in relevant environment", "Demonstrated in relevant environment",
             "Prototype in operational environment", "System complete and qualified",
             "Proven in operational environment"]
    for i in range(9):
        for r, txt, sz, bold in ((1, "TRL %d" % (i + 1), 9, True), (2, names[i], 7.8, False),
                                 (3, "✓" if i < 4 else "", 12, True)):
            c = trl.rows[r].cells[i]
            c.text = ""
            pp = c.paragraphs[0]
            pp.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _run(pp, txt, size=sz, bold=bold)
            c.width = Inches(0.71)
            if r < 3:
                pp.paragraph_format.keep_with_next = True
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    _run(p, "Current TRL: Level 4 — Technology Validated in Laboratory.", size=11,
         bold=True)


# ------------------------------------------------------------------ claims
def CLAIMS_A(doc):
    claim(doc, 1,
          "A method of correcting the length classification of a plurality of roadside "
          "vehicle detectors (110) disposed in sequence along a carriageway, carried out by "
          "a data-processing system (900) connected to the detectors by a communication "
          "link (130), each detector measuring the length of each vehicle passing it and "
          "reporting, for each of a succession of measurement intervals, the number of "
          "vehicles in each of at least three length classes separated by length-class "
          "boundaries, the method comprising: "
          "(a) receiving, from the detectors over the communication link (130), their "
          "per-interval length-class counts, together with information identifying the "
          "order of the detectors along the carriageway and the consecutive detectors "
          "between which a slip road (120) joins or leaves the carriageway; "
          "(b) selecting each pair of consecutive detectors between which no slip road "
          "joins or leaves the carriageway, and admitting the pair for a fitting period "
          "only if, over the intervals of that period in which both detectors reported, "
          "their total counts agree within a balance tolerance, whereby the two detectors "
          "of an admitted pair measure the same vehicles, whose lengths do not change "
          "between them; "
          "(c) for each admitted pair and each length-class boundary, computing over those "
          "intervals the difference between the shares of vehicles that the two detectors "
          "report below the boundary; "
          "(d) for each chain of at least two admitted pairs in which consecutive pairs "
          "share a detector, determining for every detector of the chain, separately for "
          "each boundary, a signed share bias such that for each pair of the chain the "
          "bias of its upstream detector less the bias of its downstream detector equals "
          "the share difference of that pair, the constant common to the chain being "
          "fixed by setting to zero either the median of the biases over the chain or "
          "the bias of a detector on the chain whose length measurement has been "
          "verified, which median or detector is the reference of the chain; "
          "(e) estimating a cumulative length distribution of the vehicles by "
          "interpolating, through the length-class boundaries, the shares of vehicles "
          "that the detectors report below those boundaries; "
          "(f) for each detector of the chain, determining a length offset, being the "
          "distance by which the detector over-measures each vehicle length relative to "
          "the reference of the chain, such that the change in the estimated cumulative "
          "distribution at a first boundary when that boundary is reduced by the offset "
          "equals the detector's share bias at that boundary; predicting, at each other "
          "boundary, the share bias that the same offset would produce, as the detector's "
          "share bias at the first boundary multiplied by the ratio of the slope of the "
          "estimated cumulative distribution at that boundary to its slope at the first "
          "boundary; and classifying the detector as having a length-offset fault where "
          "its share biases at the other boundaries depart from the predicted share "
          "biases by no more than a set fraction of the larger of its share bias at the "
          "first boundary and a floor, and as having a fault of another kind otherwise; "
          "(g) generating a corrected length-class count dataset (810) for measurement "
          "intervals after the fitting period by reclassifying, in each interval of each "
          "detector of the chain, a number of vehicles across each boundary equal to that "
          "detector's share bias at the boundary multiplied by the total count of the "
          "interval, the total count of every interval being left unchanged; and "
          "(h) storing the corrected dataset in a count archive (200) in place of the "
          "received length-class counts, and writing to a maintenance record (830), for "
          "each detector of the chain, its identity and fault class and, where the fault "
          "class is a length-offset fault, its length offset as a sensitivity correction "
          "for that detector unit.", independent=True)
    claim(doc, 2,
          "The method of claim 1, wherein steps (b) to (g) are repeated for successive "
          "fitting periods, the biases, offsets and fault classes determined over each "
          "fitting period being applied to the measurement intervals of the period that "
          "follows it.")
    claim(doc, 3,
          "The method of claim 1, further comprising, for each detector belonging to two "
          "admitted pairs: forming for each of the two pairs the daily series of its share "
          "difference below a boundary; locating in each series the split of largest "
          "difference of means relative to the pooled standard deviation; and, where the "
          "two splits fall within one day of each other, each step exceeds a minimum size, "
          "they are of opposite sign and their sum is less than a set fraction of the "
          "larger, writing to the maintenance record a change in the length measurement of "
          "that detector dated to the day of the steps, classifying the change as a "
          "change of length offset or a fault of another kind by applying the prediction "
          "and classification of step (f) to the change, across that day, in the "
          "detector's share biases at every boundary, and determining that detector's "
          "biases, offset and fault class from measurement intervals after that day.")
    claim(doc, 4,
          "The method of claim 1, wherein step (g) is applied to the length-class counts "
          "received from each detector after the fitting period as they are received, and "
          "the corrected counts are supplied to a downstream traffic-management process "
          "(820) in place of the received counts.")
    claim(doc, 5,
          "The method of claim 1, further comprising transmitting over the communication "
          "link (130), to the detector card (111) of each detector classified as having a "
          "length-offset fault, a sensitivity correction by which the detector card "
          "shortens each vehicle length it computes thereafter by the length offset of "
          "that detector, and determining the biases, offset and fault class of that "
          "detector afresh from the measurement intervals reported after the correction "
          "takes effect.")
    claim(doc, 6,
          "The method of claim 1, further comprising determining for each detector of the "
          "chain a speed bias, from the ratios of the mean speeds reported by the two "
          "detectors of each admitted pair solved along the chain as in step (d), and "
          "writing to the maintenance record, for each detector classified as having a "
          "length-offset fault, the part of its length offset attributable to a "
          "loop-spacing error, computed from the speed bias, and the remainder as a "
          "detection-threshold error.")


def CLAIMS_B(doc):
    claim(doc, 7,
          "A data-processing system (900) for correcting the length classification of "
          "roadside vehicle detectors (110) disposed in sequence along a carriageway, "
          "comprising: a count archive (200) receiving over a communication link (130) "
          "the per-interval length-class counts of the detectors; a pair-admission unit "
          "(300) configured to carry out step (b) of claim 1; a share-difference unit "
          "(310) configured to carry out step (c) of claim 1; a chain solver (320) "
          "configured to carry out step (d) of claim 1; an offset estimator (350) "
          "configured to carry out step (e) and the determination of the length offset "
          "of step (f) of claim 1; a fault classifier (360) configured to carry out the "
          "prediction and classification of step (f) of claim 1; a reclassification unit "
          "(400) configured to carry out step (g) of claim 1 and to write the corrected "
          "dataset to the count archive; and a maintenance record (830) holding the fault "
          "class and the length offset of each detector unit.", independent=True)
    claim(doc, 8,
          "The system of claim 7, further comprising the detectors, each detector "
          "comprising a pair of inductive loops in each lane of the carriageway and "
          "computing the length of each vehicle from the times at which the vehicle "
          "occupies the two loops.")


# ------------------------------------------------------------------ checks
def check(path, num):
    d = Document(path)
    paras = [p.text for p in d.paragraphs]
    tables = []
    for i, t in enumerate(d.tables):
        cells = [c.text for row in t.rows for c in row.cells]
        tables.append((i, cells))
    body = "\n".join(paras)
    # Tables 1 and 2 (prior art) quote titles of references, which may contain words banned
    # in the applicant's own prose; it is checked for business vocabulary only.
    prior = [c for i, cells in tables if cells[:1] == ["S.No."] and "Comments" in cells[:4]
             for c in cells]
    other = [c for i, cells in tables for c in cells if c not in prior]
    own = body + "\n" + "\n".join(other)
    everything = own + "\n" + "\n".join(prior)
    bad = [w for w in FORBIDDEN if re.search(w, everything, re.I)]
    ban = [w for w in BANNED if re.search(w, own, re.I)]
    ph = PLACEHOLDER.findall(everything)
    open_ = [p[:40] for p in paras if OPENERS.match(p)]
    phys = sum(1 for t in num.texts if PHYSICAL.search(t))
    share = phys / len(num.texts)
    problems = []
    if bad:
        problems.append("business vocabulary %s" % bad)
    if ban:
        problems.append("banned phrases %s" % ban)
    if ph:
        problems.append("placeholders %s" % ph)
    if open_:
        problems.append("banned paragraph openers %s" % open_)
    if share <= 0.6:
        problems.append("physical tie-in only %.0f %% of Sections 6-7" % (100 * share))
    if problems:
        raise SystemExit("refused: " + "; ".join(problems))
    print("checked: no business vocabulary, banned phrases, placeholders or banned "
          "openers; %d numbered paragraphs, %.0f %% name a physical element; %d words"
          % (len(num.texts), 100 * share, len(everything.split())))


if __name__ == "__main__":
    d, num = build()
    d.save(OUT)
    print("wrote", OUT)
    check(OUT, num)
