"""Build the IDF-B disclosure for stage-offset estimation at river gauging stations.

Layout: the accepted VIT IDF-B (masthead, eleven numbered sections with the subsections
of the accepted cyber-security IDF-B), through the layout helpers of the CausalInterlock
generator and the table/figure helpers of the patent 3 generator, imported unchanged.

Every measured number is read from results/*.json at build time and the operating
constants are read from the code that used them. The build refuses to run if a results
file is missing. After writing, it scans the document for placeholders, business-method
vocabulary, banned phrases and checks that more than 60 % of the numbered paragraphs
name a physical element.
"""
import importlib.util
import json
import math
import os
import re
import sys

from docx import Document
from docx.shared import Inches, Pt

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
RESULTS = os.path.join(ROOT, "results")
SRC = os.path.join(ROOT, "src")


def _load_module(name, path):
    spec = importlib.util.spec_from_file_location(name, path)
    m = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(m)
    return m


sys.path.insert(0, r"C:\adi\patents\01_02_infosec\IDF-B_CausalInterlock")
import build_idfb as B                                            # noqa: E402

P3 = _load_module("p3_make_idfb", r"C:\adi\patents\05_length_invariance\src\make_idfb.py")
sys.path.insert(0, SRC)
NET = _load_module("p4_network", os.path.join(SRC, "network.py"))
PD = _load_module("p4_probe_daily", os.path.join(SRC, "probe_daily.py"))
EX = _load_module("p4_experiments", os.path.join(SRC, "experiments.py"))
XTRA = _load_module("p4_idfb_extra", os.path.join(SRC, "idfb_extra.py"))

B.FIG = os.path.join(ROOT, "patent", "figures")
OUT = os.path.join(ROOT, "patent", "IDF-B_Stage_Offset_Neighbour_Gauge.docx")
h1, h2, para, bullets = B.h1, B.h2, B.para, B.bullets
table, figure, claim, Numbering, trl_table = P3.table, P3.figure, P3.claim, P3.Numbering, P3.trl_table

FORBIDDEN = P3.FORBIDDEN
BANNED = [b for b in P3.BANNED if b not in ("python",)] + ["python"]
OPENERS, PLACEHOLDER = P3.OPENERS, P3.PLACEHOLDER
PHYSICAL = re.compile(r"station|sensor|logger|gaug|river|channel|control \(118\)|link \(130\)|"
                      r"stage|transducer|bubbler|radar|field visit|archive", re.I)


def load(name):
    path = os.path.join(RESULTS, name)
    if not os.path.exists(path):
        raise SystemExit("results absent, refusing to build with invented figures: %s" % name)
    with open(path, encoding="utf-8") as fh:
        return json.load(fh)


FT = 0.3048           # metres per foot: USGS records are in feet, the disclosure in metres


def m3(ft):
    """A length given in feet, written in metres to three decimals."""
    return "%.3f m" % (FT * ft)


def f1(x):
    return "%.1f" % x


def f2(x):
    return "%.2f" % x


def mmp(b):
    """'median / mean / p90' of a summary block."""
    return "%s / %s / %s" % (f1(b["median"]), f1(b["mean"]), f1(b["p90"]))


def build():
    pairs = load("pairs.json")
    coup = load("coupling.json")
    daily = load("daily.json")
    e1 = load("e1_interpolation.json")
    e2 = load("e2_realtime.json")
    e3 = load("e3_injection.json")
    e4 = load("e4_ablation.json")
    e5 = load("e5_detection.json")
    e6 = load("e6_timing.json")
    e7 = load("e7_iv.json")
    e8 = load("e8_india.json")
    e9 = load("e9_bootstrap.json")
    ex = load("fig05_example.json")
    IV_OK = e7.get("n_common", 0) >= 20
    meta = e1["meta"]
    MC = meta["min_coupling"]
    n_outlets = len(NET.OUTLETS)
    n_gauges = pairs["n_mainstem_gauges"]
    n_pairs = len(pairs["pairs"])
    n_dir_c = sum(1 for r in coup if r[5] >= MC)
    n_scored = sum(1 for r in daily["results"] if r.get("rows") and r.get("coupling", 0) >= MC)
    gaps = sorted(w["gap_days"] for r in daily["results"] if r.get("rows") and r.get("coupling", 0) >= MC
                  for w in r["rows"])
    gap_med = gaps[len(gaps) // 2]
    pub, rtm = e1["published"], e1["realtime"]
    ra, rb, ro, rob = pub["all"], pub["big"], pub["open_water"], pub["open_water_big"]
    t2a, t2b, t2o = e2["realtime"]["all"], e2["realtime"]["big"], e2["realtime"]["open_water"]
    r2a, r2b = e2["published"]["all"], e2["published"]["big"]
    WIN, STEADY, MIND = meta["win_d"], meta["steady_ft"], meta["min_days"]
    INCR = NET.MAX_INCR
    GMIN, GPRI, GMAX = EX.GAIN_MIN_PAST, EX.GAIN_PRIOR, EX.GAIN_MAX
    BIG = meta["big_ft"]
    BIGM = m3(BIG)
    PR = e4["all"]["vs_practice_per_station"]          # method against practice chosen per station
    PRo = e4["open_water"]["vs_practice_per_station"]
    ipa, ipo = e1["interp"]["all"], e1["interp"]["open_water"]
    b9 = e9
    ho = e9["holdout"]
    var = e9["variants"]
    vr_all = [v["all"][1] for v in var]
    vz = [v["all_paired"]["z"] for v in var]
    slope_terc = e4["by_slope_difference"]
    # the direction of each comparison stated in words below; refuse to build if the
    # numbers no longer say so
    claims_in_words = {
        "method beats the strongest baseline in all, open water and big":
            all(e4[k]["stage_units_gain"]["mean"] < min(e4[k][b]["mean"] for b in
                                                         ("linear_time", "flood_step", "flood_ramp", "per_rise", "practice_per_station"))
                for k in ("all", "open_water", "big")),
        "flood baselines worse than linear over all visits, step better where control moved":
            ra["baselines"]["e_step"]["mean"] > ra["baseline"]["mean"] and ra["baselines"]["e_ramp"]["mean"] > ra["baseline"]["mean"]
            and rb["baselines"]["e_step"]["mean"] < rb["baseline"]["mean"],
        "real time: mean falls, median rises": t2a["method"]["mean"] < t2a["baseline"]["mean"]
            and t2a["method"]["median"] > t2a["baseline"]["median"],
        "real-time reference record: not over all visits": rtm["all"]["paired"]["z"] < 1.96,
        "finalised-to-date neighbour: mean lower, not significantly more often":
            e2["rtfinal"]["all"]["method"]["mean"] < e2["rtfinal"]["all"]["baseline"]["mean"]
            and e2["rtfinal"]["all"]["paired"]["z"] < 1.96,
        "ratio step: term smaller movement more often": e4["ratio_step_m"]["paired"]["z"] > 1.96
            and e4["ratio_step_m"]["share_over_big_with"] <= e4["ratio_step_m"]["share_over_big_without"],
        "g = 1 lower where control moved, gain lower over all":
            e4["big"]["stage_units_no_gain"]["mean"] < e4["big"]["stage_units_gain"]["mean"]
            and e4["all"]["stage_units_gain"]["mean"] < e4["all"]["stage_units_no_gain"]["mean"],
        "stage units beat ratio units in each third": all(t["paired"]["z"] > 1.96 for t in slope_terc),
        "every fit below linear and practice in every column":
            all(e4[k][fv]["mean"] < min(e4[k]["linear_time"]["mean"], e4[k]["practice_per_station"]["mean"])
                for k in ("all", "open_water", "big") for fv in ("fit_median", "fit_weighted_median", "stage_units_gain")),
    }
    wrong = [k for k, ok in claims_in_words.items() if not ok]
    if wrong:
        raise SystemExit("text no longer matches the results: %s" % wrong)

    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = B.BODY_FONT
    st.font.size = Pt(11)
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Inches(0.9)
    sec.top_margin = sec.bottom_margin = Inches(0.85)
    B.running_heads(sec)
    B.letterhead(doc)
    num = Numbering()

    # ============================================================== 1
    h1(doc, "1. Title of the Invention")
    para(doc, TITLE + ".")

    # ============================================================== 2
    h1(doc, "2. Field / Area of Invention")
    para(doc,
         "The present invention relates to river gauging stations, in which a stage sensor "
         "(a submerged pressure transducer, a gas-purge bubbler or a non-contact radar gauge) "
         "records the water level above a natural or built channel control and a data logger "
         "converts that stage to discharge through a stage-discharge rating. More particularly, "
         "it pertains to a system and method that determine, between the field visits at "
         "which a hydrographer measures discharge, the stage offset by which the true "
         "stage-discharge relation of the station has moved away from its base rating, using "
         "the discharge recorded at a neighbouring gauging station on the same river, and that "
         "correct the station's discharge record and the offset applied by its data logger.")
    para(doc,
         "A rating is a fixed curve; the channel that it describes is not. Bed scour and fill "
         "after floods, weed growth, debris on a riffle and drift of the stage sensor all "
         "move the true relation to Q = R(h + s): the discharge the station should report is "
         "the base rating read at the recorded stage h plus an offset s. The offset is "
         "measured only when a hydrographer visits the station, measures discharge with a "
         "current meter or acoustic Doppler profiler and reads the reference gauge; at the "
         "stations evaluated here the median interval between visits was %d days. Between "
         "visits a station in real time keeps the last measured offset until a new gauging "
         "updates it [N2]. For the published record the offset is interpolated between the "
         "visits: linearly in time for each stage value (Sauer [N1], section 8.4.4), or by the "
         "hydrographer's judgement with arbitrary points, for instance a step at a rise for a "
         "gravel control (Kennedy [N5]); and the same manual states that records of other "
         "stations are normally not used to distribute the offsets, so that each record stays "
         "free of the influence of other records." % gap_med)
    para(doc,
         "The resulting errors are large. At %d field visits of %d gauging stations, each left "
         "out in turn and predicted from the visits before and after it, linear interpolation "
         "gives a mean absolute discharge error of %s %% against the measured discharge (median "
         "%s %%, 90th percentile %s %%); at the %d visits where the offset had moved by %s or "
         "more from the interpolated value the mean error is %s %%, and placing the change at "
         "the largest flood between the visits, as a hydrographer would for a gravel control, "
         "still leaves %s %%. The technical problem addressed is to determine the offset of a "
         "gauging station between field visits from records that already exist, without new "
         "apparatus at the station, so that the discharge record and the discharge computed by "
         "the station itself follow the channel control as it changes."
         % (pub["n"], pub["targets"], f1(ra["baseline"]["mean"]), f1(ra["baseline"]["median"]),
            f1(ra["baseline"]["p90"]), pub["n_big"], BIGM, f1(rb["baseline"]["mean"]),
            f1(rb["baselines"]["e_step"]["mean"])))

    h2(doc, "2.1 Technical Problem Addressed by the Invention")
    bullets(doc, [
        "Between field visits a gauging station has no measurement of its stage offset; the "
        "offset moves when the channel control (118) changes, which happens mostly at "
        "floods, between visits.",
        "Interpolating the offset between visits, linearly in time or anchored to a flood, "
        "assumes when and how the control changed; where the assumption is wrong the "
        "discharge between the visits is wrong.",
        "A neighbouring station on the same river sees the same water, but its discharge "
        "cannot be copied into the record: the flow ratio between the stations changes with "
        "inflow, and a discrepancy between the stations mixes that ratio with the station's "
        "offset.",
        "A stage offset changes discharge by a share that depends strongly on stage (by the "
        "rating log-slope), so a discrepancy in discharge measured at one stage does not "
        "transfer to another stage, and is not the quantity that the station's data logger "
        "applies.",
        "Existing comparison with neighbouring stations is a visual or monthly regression "
        "check that questions a rating; it does not size the offset in the units of the "
        "rating or correct the record."])

    h2(doc, "2.2 Definitions")
    bullets(doc, [
        "Gauging station (100): a stage sensor (110) above a channel control (118), with a "
        "data logger (115) that computes discharge from stage and transmits both over a "
        "telemetry link (130).",
        "Neighbouring station (120): a gauging station on the same river, upstream or "
        "downstream, with no tributary, abstraction or storage between it and the station "
        "large enough to break the admission test.",
        "Base rating R: the stage-discharge relation of the station without offset; R(g) is the "
        "discharge at effective stage g.",
        "Stage offset s: the amount, in metres, added to the recorded stage before the base "
        "rating is read, so that the discharge is R(h + s). At a field visit the measured "
        "offset is R^-1(Q_meas) - GH_meas.",
        "Log discrepancy y: ln Q_Y(t - tau) - ln R(h(t)), the logarithmic difference between "
        "the neighbour's discharge, delayed by the travel time tau, and the base-rating "
        "discharge of the station.",
        "Rating log-slope x: d ln R / dh at the recorded stage; a stage offset s changes the "
        "log discrepancy by s times x.",
        "Flow ratio c: ln(Q_Y / Q), the part of y that comes from inflow or abstraction "
        "between the stations, so that y = c + s x.",
        "Window offset s_raw: for a window of steady intervals and a given flow ratio c, the "
        "least-squares offset sum(x (y - c)) / sum(x^2).",
        "Departure: the difference between s_raw at an intermediate time and the offset "
        "interpolated between the two field visits."])

    # ============================================================== 3
    h1(doc, "3. Prior Patents and Publications from Literature")
    para(doc,
         "Each reference below was checked against its own text before any statement about it "
         "was made: the full text of each publication, and the abstract and first claim of "
         "each patent as published.")
    h2(doc, "A. Prior Patent Literature")
    PAT = [
        ("P1", "US 11,238,356 B2", "Method of predicting streamflow data",
         "Generates synthetic streamflow for a target streamgage during an interruption of its "
         "measured data, by clustering streamflow sources on their correlations and using the "
         "groups to predict the target's flow (claim 1). It replaces the discharge; it does "
         "not estimate a stage offset of the target's rating and is not anchored to "
         "field-visit measurements."),
        ("P2", "CN 108460814 A", "Multi-station linkage stage-discharge relation fitting method",
         "Expresses the forecast stage and discharge of a test station, at a lead time not "
         "longer than the flood-wave travel time, as a function of the stage and discharge of "
         "upstream stations, for flood forecasting. It forecasts; it does not determine the "
         "offset of the station's rating from field visits or correct its record."),
        ("P3", "CN 104729605 A", "Stream flow monitoring system and water-level flow calibration method",
         "Builds the stage-discharge calibration curve of one site from repeated instantaneous "
         "flow monitoring at different water levels. Single station; no neighbouring station; "
         "no offset between visits."),
        ("P4", "US 12,529,586 B2", "Measurement system for streamflow determination in real time",
         "Estimates streamflow from a directly measured index velocity and a continuous "
         "slope-area method using stage at two points, without a conventional rating. "
         "Requires velocity instrumentation; does not estimate a rating offset."),
        ("P5", "CN 103106248 A", "Stage-discharge data assimilation flood reporting method, device and system",
         "Receives the stage reported by a measuring station, selects a stage-discharge "
         "conversion model according to the station's configuration, computes and reports "
         "discharge. Conversion only; the relation is not corrected between gaugings and no "
         "other station is used."),
        ("P6", "CN 111753965 A", "Deep-learning method and system for automatic compilation of river discharge",
         "Computes and compiles a station's discharge with a trained deep-learning model whose "
         "inputs are the station's own water level, the water levels of the upstream and "
         "downstream stations, rainfall and evaporation (claim 1). It replaces the rating by a "
         "learned model of water levels; it does not use the neighbour's discharge, does not "
         "determine a stage offset of the rating and is not anchored to the field visits "
         "between which the record is computed."),
    ]
    table(doc, ["S.No.", "Patent No.", "Title", "Comments"],
          [[a, b, c, d] for a, b, c, d in PAT],
          "Table 1: Prior patent literature",
          widths=[0.5, 1.25, 1.6, 3.35], size=9, align_left_cols=(2, 3))
    h2(doc, "B. Non-Patent Literature")
    NPL = [
        ("N1", "Sauer (2002), USGS Water-Resources Investigations Report 01-4044",
         "Standards for the analysis and processing of surface-water data and information using electronic methods",
         "Shifts are defined from discharge measurements and, between shift curves, "
         "interpolated by an unweighted linear time interpolation for each unit value "
         "(section 8.4.4). Records of nearby reference stations are used to estimate discharge "
         "for periods of missing or ice-affected record (section 13.1), not to determine "
         "shifts."),
        ("N2", "Darienzo (2021), PhD thesis, Univ. Grenoble Alpes (tel-03211343)",
         "Detection and estimation of stage-discharge rating shifts for retrospective and real-time streamflow quantification",
         "Detects shifts by segmentation of gauging residuals, from stage recessions and from a "
         "sediment-transport proxy. Reports that operational services examine regressions with "
         "upstream or downstream stations, on a monthly step, to question a rating, and that a "
         "new gauging is needed to update the curve in real time. No conversion of a "
         "neighbour discrepancy into a stage offset and no correction of the record from it."),
        ("N3", "Mansanarez et al. (2019), Water Resources Research 55 (hal-02609391)",
         "Shift happens! Adjusting stage-discharge rating curves to morphological changes at known times",
         "Stage-period-discharge model: hydraulic rating equations with parameters that change "
         "at known dates, estimated from all gaugings by Bayesian inference, replacing manual "
         "shift corrections. Uses gaugings only; the change dates must be known; no "
         "neighbouring station."),
        ("N4", "Puechberty et al. (2017), French Ministry of Environment (hal-02606950)",
         "Charte qualite de l'hydrometrie: guide de bonnes pratiques",
         "Hydrometry quality charter. Rating changes are suspected from the hydrograph, flood "
         "occurrence or works, and confirmed by a gauging (5.3); hydrographs of upstream and "
         "downstream stations are superposed, a gap after a flood prompts questioning both "
         "ratings, and transfer time may reveal sensor or rating drift (6.2.2); a monthly "
         "linear correlation models one station from another (6.2.3). Visual or monthly "
         "checks; no offset sized in stage units, no anchoring to visits, no correction "
         "applied to the station."),
        ("N5", "Kennedy (1983), USGS Techniques of Water-Resources Investigations 3-A13",
         "Computation of continuous records of streamflow",
         "Shift adjustments between measurements are distributed by the hydrographer's "
         "judgement: by inspection, by shift-variation diagrams in time or stage, with "
         "arbitrary data points (for a gravel control that changes only during rises, a step "
         "at the rise), or by varying the shifts until a smooth daily hydrograph is obtained. "
         "States that a record should be free of the influence of other streamflow records, "
         "so the use of records of other stations to distribute shift adjustments is normally "
         "avoided."),
        ("N6", "U.S. Geological Survey (2019), training presentation (title and date from the file's metadata)",
         "Stage-discharge relations: basic concepts",
         "Shifts are often prorated to a positive shift on a rise until soon after the peak, "
         "and to a negative shift on a recession as sediment falls out. Hydrograph-guided "
         "proration from the station's own record; no neighbouring station."),
        ("N7", "Vlah et al. (2024), Hydrology and Earth System Sciences 28, 545",
         "Leveraging gauge networks and strategic discharge measurements to aid the development of continuous streamflow records",
         "Reconstructs the continuous discharge of a target site (\"virtual gauge\") by linear "
         "regression of its field-measured discharges on the discharge of nearby donor gauges, "
         "or by neural networks; the target's own water level is deliberately not used because "
         "of sensor drift. The regression is fitted at the target's field visits, so the "
         "relation between donor and target is anchored to measurements there; but its output "
         "replaces the target's discharge. It does not keep the target's stage record and "
         "rating, does not separate a stage offset from the flow ratio by their different "
         "dependence on stage, and does not size any correction in units of stage."),
        ("N8", "Rantz et al. (1982), USGS Water-Supply Paper 2175, vol. 2",
         "Measurement and computation of streamflow: computation of discharge",
         "Shifts are prorated with time, or treated as discrete shifts at peak flows and "
         "prorated between rises (pp. 349, 354). Hydrographs of nearby stations are compared "
         "as a test for consistency of the computed record and used to estimate discharge for "
         "days without gauge-height record or with ice (pp. 572-576). The comparison checks or "
         "replaces the record; it does not size a stage offset."),
    ]
    table(doc, ["S.No.", "Reference", "Title", "Comments"],
          [[a, b, c, d] for a, b, c, d in NPL],
          "Table 2: Non-patent literature",
          widths=[0.5, 1.45, 1.6, 3.15], size=9, align_left_cols=(1, 2, 3))

    # ============================================================== 4
    h1(doc, "4. Summary and Background of the Invention")
    h2(doc, "4.1 Background")
    para(doc,
         "Discharge records of rivers drive flood warning, reservoir release, intake and "
         "outfall control and the design of hydraulic structures. Almost all of them are "
         "computed, not measured: a stage sensor (110) records water level every 15 minutes, "
         "and a data logger (115) or a central system converts it with a rating. The rating "
         "is checked against a measured discharge at each field visit, and the difference, "
         "expressed as a stage offset, is applied from then on. Sand-bed and gravel rivers "
         "change their controls at floods. At the stations evaluated here the offset "
         "measured at a visit differed from that of the visit before by %s or more at %d of "
         "%d visits, and an offset of %s moves discharge by tens of percent at low stage."
         % (BIGM, e2["published"]["n_big"], e2["published"]["n"], BIGM))
    para(doc,
         "Two remedies exist. A field visit measures the offset exactly but only at one "
         "instant; between visits the offset is interpolated in time or by judgement from the "
         "station's own hydrograph [N1, N5, N6]. A neighbouring station sees the same water, "
         "and hydrographers use its hydrograph to fill gaps or estimate ice periods [N1] or to "
         "question a rating [N2, N4], while the offsets themselves are normally distributed "
         "without it [N5]; P1 builds synthetic discharge from correlated stations. None of "
         "these converts the neighbour's information into the quantity the station's rating "
         "uses, the stage offset, and none separates the flow ratio between the stations from "
         "the offset by anchoring it to the station's own field visits.")
    h2(doc, "4.2 Distinctions from Prior Art")
    bullets(doc, [
        "From N1, N5, N6 and N8: these distribute the offset between visits from the station's own "
        "data, linearly in time, by judgement or tied to the hydrograph of the station; N5 "
        "normally avoids other stations' records for this. The invention takes the timing and "
        "size of the departure from a neighbour, after removing the flow ratio fixed at the "
        "station's own visits, and on the same visits beats each of these practices and the "
        "practice chosen per station from its own history (Section 8.3).",
        "From N2 and N3: N2 detects that a shift may have occurred, from gaugings or monthly "
        "regressions, and N3 re-estimates the rating from gaugings once the dates of change "
        "are known; the invention sizes the offset in metres between gaugings, from a "
        "neighbour, without knowing when the change happened, and writes it into the record "
        "and the data logger.",
        "From N4: N4 superposes upstream and downstream hydrographs, or fits a monthly linear "
        "relation between stations, to prompt questioning of the ratings; the invention turns "
        "the comparison into a daily offset in stage units anchored to the visits.",
        "From P1: P1 substitutes the target's discharge with a prediction from other gauges; "
        "the invention keeps the target's own stage record and rating and corrects only the "
        "offset, so the station's own stage dynamics are preserved.",
        "From P2: P2 forecasts the target from upstream stations; the invention estimates the "
        "present state of the target's rating and is anchored by measured offsets.",
        "From P6 and N7: these compute the discharge of a station from other stations, by a "
        "learned model of water levels (P6) or by regression of the station's gaugings on a "
        "donor's discharge (N7). N7 thereby anchors a donor-target relation at the field "
        "visits, but uses it to replace the target's discharge and sets the target's stage "
        "record aside. The invention keeps the station's own stage record and rating; it "
        "separates the flow ratio from the stage offset by their different dependence on "
        "stage (the offset acts through the rating log-slope, the flow ratio equally at every "
        "stage), uses the visits to fix the flow ratio from the offset measured there, and "
        "returns an offset in metres that the station's data logger applies to its own "
        "stage.",
        "From P3, P4 and P5: single-station calibration, replacement of the rating by velocity "
        "instrumentation, or plain conversion of stage to discharge; the invention needs no "
        "new instrument and corrects the conversion between gaugings."])
    h2(doc, "4.3 Inventive Combination")
    para(doc,
         "The combination is: (i) the log discrepancy y between the neighbour's lagged "
         "discharge and the station's base-rating discharge, in which a stage offset s and the "
         "flow ratio c appear as y = c + s x; (ii) the flow ratio fixed at each field visit as "
         "the value for which the window next to the visit returns the offset measured there, "
         "and interpolated between the visits; (iii) the offset of each window with that flow "
         "ratio removed, fitted on the rating log-slope (least squares in the example), in metres; (iv) a weight on "
         "its departure from the interpolated offset, in the example a gain learned from the "
         "station's own earlier visits; and (v) the resulting offset written into the record and into the station's "
         "data logger. With the least-squares fit, each element carries measured weight on "
         "the same %d visits (Table 10): "
         "without the flow-ratio term the error is higher at %d of %d visits where the two "
         "differ (z = %s); in discharge-ratio units instead of stage units, higher at %d of %d "
         "(z = %s); with the gain fixed at 1, the mean error rises from %s %% to %s %%."
         % (pub["n"], e4["all"]["with_vs_without_ratio_term"]["better"],
            e4["all"]["with_vs_without_ratio_term"]["of"], f2(e4["all"]["with_vs_without_ratio_term"]["z"]),
            e4["all"]["stage_vs_ratio"]["better"], e4["all"]["stage_vs_ratio"]["of"],
            f2(e4["all"]["stage_vs_ratio"]["z"]), f1(e4["all"]["stage_units_gain"]["mean"]),
            f1(e4["all"]["stage_units_no_gain"]["mean"])))
    h2(doc, "4.4 Feature Mapping Against Prior Art")
    FEAT = [
        ("F1", "Station's own recorded stage and base rating used between visits", "No", "Partial", "Yes", "Partial", "Yes", "No"),
        ("F2", "Discharge of a neighbouring station on the same river", "Yes", "Yes", "Partial", "Yes", "No", "Yes"),
        ("F3", "Travel-time lag and coupling test admitting the pair", "Partial", "Partial", "No", "Partial", "No", "Partial"),
        ("F4", "Flow ratio fixed at each field visit from the measured offset", "No", "No", "No", "No", "No", "Partial"),
        ("F5", "Discrepancy converted to an offset in metres through the rating log-slope", "No", "No", "No", "No", "No", "No"),
        ("F6", "Gain learned from the station's earlier visits", "No", "No", "No", "No", "No", "No"),
        ("F7", "Offset written to the record and to the station's data logger", "No", "No", "Partial", "No", "Partial", "No"),
    ]
    table(doc, ["#", "Feature", "P1", "P2", "N1", "N4", "N5", "N7"],
          [list(r) for r in FEAT], "Table 3: Feature mapping of the invention against the nearest art",
          widths=[0.4, 2.85, 0.58, 0.58, 0.58, 0.58, 0.58, 0.58], size=9, align_left_cols=(1,))
    h2(doc, "4.5 Technical Effects")
    bullets(doc, [
        "Discharge record between field visits, over all visits: at %d field visits of %d "
        "stations, each left out and scored against the hydrographer's measured discharge, "
        "the mean absolute error of the record falls from %s %% (linear interpolation) to %s %% "
        "and the 90th percentile from %s %% to %s %%; the error is smaller at %d of the %d "
        "visits where the two differ (sign test z = %s). Resampling whole stations, the "
        "reduction is %s points (95 %% interval %s to %s)."
        % (pub["n"], pub["targets"], f1(ra["baseline"]["mean"]), f1(ra["method"]["mean"]),
           f1(ra["baseline"]["p90"]), f1(ra["method"]["p90"]), ra["paired"]["better"],
           ra["paired"]["of"], f2(ra["paired"]["z"]), f1(b9["e1_all"]["reduction_pts"]),
           f1(b9["e1_all"]["ci95"][0]), f1(b9["e1_all"]["ci95"][1])),
        "In open water (April to November, no ice at the visit): mean %s %% to %s %%, better "
        "at %d of %d (z = %s), reduction %s points (%s to %s)."
        % (f1(ro["baseline"]["mean"]), f1(ro["method"]["mean"]), ro["paired"]["better"],
           ro["paired"]["of"], f2(ro["paired"]["z"]), f1(b9["e1_open_water"]["reduction_pts"]),
           f1(b9["e1_open_water"]["ci95"][0]), f1(b9["e1_open_water"]["ci95"][1])),
        "Against hydrographer practice chosen per station: for each visit, the practice "
        "(linear in time, the change at the largest flood, prorated to it, or split over the "
        "rises) with the lowest error at the station's earlier visits gives a mean of %s %% "
        "over all visits; the method gives %s %%, better at %d of %d visits (z = %s), "
        "reduction %s points (%s to %s); in open water, z = %s, reduction %s points (%s to %s)."
        % (f1(e4["all"]["practice_per_station"]["mean"]), f1(e4["all"]["stage_units_gain"]["mean"]),
           PR["better"], PR["of"], f2(PR["z"]), f1(b9["e1_vs_practice"]["reduction_pts"]),
           f1(b9["e1_vs_practice"]["ci95"][0]), f1(b9["e1_vs_practice"]["ci95"][1]), f2(PRo["z"]),
           f1(b9["e1_open_water_vs_practice"]["reduction_pts"]),
           f1(b9["e1_open_water_vs_practice"]["ci95"][0]), f1(b9["e1_open_water_vs_practice"]["ci95"][1])),
        "With the neighbour's record rebuilt from its own stage, base rating and field visits "
        "(offsets interpolated between its visits, claim 6) instead of its published record: "
        "mean %s %% to %s %% over %d visits (z = %s), reduction %s points (%s to %s)."
        % (f1(ipa["baseline"]["mean"]), f1(ipa["method"]["mean"]), ipa["baseline"]["n"], f2(ipa["paired"]["z"]),
           f1(b9["e1_ip_all"]["reduction_pts"]), f1(b9["e1_ip_all"]["ci95"][0]), f1(b9["e1_ip_all"]["ci95"][1])),
        "Offset applied by the data logger before the next visit, with the neighbour's "
        "discharge computed as it would be in real time: mean error at the next visit %s %% "
        "against %s %% for the last offset carried forward (reduction %s points, %s to %s), "
        "and %s %% against %s %% where the offset changed by %s or more (z = %s); the median "
        "error rises from %s %% to %s %% and in open water the two are level, so the gain is "
        "in the large errors."
        % (f1(t2a["method"]["mean"]), f1(t2a["baseline"]["mean"]), f1(b9["e2_rt_all"]["reduction_pts"]),
           f1(b9["e2_rt_all"]["ci95"][0]), f1(b9["e2_rt_all"]["ci95"][1]), f1(t2b["method"]["mean"]),
           f1(t2b["baseline"]["mean"]), BIGM, f2(t2b["paired"]["z"]), f1(t2a["baseline"]["median"]),
           f1(t2a["method"]["median"])),
        "Stage-independent sizing: a stage offset of %s injected into real stage records is "
        "recovered as a median %s (10th-90th percentile %s to %s) whatever the stage, while "
        "the same injection read as a discharge ratio varies with stage (coefficient of "
        "variation %s against %s)."
        % (m3(0.10), m3(e3["0.10"]["median_ft"]), m3(e3["0.10"]["p10_ft"]), m3(e3["0.10"]["p90_ft"]),
           f2(e3["0.10"]["ratio_change_cv"]), f2(e3["0.10"]["stage_units_cv"])),
        "No new apparatus: the inputs are the stage record, the rating, the field-visit "
        "measurements and the neighbour's discharge, all already produced; the computation "
        "takes under %s ms per station-day on one processor core."
        % ("0.1" if e6["ms_per_station_day"] < 0.1 else "%.1f" % (e6["ms_per_station_day"] + 0.05))])
    h2(doc, "4.6 Applicability")
    para(doc,
         "The method applies to any stage-discharge gauging station with a neighbouring "
         "station on the same river: national hydrometric networks, reservoir and canal "
         "operators, flood-warning services and hydropower intakes. It is most useful on "
         "alluvial rivers, whose controls move at floods, and at low flow, where a small "
         "offset is a large share of discharge. In the open discharge records of India's "
         "Central Water Commission (V8), %d pairs of sites on the same river in the %s basins "
         "pass the same admission test as the United States pairs, so neighbouring stations of "
         "the kind the method needs exist there; the method itself has not yet been run on "
         "Indian stage records (Section 8.12)." % (e8["admitted_pairs"], XTRA.basins_text(e8)))
    # ============================================================== 5
    h1(doc, "5. Objective(s) of the Invention")
    para(doc,
         "The primary objective is to provide a system that determines the stage offset of a "
         "river gauging station between field visits from the discharge recorded at a "
         "neighbouring station, and corrects the discharge record and the offset applied by "
         "the station's data logger.")
    bullets(doc, [
        "(a) To separate, in the discrepancy between a neighbour's discharge and the station's "
        "base-rating discharge, the flow ratio between the stations from the station's stage "
        "offset, by fixing the flow ratio at the field visits where the offset is measured.",
        "(b) To convert the remaining discrepancy into an offset in metres through the rating "
        "log-slope, the quantity the data logger applies at any stage.",
        "(c) To admit a neighbouring station only when the travel-time lag and the coupling "
        "of the two discharge records show that both see the same water.",
        "(d) To learn, from the station's own earlier visits, how much of the departure to "
        "apply.",
        "(e) To place an offset change at a flood when it happened rather than where an "
        "interpolation assumes it.",
        "(f) To transmit the offset in force to the data logger of the station so that the "
        "discharge it reports follows the control.",
        "(g) To evidence the correction against discharge measured independently at field "
        "visits."])

    # ============================================================== 6
    h1(doc, "6. Working Principle of the Invention (in brief)")
    num(doc,
        "A gauging station (100) records the stage h of a river with a stage sensor (110) "
        "above a channel control (118). Its data logger (115) converts stage to discharge with "
        "a base rating R and a stage offset s, so that Q = R(h + s), and transmits stage and "
        "discharge over a telemetry link (130) to a record archive (200). A neighbouring "
        "station (120) on the same river does the same with its own sensor (121), logger (125) "
        "and control (128).")
    num(doc,
        "Water is conserved along the reach between the two stations. Where no tributary, "
        "abstraction or storage lies between them, the neighbour's discharge delayed by the "
        "travel time tau equals the station's discharge up to a flow ratio that changes "
        "slowly with the inflow between them. The log discrepancy y(t) = ln Q_Y(t - tau) - "
        "ln R(h(t)) of the station therefore contains two parts: the logarithm c of that flow "
        "ratio, and the effect of the station's stage offset.")
    num(doc,
        "A stage offset s changes ln R by s times the rating log-slope x = d ln R/dh at the "
        "recorded stage, so y = c + s x. The log-slope is large at low stage and small at high "
        "stage, so the same offset moves discharge by tens of percent at low flow and by a "
        "few percent at high flow, while the flow ratio moves y by the same amount at every "
        "stage. A least-squares offset over a window that ignores c, sum(x y)/sum(x^2), "
        "returns s + c sum(x)/sum(x^2): the flow ratio enters it in proportion to the inverse "
        "of the window's mean log-slope, so it changes with the stage of the window even when "
        "nothing at the station has changed.")
    num(doc,
        "The invention removes c with the field visits. At visit P the hydrographer measures "
        "the offset s(P); the flow ratio of the window next to P is then the value c(P) for "
        "which the least-squares offset of y - c on x returns s(P). The same is done at visit "
        "N. Between them the flow ratio, which follows the slowly varying inflow, is "
        "interpolated in time, and the window around each intermediate time t gives the "
        "offset s_raw(t) = sum x (y - c(t)) / sum(x^2) in metres. Where the stage of the window "
        "differs from the stage at the visits, the flow ratio no longer leaks into the "
        "offset.")
    num(doc,
        "The offset between the visits is first interpolated from the measured offsets, as is "
        "done now, giving s_lin(t). Its departure s_raw(t) - s_lin(t) is the part of the "
        "change that the interpolation has not placed. The departure is multiplied by a gain "
        "learned from the station's earlier visits: the regression coefficient, through the "
        "origin, of the measured departures of the offset from interpolation on the estimated "
        "departures. The gain absorbs what the neighbour cannot see, for instance inflow that "
        "changes with the same storms that move the control; it is %s until the station has "
        "%d earlier visits, bounded between 0 and %s." % (f1(GPRI), GMIN, f1(GMAX)))
    num(doc,
        "The offset s(t) so obtained is applied to every interval's recorded stage: the "
        "corrected discharge R(h(t) + s(t)) replaces the provisional record in the record "
        "archive (200) as the corrected record (410). Before the next visit has been made the same principle runs forward from "
        "visit P alone: the flow ratio is fixed at P, s_raw(t) is computed for the window "
        "ending at t, and the offset s(P) + g'(s_raw(t) - s(P)) is sent over the link (130) to "
        "the data logger (115), which applies it to the discharge it reports.")
    h2(doc, "6.1 Claimed Function Against Example Implementation")
    table(doc, ["Claimed function", "Example implementation (Section 8)"], [
        ["Receive stage and visit offsets (claim 1 (a))",
         "USGS 15-minute stage reduced to daily means; expanded shift-adjusted rating table; "
         "field visits with measured discharge and mean gauge height"],
        ["Neighbouring station and lag (claim 1 (b)); admission, claim 3",
         "Consecutive gauges on one mainstem from river-network navigation; incremental "
         "drainage area <= %d %%; lag -2..+2 days maximising correlation of daily changes; "
         "coupling of daily ln Q >= %s" % (round(100 * INCR), f2(MC))],
        ["Log discrepancy and log-slope (claim 1 (c)); steady intervals, claim 4",
         "Days whose mean stage changes by <= %s from the day before and to the day after; "
         "x by central difference over +-%s" % (m3(STEADY), m3(0.05))],
        ["Flow ratio at the visits (claim 1 (d)); window, claim 4", "Window of %d days after P and before N, at least %d steady days" % (WIN, MIND)],
        ["Window offset (claim 1 (e)); window and fit, claims 4 and 5",
         "Window of %d days before and %d days after the time, pooled" % (WIN, WIN)],
        ["Offset between visits and weight (claim 1 (f)); learned gain, claim 5",
         "Linear interpolation in time plus gain x departure; gain %s until %d earlier "
         "visits, bounded 0 to %s" % (f1(GPRI), GMIN, f1(GMAX))],
        ["Corrected record (claim 1 (g)); offset to the logger, claim 2",
         "Q = R(h + s) per interval; offset in force written for the logger"],
    ], "Table 4: Claimed functional combination against the example implementation",
        widths=[2.3, 4.4], size=9, align_left_cols=(0, 1))
    h2(doc, "6.2 Technical Contribution")
    para(doc,
         "The contribution is a change in how a gauging station's offset is determined "
         "between visits: from an assumption about when the control changed (constant, linear "
         "in time, or at a flood) to a measurement-based estimate expressed in the rating's own "
         "units, with the flow ratio between the stations fixed by the station's own field "
         "measurements, and applied by the station's own data logger. The effect is on the "
         "accuracy of the discharge the instrument reports, evidenced against discharge "
         "measured independently by current meter or acoustic Doppler profiler at the field "
         "visits, not against another computed record.")
    h2(doc, "6.3 Drawings")
    bullets(doc, [
        "Figure 1: architecture: the two gauging stations, field measurement, telemetry link "
        "and the data-processing system.",
        "Figure 2: flowchart of the method.",
        "Figure 3: physics and sensitivity: rating log-slope, injected offsets and departures.",
        "Figure 4: accuracy at field visits against current practice, baselines and ablation.",
        "Figure 5: offset between field visits at one station.",
        "Figure 6: pair admission on Indian gauge sites, noise floor of the neighbour "
        "departure, and station-level uncertainty of the main effects.",
        "Figure 7: claim tree."])
    figure(doc, "fig01.png",
           "Figure 1: System architecture for stage-offset determination from a neighbouring gauging station",
           "Stage records of station 100 and discharge of station 120 reach the archive (200) over "
           "link 130; field visits (600) supply measured offsets; the estimator (300) returns the "
           "corrected record (410) and the offset command (420) to logger 115.")
    figure(doc, "fig02.png",
           "Figure 2: Flowchart of the method",
           "S1-S5 fix the flow ratio at the visits, compute the window offsets and the gain; S6a "
           "applies them between two visits (claim 1), S6b forward from the last visit (claim 2); "
           "S7 writes the record and the logger offset.", width=5.2)

    # ============================================================== 7
    h1(doc, "7. Description of the Invention in Detail")
    h2(doc, "7.1 Gauging Station Hardware")
    num(doc,
        "The stage sensor (110) is a submerged pressure transducer in a stilling well or on "
        "the bed, a gas-purge bubbler with an orifice line, or a radar gauge above the water. "
        "The data logger (115) samples it every 15 minutes, applies any gauge-height "
        "correction, computes discharge as R(h + s) from its stored rating table and offset, "
        "and transmits stage and discharge by satellite or cellular link (130). The channel "
        "control (118) is the riffle, bar, weir or channel section whose geometry fixes the "
        "stage-discharge relation; its changes are what the offset represents.")
    h2(doc, "7.2 Records and Rating")
    num(doc,
        "The archive (200) holds the stage series (210), the base rating (220) as a table of "
        "effective stage and discharge, and the visit offsets (230). The offset measured at a "
        "visit is the stage at which the base rating returns the measured discharge less the "
        "mean gauge height recorded during the measurement. Where the published rating table "
        "lists the offset in force at each stage, the base rating is recovered by tabulating "
        "the discharge against stage plus offset.")
    h2(doc, "7.3 Pair Admission (310)")
    num(doc,
        "Neighbouring stations are found from the river network, not chosen by hand: each "
        "station's mainstem is navigated upstream and downstream, gauges on it are ordered by "
        "drainage area, and consecutive gauges whose drainage areas differ by no more than %d "
        "%% form a candidate pair. In the example %d river outlets gave %d mainstem gauges and "
        "%d candidate pairs." % (round(100 * INCR), n_outlets, n_gauges, n_pairs))
    num(doc,
        "For each direction (target station, reference station) the travel-time lag is the "
        "shift, in whole days from -2 to +2, that maximises the correlation of the daily "
        "changes of ln Q at the two stations. The pair is admitted when the correlation of "
        "daily ln Q, after the lag, is at least %s, which shows that the stations see the same "
        "water; %d directions passed." % (f2(MC), n_dir_c))
    h2(doc, "7.4 Discrepancy and Log-Slope (320, 330)")
    num(doc,
        "For each day at which the station's mean stage changes by no more than %s from "
        "the previous day and to the next day (steady flow, so that the daily mean stage "
        "represents the day and travel time is not distorted), the discrepancy unit (320) "
        "computes y = ln Q_Y(d - tau) - ln R(h(d)) and the log-slope unit (330) computes x as "
        "the central difference of ln R over +-%s about h(d)." % (m3(STEADY), m3(0.05)))
    h2(doc, "7.5 Flow Ratio and Window Offset (340)")
    num(doc,
        "For each window the flow-ratio and window-offset unit (340) keeps the number of steady days and the sums of "
        "x, y, x^2 and x y. For the window of %d days after visit P (at least %d steady days) "
        "it computes c(P) = (sum(x y) - s(P) sum(x^2)) / sum(x), the flow ratio for which the "
        "window returns the measured offset s(P), and likewise c(N) for the window before "
        "visit N. For an intermediate time t a fraction f of the way from P to N, c(t) = "
        "c(P) + f (c(N) - c(P)), and the window offset is s_raw(t) = (sum(x y) - c(t) sum(x)) "
        "/ sum(x^2) over the %d days before and the %d days after t. Weighting by x^2 gives most "
        "weight to low-stage days, where the stage sensor (110) and the control (118) "
        "determine discharge most sharply." % (WIN, MIND, WIN, WIN))
    num(doc,
        "Least squares is one way of fitting the stage offset to y - c against x. The same "
        "steps work with any fit that returns an offset in metres from the days of a window "
        "for a given flow ratio and that decreases as the flow ratio increases, so that the "
        "flow ratio at a visit can be found from the measured offset: for instance the median "
        "of the per-day offsets (y - c)/x of the steady days at the stage sensor (110), or "
        "their median weighted by x^2. With either, the flow ratio at a visit is found by "
        "bisection, and the results on the same field visits are given in Table 10.")
    figure(doc, "fig03.png",
           "Figure 3: Physical basis and sensitivity of the stage-unit offset",
           "(a) rating log-slope of one station; (b) response to an injected %s offset in stage "
           "units and as a discharge ratio, against the window's log-slope; (c) recovery of %s "
           "and %s; (d) estimated against measured departure at left-out visits."
           % (m3(0.10), m3(0.10), m3(0.30)))
    h2(doc, "7.6 Offset Estimator (350)")
    num(doc,
        "Between visits P and N the estimator computes s_lin(t) = s(P) + f (s(N) - s(P)) and "
        "s(t) = s_lin(t) + g (s_raw(t) - s_lin(t)). The gain g is the regression coefficient "
        "through the origin of the measured departures s_meas - s_lin on the estimated "
        "departures s_raw - s_lin over the station's earlier visits whose following visit has "
        "already been made, %s until %d are available, bounded to [0, %s]."
        % (f1(GPRI), GMIN, f1(GMAX)))
    num(doc,
        "In real time, after visit P and before N, c(P) is fixed from the window after P, "
        "s_raw(t) is computed for the window ending at t, and s(t) = s(P) + g' (s_raw(t) - "
        "s(P)), with g' learned in the same way from earlier visits (measured change of "
        "offset since the previous visit on the estimated change). The estimate is refreshed "
        "each day and sent to the data logger (115).")
    h2(doc, "7.7 Correction Unit (400)")
    num(doc,
        "The correction unit writes R(h(t) + s(t)) for every 15-minute stage value of the "
        "station to the corrected record (410), in place of the provisional discharge, "
        "keeping the station's own stage dynamics, and sends the offset in force over the "
        "link (130) as an offset command (420) that the data logger (115) stores in place of "
        "the last measured offset. A new field visit replaces the estimate by the measured "
        "offset and becomes visit P for the next interval.")
    num(doc,
        "The neighbour's discharge used between two visits of the station is a record "
        "finalised with offsets from the neighbour's own field visits before and after each "
        "day: its published record, or a record rebuilt from its stage, base rating and visits "
        "with offsets interpolated between them (claim 6). A reference whose own offsets are "
        "anchored by its own visits on both sides carries little unmeasured drift, while a "
        "neighbour computed in real time passes on its own drift (Section 8.3).")
    h2(doc, "7.8 Reference Numerals")
    table(doc, ["Numeral", "Element", "Numeral", "Element"], [
        ["100", "Target gauging station", "300", "Offset estimation"],
        ["110", "Stage sensor", "310", "Pair-admission unit"],
        ["115", "Data logger", "320", "Discrepancy unit"],
        ["118", "Channel control", "330", "Log-slope unit"],
        ["120", "Neighbouring gauging station", "340", "Flow-ratio and window-offset unit"],
        ["121, 125, 128", "Its sensor, logger, control", "350", "Offset estimator"],
        ["130", "Telemetry link", "400", "Correction unit"],
        ["200", "Record archive", "410", "Corrected discharge record"],
        ["210", "Stage and discharge series", "420", "Offset command"],
        ["220", "Base rating", "600", "Field visit"],
        ["230", "Visit offsets", "610", "Discharge measurement"],
        ["900", "Data-processing system", "620", "Field-visit record"],
    ], "Table 5: Reference numerals", widths=[0.9, 2.45, 0.9, 2.45], size=9, align_left_cols=(1, 3))

    # ============================================================== 8
    h1(doc, "8. Experimental Validation Results")
    h2(doc, "8.1 Data and Configuration")
    para(doc,
         "All inputs are public records of the United States Geological Survey, read without "
         "registration: daily mean stage and discharge for water years 2021-2026 (%s to %s), "
         "the expanded shift-adjusted rating of each station, field visits with measured "
         "discharge and mean gauge height, and the river network for topology. The records are "
         "in feet and are converted here to metres (1 ft = 0.3048 m). Only visits made after "
         "the start date of the station's current rating are scored. Of the %d admitted "
         "directions, %d had retrievable field visits and a current rating covering enough "
         "visits to be scored. The error at a visit is |R(GH + s_est)/Q_meas - 1|: the base "
         "rating read at the visit's gauge height plus the estimated offset, against the "
         "discharge measured at the visit. Table 6 lists the operating constants, read from "
         "the code that used them; they were fixed before the comparisons with practice "
         "baselines, the station halves and the sensitivity runs of Section 8.9, and were "
         "not changed for them."
         % (meta["period"][0], meta["period"][1], n_dir_c, n_scored))
    table(doc, ["Parameter", "Value", "Role"], [
        ["Incremental drainage area", "<= %d %%" % round(100 * INCR), "candidate pair"],
        ["Coupling of daily ln Q", ">= %s" % f2(MC), "pair admission"],
        ["Travel-time lag", "-2 .. +2 days", "maximum correlation of daily changes"],
        ["Steady day", "<= %s change" % m3(STEADY), "interval selection"],
        ["Window", "%d days, >= %d steady days" % (WIN, MIND), "flow ratio and window offset"],
        ["Gain", "%s until %d visits; 0 .. %s" % (f1(GPRI), GMIN, f1(GMAX)), "departure weight"],
        ["Large departure", ">= %s" % BIGM, "subset reported separately"],
    ], "Table 6: Configuration parameters", widths=[2.1, 1.9, 2.7], size=9, align_left_cols=(0, 2))
    h2(doc, "8.2 Tests")
    table(doc, ["Test", "Question", "Ground truth"], [
        ["V1", "Record between visits: is the offset at a left-out visit better predicted with the "
               "neighbour than by linear interpolation, or by placing the change at the largest "
               "flood between the visits?", "Measured offset at the left-out visit"],
        ["V2", "Same, with the neighbour's discharge computed as in real time", "Same"],
        ["V3", "Real time: is the offset at the next visit better predicted than by carrying the last offset forward?", "Measured offset at the next visit"],
        ["V4", "Injected offsets and a synthetic flow-ratio change: is a known offset recovered "
               "in metres at every stage, and is a flow-ratio change kept out of it?",
         "Injected %s and %s" % (m3(0.10), m3(0.30))],
        ["V5", "Ablation: flow-ratio term, stage units against discharge-ratio units, gain", "As V1"],
        ["V6", "Processing time", "Wall clock"],
        ["V7", "Same as V1 with the rating applied to every 15-minute stage", "As V1"],
        ["V8", "Indian gauge sites: are neighbouring sites admitted, and how large is the "
               "neighbour's departure noise?", "Observed CWC discharge"],
        ["V9", "Do the effects hold when whole stations are resampled, on two disjoint halves "
               "of the stations, and when each constant is changed?", "As V1 and V3"],
    ], "Table 7: Validation tests", widths=[0.5, 4.0, 2.2], size=9, align_left_cols=(1, 2))

    h2(doc, "8.3 Record Between Visits (V1, V2)")
    para(doc,
         "Each visit T with a visit before (P) and after (N) is left out, its offset is "
         "estimated from P, N and the records, and the discharge given by the estimated offset "
         "at T's gauge height is compared with the discharge measured at T. The gain uses only "
         "earlier visits whose following visit precedes T, so the measurement at T never "
         "enters its own estimate. %d visits at %d stations (%d directions) were scored. The "
         "estimated departure correlates with the measured departure of the offset from "
         "linear interpolation at r = %s (Figure 3(d)); Table 8 and Figure 4 give the errors."
         % (pub["n"], pub["targets"], pub["directions"], f2(pub["r_dev"])))

    def row(lab, blk):
        return [lab, str(blk["baseline"]["n"]), mmp(blk["baseline"]), mmp(blk["method"]),
                "%d of %d" % (blk["paired"]["better"], blk["paired"]["of"])]
    table(doc, ["Subset", "Visits", "Linear interpolation: median / mean / p90 (%)",
                "This method: median / mean / p90 (%)", "Better at"], [
        row("All visits", ra),
        row("Open water (April-November, no ice at visit)", ro),
        row("Departure >= %s" % BIGM, rb),
        row("Open water, departure >= %s" % BIGM, rob),
        row("All, neighbour record rebuilt from its own visits", ipa),
        row("Open water, neighbour record rebuilt", ipo),
        row("All, real-time reference", rtm["all"]),
        row(">= %s, real-time reference" % BIGM, rtm["big"]),
    ], "Table 8: Absolute discharge error at left-out field visits",
        widths=[1.45, 0.6, 1.85, 1.85, 0.95], size=9, align_left_cols=(0,))
    para(doc,
         "Over all visits the error falls at more visits than it rises (z = %s), in open water "
         "too (z = %s), and on both halves of the interval lengths (P to N up to %d days: z = "
         "%s; longer: z = %s). The departure subset is defined by the measured offset and so "
         "cannot be selected in operation; it shows where the benefit comes from, not a rate "
         "to expect. The mean offset error falls from %s to %s. The hydrographer's alternatives "
         "were scored on the same visits (Table 10): placing the whole change at the largest "
         "flood between the visits, or prorating up to it, is worse than linear interpolation "
         "over all visits (mean %s %% and %s %%) and better where the control moved (%s %%). "
         "Choosing for each station, from its earlier visits only, the practice that served it "
         "best gives %s %% over all visits, and the method is better than that choice at %d of "
         "%d visits (z = %s) and better than every practice baseline in every column of Table 10."
         % (f2(ra["paired"]["z"]), f2(ro["paired"]["z"]), int(pub["gap_median_days"]),
            f2(pub["short_gap"]["paired"]["z"]), f2(pub["long_gap"]["paired"]["z"]),
            m3(pub["abs_offset_err_m"]["lin"]["mean"] / FT), m3(pub["abs_offset_err_m"]["method"]["mean"] / FT),
            f1(ra["baselines"]["e_step"]["mean"]), f1(ra["baselines"]["e_ramp"]["mean"]),
            f1(rb["baselines"]["e_step"]["mean"]), f1(e4["all"]["practice_per_station"]["mean"]),
            PR["better"], PR["of"], f2(PR["z"])))
    para(doc,
         "With the neighbour's discharge computed as in real time (its base rating and its "
         "own last measured offset), the neighbour carries its own unmeasured offset change: "
         "the error still falls where the control moved (better at %d of %d visits, z = %s) "
         "but not over all visits (z = %s). The record between visits is therefore computed "
         "with a neighbour record finalised with the neighbour's own field visits on both sides "
         "(step (b) of claim 1): with the published record, and equally with a record "
         "rebuilt from the neighbour's stage, base rating and visits (claim 6; %d visits, z = %s, open water "
         "z = %s). "
         "At the %d visits with ice reported at the control the error falls from %s %% to "
         "%s %%; ice-affected records of the neighbour may themselves have been estimated from "
         "nearby stations [N1], so these visits are not used as evidence. Figure 5 shows the "
         "offset between visits for the direction with the most scored visits (%d visits, "
         "mean error %s %% with linear interpolation and %s %% with the method)."
         % (rtm["big"]["paired"]["better"], rtm["big"]["paired"]["of"], f2(rtm["big"]["paired"]["z"]),
            f2(rtm["all"]["paired"]["z"]), ipa["baseline"]["n"], f2(ipa["paired"]["z"]),
            f2(ipo["paired"]["z"]), pub["ice"]["baseline"]["n"], f1(pub["ice"]["baseline"]["mean"]),
            f1(pub["ice"]["method"]["mean"]), ex["n"], f1(ex["e_lin_mean"]), f1(ex["e_hat_mean"])))
    figure(doc, "fig04.png",
           "Figure 4: Accuracy of the discharge record at field visits",
           "(a) cumulative error, linear interpolation against this method; (b) record between "
           "visits; (c) real time against the last offset carried forward; (d) practice "
           "baselines and ablation on the same visits.")
    figure(doc, "fig05.png",
           "Figure 5: Offset of one gauging station between field visits",
           "(a) station %s with neighbour %s: offsets measured at visits (dots), linear "
           "interpolation in time (red) and this method with the station's gain of %s (blue); "
           "(b) daily mean stage. Between visits the blue line departs from the red where the "
           "neighbour's record shows the offset moving between the visits."
           % (ex["x"], ex["y"], f2(ex["gain"])))

    h2(doc, "8.4 Real Time (V3)")
    para(doc,
         "Before the next visit only data before it are available. The offset applied is the "
         "last measured offset plus the gain times the change of the window offset since that "
         "visit, the flow ratio fixed at that visit. With the neighbour's discharge computed "
         "as in real time the estimated change correlates with the measured change at r = %s "
         "over %d visits. The mean error falls over all visits, with an interval from station "
         "resampling of %s to %s points, while the median rises from %s %% to %s %%: the "
         "method removes large errors where the control moved and adds small ones elsewhere. "
         "In open water alone the two are level (z = %s). Using, for the neighbour, its record "
         "finalised to date (offsets interpolated between its own visits already made, the "
         "last offset carried after them) lowers the mean error to %s %% but does not make "
         "the improvement more frequent (z = %s): the neighbour's drift since its own last "
         "visit, which no real-time record can remove, limits this use."
         % (f2(e2["realtime"]["r_change"]), e2["realtime"]["n"], f1(b9["e2_rt_all"]["ci95"][0]),
            f1(b9["e2_rt_all"]["ci95"][1]), f1(t2a["baseline"]["median"]), f1(t2a["method"]["median"]),
            f2(t2o["paired"]["z"]), f1(e2["rtfinal"]["all"]["method"]["mean"]),
            f2(e2["rtfinal"]["all"]["paired"]["z"])))
    table(doc, ["Reference record", "Subset", "Visits", "Carried forward: mean / p90 (%)",
                "This method: mean / p90 (%)", "Better at"], [
        [ref, sub, str(b["baseline"]["n"]),
         "%s / %s" % (f1(b["baseline"]["mean"]), f1(b["baseline"]["p90"])),
         "%s / %s" % (f1(b["method"]["mean"]), f1(b["method"]["p90"])),
         "%d of %d" % (b["paired"]["better"], b["paired"]["of"])]
        for ref, sub, b in (("real time", "all", t2a), ("real time", "open water", t2o),
                            ("real time", "change >= %s" % BIGM, t2b),
                            ("corrected", "all", r2a), ("corrected", "change >= %s" % BIGM, r2b))],
        "Table 9: Real-time offset at the next field visit",
        widths=[1.0, 1.2, 0.6, 1.55, 1.45, 0.9], size=9, align_left_cols=(0, 1))

    h2(doc, "8.5 Injected Offsets (V4)")
    para(doc,
         "A stage offset was injected into the real stage record of the window before each "
         "visit (the sensor reading low by the injected amount) and the change of the window "
         "offset, with the flow ratio fixed at the previous visit as in claim 2, recovered from "
         "the same neighbour data. Because an offset moves y by s x except for the curvature "
         "of the rating, this test checks the conversion and its independence of stage, not "
         "the separation from the flow ratio, which the second test checks. For %s the recovered change has "
         "median %s (10th-90th percentile %s to %s, n = %d); for %s, median %s (%s to %s). The "
         "same injection read as a change of the median discharge ratio varies with the "
         "stage of the window (coefficient of variation %s and %s, against %s and %s in stage "
         "units; Figure 3(b), 3(c)). The shortfall at the larger offset comes from the "
         "curvature of the rating: the log-slope read at the lowered stage is steeper than its "
         "average over the step, so the linearised conversion slightly under-sizes large "
         "offsets."
         % (m3(0.10), m3(e3["0.10"]["median_ft"]), m3(e3["0.10"]["p10_ft"]), m3(e3["0.10"]["p90_ft"]),
            e3["0.10"]["n"], m3(0.30), m3(e3["0.30"]["median_ft"]), m3(e3["0.30"]["p10_ft"]),
            m3(e3["0.30"]["p90_ft"]), f2(e3["0.10"]["ratio_change_cv"]),
            f2(e3["0.30"]["ratio_change_cv"]), f2(e3["0.10"]["stage_units_cv"]),
            f2(e3["0.30"]["stage_units_cv"])))
    RL, RS = e4["ratio_leak_m"], e4["ratio_step_m"]
    para(doc,
         "In the second test a flow-ratio change of %s %% was added to the log discrepancies "
         "between visits P and N, with no change of offset, in two forms. Growing linearly in "
         "time, the form the interpolation of the flow ratio assumes, it moves the offset "
         "estimated at the left-out visits by a median of %s with the flow-ratio term and %s "
         "without it; this form shows only that the term behaves as designed. As a step at "
         "the largest flood between the visits, as when a storm changes the inflow between the "
         "stations, it moves the offset by a median of %s with the term and %s without (mean "
         "%s and %s); the term gives the smaller movement at %d of %d visits (z = %s), and the "
         "movement exceeds %s at %s %% of visits with the term and %s %% without. The term "
         "therefore reduces, but does not remove, the effect of a flow-ratio change that does "
         "not follow the interpolation."
         % (f1(100 * (math.exp(RL["ratio_change_ln"]) - 1)), "%.4f m" % RL["with_ratio_term"]["median"],
            "%.4f m" % RL["without_ratio_term"]["median"], "%.4f m" % RS["with_ratio_term"]["median"],
            "%.4f m" % RS["without_ratio_term"]["median"], "%.4f m" % RS["with_ratio_term"]["mean"],
            "%.4f m" % RS["without_ratio_term"]["mean"], RS["paired"]["better"], RS["paired"]["of"],
            f2(RS["paired"]["z"]), BIGM, f1(100 * RS["share_over_big_with"]), f1(100 * RS["share_over_big_without"])))

    h2(doc, "8.6 Baselines and Ablation (V5) and Processing Time (V6)")
    A4 = e4

    def arow(lab, k):
        return [lab, mmp(A4["all"][k]), f1(A4["open_water"][k]["mean"]), f1(A4["big"][k]["mean"])]
    table(doc, ["Variant", "All visits: median / mean / p90 (%)", "Open water: mean (%)",
                "Departure >= %s: mean (%%)" % BIGM], [
        arow("Linear interpolation in time [N1]", "linear_time"),
        arow("Whole change at the largest flood [N5]", "flood_step"),
        arow("Prorated up to the largest flood [N6]", "flood_ramp"),
        arow("Change split over the rises of %s or more [N5, N8]" % m3(EX.RISE_FT), "per_rise"),
        arow("Practice chosen per station from earlier visits", "practice_per_station"),
        arow("Neighbour, no flow-ratio term", "no_ratio_term"),
        arow("Neighbour, discharge-ratio units", "ratio_units"),
        arow("Neighbour, stage units, median fit", "fit_median"),
        arow("Neighbour, stage units, x^2-weighted median fit", "fit_weighted_median"),
        arow("Neighbour, stage units, g = 1", "stage_units_no_gain"),
        arow("Neighbour, stage units, gain (claim 5)", "stage_units_gain"),
    ], "Table 10: Practice baselines and ablation on the same left-out visits",
        widths=[2.3, 2.1, 1.1, 1.2], size=9, align_left_cols=(0,))
    para(doc,
         "The fit of the offset is not what carries the effect. With the median of the per-day "
         "offsets, or their x^2-weighted median, in place of least squares, the mean error over "
         "all visits is %s %% and %s %% (least squares %s %%), and %s %% and %s %% where the "
         "control moved by %s or more (least squares %s %%); every fit has a lower mean error "
         "than linear interpolation and than the practice chosen per station in every column of "
         "Table 10, and against that practice over all visits the median fit is better at %d of "
         "%d visits (z = %s) and the weighted median at %d of %d (z = %s). The weight of the "
         "flow-ratio term and of stage units (Table 10, least squares) was measured with the "
         "least-squares fit only."
         % (f1(A4["all"]["fit_median"]["mean"]), f1(A4["all"]["fit_weighted_median"]["mean"]),
            f1(A4["all"]["stage_units_gain"]["mean"]), f1(A4["big"]["fit_median"]["mean"]),
            f1(A4["big"]["fit_weighted_median"]["mean"]), BIGM, f1(A4["big"]["stage_units_gain"]["mean"]),
            A4["all"]["e_median_vs_practice"]["better"], A4["all"]["e_median_vs_practice"]["of"],
            f2(A4["all"]["e_median_vs_practice"]["z"]), A4["all"]["e_wmedian_vs_practice"]["better"],
            A4["all"]["e_wmedian_vs_practice"]["of"], f2(A4["all"]["e_wmedian_vs_practice"]["z"])))
    para(doc,
         "Fixing the flow ratio at the visits lowers the error against the same estimator "
         "without it (better at %d of %d visits, z = %s). Stage units beat discharge-ratio "
         "units (better at %d of %d, z = %s), and they do so in each third of the visits "
         "ordered by how far the log-slope at the visit differs from that of the window (z = "
         "%s, %s and %s). The advantage is about the same in each third, so it is not explained "
         "by the slope difference alone; the stage-unit form weights each day by its "
         "log-slope, so that low-stage days, where an offset is most visible, count most, and "
         "it gives the offset in metres that the data logger (115) applies. Over all visits the gain lowers the mean error from %s %% "
         "(g = 1) to %s %%, because it scales the departure to what the station's own history "
         "supports; where the control moved by %s or more, g = 1 is lower (%s %% against %s %%), "
         "so the gain trades some of the large-change benefit for protection of the other "
         "visits. The full computation for one station over six years takes under one second "
         "on one processor core (median %s s over %d stations, %s ms per station-day in this "
         "run; wall-clock times vary from run to run)."
         % (A4["all"]["with_vs_without_ratio_term"]["better"], A4["all"]["with_vs_without_ratio_term"]["of"],
            f2(A4["all"]["with_vs_without_ratio_term"]["z"]), A4["all"]["stage_vs_ratio"]["better"],
            A4["all"]["stage_vs_ratio"]["of"], f2(A4["all"]["stage_vs_ratio"]["z"]),
            f2(slope_terc[0]["paired"]["z"]), f2(slope_terc[1]["paired"]["z"]), f2(slope_terc[2]["paired"]["z"]),
            f1(A4["all"]["stage_units_no_gain"]["mean"]), f1(A4["all"]["stage_units_gain"]["mean"]),
            BIGM, f1(A4["big"]["stage_units_no_gain"]["mean"]), f1(A4["big"]["stage_units_gain"]["mean"]),
            f2(e6["seconds_per_direction_median"]), e6["directions"], e6["ms_per_station_day"]))

    h2(doc, "8.7 15-Minute Records (V7)")
    para(doc,
         "On the directions whose 15-minute records were retrieved (water years 2024-2026), "
         "the base rating and its log-slope were applied to every 15-minute stage, the "
         "neighbour's discharge lagged in hours, and the values averaged per day before the "
         "same windows were formed. On the %d left-out visits common to both versions (%d "
         "stations) the method's mean error is %s %% with 15-minute records and %s %% with "
         "daily means, against %s %% for linear interpolation."
         % (e7["n_common"], e7["targets"], f1(e7["iv"]["method"]["mean"]),
            f1(e7["daily"]["method"]["mean"]), f1(e7["iv"]["linear"]["mean"]))
         if IV_OK else
         "The 15-minute records were retrieved for only %d left-out visits, too few to report; "
         "the daily-mean results above stand for the method." % e7.get("n_common", 0))
    XTRA.india_section(doc, h2, para, figure, e8, e9, pub, BIG, var, ho)
    h2(doc, "8.10 Technical Effect Summary")
    table(doc, ["Effect", "Mechanism", "Evidence"], [
        ["Lower error of the discharge record between visits", "Neighbour departure in stage units, flow ratio fixed at visits, gained",
         "Table 8, all visits: mean %s -> %s %%; %s points (%s to %s)"
         % (f1(ra["baseline"]["mean"]), f1(ra["method"]["mean"]), f1(b9["e1_all"]["reduction_pts"]),
            f1(b9["e1_all"]["ci95"][0]), f1(b9["e1_all"]["ci95"][1]))],
        ["Same in open water", "As above",
         "Table 8: mean %s -> %s %%; %s points (%s to %s)"
         % (f1(ro["baseline"]["mean"]), f1(ro["method"]["mean"]), f1(b9["e1_open_water"]["reduction_pts"]),
            f1(b9["e1_open_water"]["ci95"][0]), f1(b9["e1_open_water"]["ci95"][1]))],
        ["Record better than hydrographer practice chosen per station", "Departure from the interpolated offset",
         "Table 10, all visits: %s -> %s %%; %s points (%s to %s)"
         % (f1(e4["all"]["practice_per_station"]["mean"]), f1(e4["all"]["stage_units_gain"]["mean"]),
            f1(b9["e1_vs_practice"]["reduction_pts"]), f1(b9["e1_vs_practice"]["ci95"][0]),
            f1(b9["e1_vs_practice"]["ci95"][1]))],
        ["Flow ratio separated from the offset", "Flow ratio fixed at the visits",
         "Table 10: %d of %d visits better with the term (z = %s); V4 step: smaller movement at %s %% of visits"
         % (e4["all"]["with_vs_without_ratio_term"]["better"], e4["all"]["with_vs_without_ratio_term"]["of"],
            f2(e4["all"]["with_vs_without_ratio_term"]["z"]), f1(100 * e4["ratio_step_m"]["paired"]["share"]))],
        ["Better offset in the data logger before the next visit", "Change of window offset since the last visit, gained",
         "Table 9, real-time reference, all: mean %s -> %s %%; %s points (%s to %s); median "
         "%s -> %s %%; open water level"
         % (f1(t2a["baseline"]["mean"]), f1(t2a["method"]["mean"]), f1(b9["e2_rt_all"]["reduction_pts"]),
            f1(b9["e2_rt_all"]["ci95"][0]), f1(b9["e2_rt_all"]["ci95"][1]), f1(t2a["baseline"]["median"]),
            f1(t2a["method"]["median"]))],
        ["Offset sized in metres at any stage", "Least squares on the rating log-slope",
         "V4: %s recovered as %s" % (m3(0.10), m3(e3["0.10"]["median_ft"]))],
    ], "Table 11: Technical effects, mechanism and evidence", widths=[2.0, 2.2, 2.5], size=9, align_left_cols=(0, 1, 2))
    h2(doc, "8.11 Insights")
    bullets(doc, [
        "The offsets of alluvial stations move a lot between visits: %d of %d left-out visits "
        "depart from linear interpolation by %s or more." % (pub["n_big"], pub["n"], BIGM),
        "The neighbour sees those movements: the estimated and measured departures correlate "
        "at r = %s." % f2(pub["r_dev"]),
        "Fixing the flow ratio at the visits lowers the error on real records (Table 10) and "
        "reduces, without removing, the part of a flow-ratio change that reaches the offset "
        "(Section 8.5).",
        "A neighbour that is itself drifting passes its drift on; the gain learned per "
        "station limits the damage, and a neighbour record finalised with its own visits "
        "gives the benefit of Table 8."])
    h2(doc, "8.12 Limitations and Acceptance Criteria")
    bullets(doc, [
        "Leaving a visit out doubles the interval over which the offset is interpolated (median "
        "%d days from P to N, against %d days between consecutive visits), so the errors of "
        "Table 8 are larger than those of a record computed with every visit. The reduction is "
        "of similar size on the shorter and the longer half of the intervals (Section 8.3)."
        % (int(pub["gap_median_days"]), gap_med),
        "In real time over all visits the mean error falls but the median rises (Section 8.4). "
        "Acceptance: before use in a data logger, the station's own earlier visits must show "
        "a gain above zero.",
        "Detecting, before the visit, which intervals hold a large offset change is weak "
        "(area under the curve %s); the method sizes changes better than it flags them, and "
        "no field-visit trigger is claimed." % f2(e5["auc_abs"]),
        ("Evaluated on United States stations; the 15-minute check (V7) covers %d visits. "
         "Indian gauge sites were tested only for pair admission and departure noise (V8), "
         "because the open Indian record has discharge but no stage; the departure noise is %s "
         "times that of the United States pairs. Acceptance: on each new network, and with the "
         "stage records and gaugings of admitted Central Water Commission sites, V1 must show a "
         "lower mean error than linear interpolation."
         % (e7.get("n_common", 0), f1(e8["india_floor"]["median"] / e8["us_floor_upper"]["median"]))),
        "Ice on the control is excluded from the evidence; the open-water results of Table 8 "
        "are the basis of the acceptance criteria.",
        "Stations whose daily discharge correlates with the neighbour below %s are not "
        "admitted; in this example %d of %d candidate directions were." % (f2(MC), n_dir_c, 2 * n_pairs),
        "A stage-dependent (variable) offset is represented by one value for the stages of "
        "each window; where a station's offset varies strongly with stage the window should "
        "be restricted to stages near the one at which the offset is applied."])

    # ============================================================== 9
    h1(doc, "9. Discussion")
    para(doc,
         "The invention turns a neighbouring station from a picture that a hydrographer looks "
         "at into a measurement of the station's own rating offset. Two physical facts make "
         "this possible. A stage offset changes the log discrepancy in proportion to the "
         "rating log-slope, while the flow ratio between the stations changes it by the same "
         "amount at every stage; and at a field visit the offset is known, so the flow ratio "
         "there can be computed rather than assumed. Fixing the flow ratio at the visits and "
         "converting the rest through the log-slope gives an offset in metres that can be "
         "compared with, and applied like, a measured one. The results show the gain over all "
         "visits and in open water, against linear interpolation and against the practice "
         "best suited to each station, and they show the boundary honestly: a reference that is itself "
         "drifting in real time limits the benefit.")

    # ============================================================== 10
    h1(doc, "10. What Aspect(s) of the Invention Need(s) Protection?")
    bullets(doc, [
        "Fixing the flow ratio between a station and its neighbour at the station's field "
        "visits from the offsets measured there.",
        "Conversion of the remaining discrepancy into a stage offset in metres through the "
        "rating log-slope.",
        "Determination of the offset between two field visits as the interpolated offset "
        "plus a gain times the departure of the window offset from it, the gain learned from "
        "the station's earlier field visits.",
        "Admission of the neighbouring station by travel-time lag and coupling.",
        "Replacing the discharge record between visits, and transmitting the offset in force "
        "to the data logger of the gauging station before the next visit.",
        "Using as the neighbour's discharge a record finalised with offsets from the "
        "neighbour's own field visits on both sides."])
    figure(doc, "fig07.png", "Figure 7: Claim tree",
           "Independent method claim 1 (record between visits), its dependents 2 (offset sent to "
           "the data logger before the next visit) and 3 to 6, and system claim 7 carrying out "
           "the method.")
    h2(doc, "Claims")
    CLAIMS(doc, MC, INCR, STEADY, WIN, MIND, GPRI, GMIN, GMAX)

    # ============================================================== 11
    h1(doc, "11. Technology Readiness Level of the Invention")
    para(doc,
         "The method runs end to end on public records of operating gauging stations and is "
         "scored against independent field measurements; it has not yet been connected to a "
         "station's data logger.")
    trl_table(doc)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    B._run(p, "END OF THE DOCUMENT", size=11, bold=True)
    p.alignment = 1
    return doc, num


TITLE = ("System and Method for Correcting a River Gauging Station's Stage Offset from "
         "Neighbouring-Station Discharge")


def CLAIMS(doc, MC, INCR, STEADY, WIN, MIND, GPRI, GMIN, GMAX):
    """One independent method claim, five dependents, one system claim.

    The independent claim carries only what the evidence needs: the neighbour record
    finalised with its own visits (without it the all-visit gain is lost, Section 8.3),
    the log discrepancy on the rating log-slope, the flow ratio fixed at the visits, the
    window offset in metres and the corrected record. Pair admission, steady intervals,
    window constants, the least-squares fit and the learned gain are dependents."""
    station = (
        "a river gauging station (100) comprising a stage sensor (110) that records the stage "
        "of the river above a channel control (118) and a data logger (115) that stores a "
        "base stage-discharge rating R and a stage offset and computes discharge by reading "
        "the base rating at the recorded stage plus the stage offset")
    claim(doc, 1,
          "A method of correcting the discharge record of " + station + ", between a first and "
          "a later second field visit, the data logger transmitting the recorded stage and the "
          "discharge over a telemetry link (130) to a record archive (200), the method being "
          "carried out by a data-processing system (900) and comprising: "
          "(a) receiving over the link (130) the stage h recorded by the stage sensor (110) at "
          "successive intervals, and the stage offset measured at each of the first and second "
          "field visits as the stage at which the base rating gives the discharge measured at "
          "that visit by a current meter or acoustic Doppler profiler less the stage recorded "
          "at that visit; "
          "(b) receiving a discharge record of a neighbouring gauging station (120) on the same "
          "river, computed with stage offsets determined from offsets measured at field visits "
          "of that station made before and after each interval, and delaying it by the "
          "travel-time lag between the two stations; "
          "(c) computing, for the intervals, a log discrepancy y, being the logarithm of the "
          "delayed discharge of the neighbouring station less the logarithm of R(h), and a "
          "rating log-slope x, being the derivative of the logarithm of R with respect to stage "
          "at h, whereby a stage offset s and a flow ratio c between the stations give "
          "y = c + s x; "
          "(d) determining, at each of the two field visits, the flow ratio c for which a stage "
          "offset fitted to y - c against x over the intervals of a window next to that visit "
          "equals the offset measured at that visit, and interpolating the flow ratio in time "
          "between the two visits; "
          "(e) computing, for each intermediate time, a window offset in units of stage by "
          "fitting a stage offset to y - c against x over the intervals of a window containing "
          "that time, c being the interpolated flow ratio at that time; "
          "(f) estimating the stage offset at each intermediate time as an offset interpolated "
          "between the offsets measured at the two visits plus a weight times the difference "
          "between the window offset and the interpolated offset; and "
          "(g) computing the discharge of each interval between the two visits by reading the "
          "base rating at the recorded stage plus the estimated offset, and replacing in the "
          "record archive (200) the discharge computed by the data logger (115) for those "
          "intervals with the discharge so computed.", independent=True)
    claim(doc, 2,
          "The method of claim 1, further comprising, after the second field visit and before "
          "a next field visit: receiving the discharge of the neighbouring station (120) as "
          "computed to date, with offsets interpolated between its own field visits already "
          "made and the offset of its latest visit carried forward after it; determining the "
          "flow ratio at the second visit as in step (d) from a window following it; computing "
          "the window offset for a window ending at a current time with that flow ratio; "
          "estimating the stage offset at the current time as the offset measured at the "
          "second visit plus a weight times the difference between the window offset and that "
          "offset; and transmitting the estimated offset over the link (130) to the data logger "
          "(115), which stores it in place of the last measured offset and computes the "
          "discharge it reports with it until the next field visit.")
    claim(doc, 3,
          "The method of claim 1 or claim 2, wherein the neighbouring station (120) lies on the "
          "same mainstem with a drainage area differing from that of the station by no more "
          "than %d %%, the travel-time lag is the whole number of days from -2 to +2 that "
          "maximises the correlation of the daily changes of the logarithms of the two "
          "discharges, and the neighbouring station is admitted only when the correlation of "
          "the logarithms of the two daily mean discharges after the lag is at least %s."
          % (round(100 * INCR), "%.2f" % MC))
    claim(doc, 4,
          "The method of claim 1 or claim 2, wherein the intervals are days and steps (c) to "
          "(e) use only days whose mean stage differs from that of the day before and of the "
          "day after by no more than %s; the window next to a visit spans the %d days after "
          "the first visit or before the second visit, and the window containing an "
          "intermediate time spans the %d days before and the %d days after it; a flow ratio "
          "or window offset is computed only from a window holding at least %d such days; and "
          "the interpolated offset and the flow ratio are each interpolated linearly in time "
          "between the two visits." % (m3(STEADY), WIN, WIN, WIN, MIND))
    claim(doc, 5,
          "The method of claim 1 or claim 2, wherein each stage offset fitted to y - c against "
          "x is the sum of x (y - c) divided by the sum of x squared over the intervals of the "
          "window, and the weight is a gain learned from earlier field visits of the station "
          "as the regression coefficient, through the origin, of the measured differences "
          "between offset and interpolated offset, or for claim 2 the measured changes of "
          "offset between consecutive visits, on the corresponding estimated differences or "
          "changes, the gain being %s until at least %d earlier visits are available and "
          "bounded between 0 and %s." % ("%.1f" % GPRI, GMIN, "%.1f" % GMAX))
    claim(doc, 6,
          "The method of claim 1, wherein the discharge record of the neighbouring station "
          "(120) is computed by reading the base rating of that station at its recorded stage "
          "plus stage offsets interpolated linearly in time between the offsets measured at "
          "its own field visits before and after each interval.")
    claim(doc, 7,
          "A system for correcting the discharge record of a river gauging station between "
          "field visits, comprising: " + station + "; a telemetry link (130); a record archive "
          "(200) receiving over the link the recorded stage and discharge of the station, the "
          "discharge record of a neighbouring gauging station (120) on the same river and the "
          "stage offsets measured at field visits (230); and a data-processing system (900) "
          "comprising a discrepancy unit (320), a log-slope unit (330), a flow-ratio and "
          "window-offset unit (340), an offset estimator (350) and a correction unit (400), "
          "configured to carry out the method of any one of claims 1 to 6.", independent=True)


def check(path, num):
    d = Document(path)
    paras = [p.text for p in d.paragraphs]
    tables = [[c.text for row in t.rows for c in row.cells] for t in d.tables]
    body = "\n".join(paras)
    prior = [c for cells in tables if cells[:1] == ["S.No."] and "Comments" in cells[:4] for c in cells]
    other = [c for cells in tables for c in cells if c not in prior]
    own = body + "\n" + "\n".join(other)
    everything = own + "\n" + "\n".join(prior)
    bad = [w for w in FORBIDDEN if re.search(w, everything, re.I)]
    ban = [w for w in BANNED if re.search(w, own, re.I)]
    ph = PLACEHOLDER.findall(everything)
    open_ = [p[:40] for p in paras if OPENERS.match(p)]
    phys = sum(1 for t in num.texts if PHYSICAL.search(t))
    share = phys / len(num.texts)
    refs = set(re.findall(r"Table (\d+)", everything))
    caps = set(re.findall(r"^Table (\d+):", "\n".join(paras), re.M))
    figs = set(re.findall(r"Figure (\d+)", everything))
    problems = []
    if bad:
        problems.append("business vocabulary %s" % bad)
    if ban:
        problems.append("banned phrases %s" % ban)
    if ph:
        problems.append("placeholders %s" % ph)
    if open_:
        problems.append("banned paragraph openers %s" % open_)
    if share <= 0.6:
        problems.append("physical tie-in only %.0f %%" % (100 * share))
    if refs - caps:
        problems.append("tables referenced without caption %s" % sorted(refs - caps))
    if problems:
        raise SystemExit("refused: " + "; ".join(problems))
    print("checked: %d numbered paragraphs, %.0f %% name a physical element; tables %s; figures %s; "
          "%d words" % (len(num.texts), 100 * share, sorted(caps, key=int), sorted(figs, key=int),
                        len(everything.split())))


if __name__ == "__main__":
    doc, num = build()
    for p in doc.paragraphs:          # keep each figure on the page of its caption
        if p._p.xpath(".//w:drawing"):
            p.paragraph_format.keep_with_next = True
    doc.save(OUT)
    print("wrote", OUT)
    check(OUT, num)
